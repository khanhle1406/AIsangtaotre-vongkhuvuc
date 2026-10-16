"""Observable scope/ranking invariants, independent of reference answers."""

import datetime
import re
from typing import Any

from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries

from .executor import ExecutionResult


def canonical_label(value: Any) -> str:
    return " ".join(re.sub(r"(?<=[A-Za-z])-\s*\n\s*(?=[a-z])", "", str(value)).split())


def cell_region(sheet: dict[str, Any], row: int, col: int) -> dict[str, Any] | None:
    matches = [
        table
        for table in sheet.get("tables", [])
        if (bounds := range_boundaries(table["range"]))[0] <= col <= bounds[2]
        and bounds[1] <= row <= bounds[3]
    ]
    return matches[0] if len(matches) == 1 else None


def source_headers(sheet: dict[str, Any], address: str) -> dict[str, list[str]]:
    row, col = coordinate_to_tuple(address)
    table = cell_region(sheet, row, col)
    return (table.get("headers") if table else None) or sheet.get("headers", {})


def read_header_paths(
    sheet: dict[str, Any], addresses: list[str]
) -> dict[str, list[str]]:
    headers: dict[str, list[str]] = {}
    seen: set[str | None] = set()
    for address in addresses:
        table = cell_region(sheet, *coordinate_to_tuple(address))
        key = table["range"] if table else None
        if key in seen:
            continue
        seen.add(key)
        for col, path in source_headers(sheet, address).items():
            headers.setdefault(col, []).extend(
                value for value in path if value not in headers.get(col, [])
            )
    return headers


def row_label_index(sheet: dict[str, Any]) -> dict[int, list[tuple[int, str]]]:
    result: dict[int, list[tuple[int, str]]] = {}
    for address, entry in sheet["cells"].items():
        if entry["kind"] == "string" and entry["value"]:
            row, col = coordinate_to_tuple(address)
            result.setdefault(row, []).append((col, canonical_label(entry["value"])))
    return result


def labels_before(
    sheet: dict[str, Any],
    row: int,
    col: int,
    index: dict[int, list[tuple[int, str]]],
) -> list[str]:
    table = cell_region(sheet, row, col)
    left = range_boundaries(table["range"])[0] if table else 1
    return [text for position, text in index.get(row, []) if left <= position < col]


def age_reference_columns(sheet: dict[str, Any], query: str) -> tuple[str, str] | None:
    """An explicitly named reference date wins over the Admission default."""
    if not re.search(r"\b(oldest|youngest)\b", query, re.IGNORECASE):
        return None
    if re.search(
        r"\b(?:as of|on)\s+\d{4}[-/]\d|\bby (?:the )?(?:birth date|date of birth)|\bbased on (?:the )?(?:birth date|date of birth)",
        query,
        re.IGNORECASE,
    ):
        return None
    date_columns = {
        get_column_letter(coordinate_to_tuple(address)[1])
        for address, entry in sheet["cells"].items()
        if entry["kind"] == "datetime"
    }
    headers = sheet.get("headers", {})
    birth = [
        col
        for col, path in headers.items()
        if col in date_columns
        and any(re.search(r"birth|dob", value, re.IGNORECASE) for value in path)
    ]
    if len(birth) != 1:
        return None
    references = {}
    for col in date_columns - {birth[0]}:
        label = " ".join(headers.get(col, []))
        label = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", label).lower()
        tokens = set(re.findall(r"[a-z]+", label)) - {
            "date",
            "of",
            "on",
            "at",
            "the",
            "time",
            "year",
        }
        if tokens:
            references[col] = tokens
    contextual = [
        col
        for col, tokens in references.items()
        if any(
            re.search(
                rf"\b(?:at|on|relative to|with respect to)\b[^.?!]{{0,45}}\b{re.escape(token)}\b",
                query,
                re.IGNORECASE,
            )
            for token in tokens
        )
    ]
    mentioned = [
        col
        for col, tokens in references.items()
        if any(
            re.search(rf"\b{re.escape(token)}\b", query, re.IGNORECASE)
            for token in tokens
        )
    ]
    choices = contextual or mentioned
    if len(choices) == 1:
        return birth[0], choices[0]
    if choices:
        return None
    admission = [col for col, tokens in references.items() if "admission" in tokens]
    return (birth[0], admission[0]) if len(admission) == 1 else None


def row_scope_errors(
    workbook: dict[str, Any], reads: dict[str, Any], query: str
) -> list[str]:
    errors = []
    ranges = re.findall(r"\b(\d+)\s*(?:to|[-–])\s*(\d+)\b", query.lower())
    for name, sheet in workbook.items():
        selected = [
            (coordinate_to_tuple(key[len(name) + 1 :]), entry)
            for key, entry in reads.items()
            if key.startswith(name + "!")
            and re.fullmatch(r"[A-Z]+[1-9]\d*", key[len(name) + 1 :])
        ]
        numeric_rows = {
            position[0] for position, entry in selected if entry["kind"] == "number"
        }
        numeric_columns = {
            get_column_letter(position[1])
            for position, entry in selected
            if entry["kind"] == "number"
        }
        data_rows = {
            coordinate_to_tuple(address)[0]
            for address, entry in sheet["cells"].items()
            if entry["kind"] == "number"
        }
        for lower, upper in ranges:
            pattern = rf"\b{lower}\s*(?:to|[-–])\s*{upper}\b"
            matching = {
                coordinate_to_tuple(address)[0]
                for address, entry in sheet["cells"].items()
                if entry["kind"] == "string"
                and coordinate_to_tuple(address)[0] in data_rows
                and re.search(pattern, str(entry["value"]), re.IGNORECASE)
            }
            if matching and numeric_rows and not numeric_rows & matching:
                errors.append(
                    f"Numeric reads omit the requested row range {lower} to {upper} in {name}; matching rows: {sorted(matching)}"
                )

        headers = sheet.get("headers", {})
        numeric_columns_by_row: dict[int, list[int]] = {}
        for address, entry in sheet["cells"].items():
            if entry["kind"] == "number":
                row, col = coordinate_to_tuple(address)
                numeric_columns_by_row.setdefault(row, []).append(col)
        leading_labels = []
        for address, entry in sheet["cells"].items():
            if entry["kind"] == "string" and entry["value"]:
                row, col = coordinate_to_tuple(address)
                if any(other > col for other in numeric_columns_by_row.get(row, [])):
                    leading_labels.append(((row, col), canonical_label(entry["value"])))
        # Both explicitly named measures must be read for a difference. Searching
        # labels alone is not evidence that their numeric operands were selected.
        if re.search(r"\bdifference between\b", query, re.IGNORECASE):
            lower = canonical_label(query).casefold()
            named: dict[str, set[int]] = {}
            for (row, _), label in leading_labels:
                text = label.casefold()
                if len(re.findall(r"\w+", text)) >= 2 and re.search(
                    rf"(?<!\w){re.escape(text)}(?!\w)", lower
                ):
                    named.setdefault(text, set()).add(row)
            maximal = {
                label: matches
                for label, matches in named.items()
                if not any(label != other and label in other for other in named)
            }
            if len(maximal) >= 2:
                for label, matches in maximal.items():
                    if not matches & numeric_rows:
                        errors.append(
                            f"Difference operands omit explicitly requested row measure {label!r}; read its numeric value in the requested section, not a neighboring metric."
                        )
        # Across ALL displayed categories is a column-coverage contract when the
        # metric is a row label. Explicit leaf/weighted qualifiers override it.
        if re.search(
            r"\b(?:average|mean)\b.*\bacross all\b", query, re.IGNORECASE
        ) and not re.search(
            r"\b(?:leaf|mutually exclusive|weighted)\b", query, re.IGNORECASE
        ):
            for row in numeric_rows:
                if any(
                    position[0] == row and label.casefold() in query.casefold()
                    for position, label in leading_labels
                ):
                    available = set(mean_scope_cells(sheet, row, query) or [])
                    observed = {
                        key[len(name) + 1 :]
                        for key, entry in reads.items()
                        if key.startswith(name + "!") and entry["kind"] == "number"
                    }
                    if available - observed:
                        errors.append(
                            f"Average across all displayed categories omitted cells {sorted(available - observed)}. Include all displayed data columns, including parent Totals, unless the question explicitly restricts to leaves/weighted averaging."
                        )
        options = re.search(
            r"rank the following [^:]+:\s*(.*?)\.\s*answer", query, re.IGNORECASE
        )
        if options:
            for option in options[1].split(","):
                text = re.sub(
                    r"^and\s+", "", option.strip(), flags=re.IGNORECASE
                ).casefold()
                matching = {
                    position[0]
                    for position, label in leading_labels
                    if label.casefold() == text
                    or label.casefold().startswith(text + " ")
                }
                if matching and not numeric_rows & matching:
                    errors.append(
                        f"Ranking reads omit the requested option {option.strip()!r}; use find_rows() to locate each label separately, not adjacent row assumptions."
                    )
        headers = sheet.get("headers", {})
        if (
            "across all industries" in query.lower()
            and len(numeric_columns) == 1
            and len(headers) >= 4
        ):
            paths = [
                " ".join(headers.get(column, [])).lower() for column in numeric_columns
            ]
            if paths and not any(
                "all industries" in path or "total" in path for path in paths
            ):
                errors.append(
                    "The selected column is a single industry, not an overall total. Aggregate disjoint industry columns when no total column exists."
                )
    return errors


def age_sort_errors(
    workbook: dict[str, Any], execution: ExecutionResult, query: str
) -> list[str]:
    """Audit elapsed-age order/index without computing a reference answer."""
    lower = query.lower()
    for name, sheet in workbook.items():
        reference = age_reference_columns(sheet, query)
        if reference is None:
            continue
        birth_col, reference_col = reference
        ages = []
        for key, entry in execution.reads.items():
            if key.startswith(name + "!") and entry["kind"] == "datetime":
                row, col = coordinate_to_tuple(key[len(name) + 1 :])
                if get_column_letter(col) != birth_col:
                    continue
                other = execution.reads.get(f"{name}!{reference_col}{row}")
                if other and other["kind"] == "datetime":
                    ages.append(
                        (
                            datetime.datetime.fromisoformat(other["value"])
                            - datetime.datetime.fromisoformat(entry["value"])
                        ).days
                    )
        words = {
            "first": 1,
            "second": 2,
            "third": 3,
            "fourth": 4,
            "fifth": 5,
            "sixth": 6,
            "one": 1,
            "two": 2,
            "three": 3,
            "four": 4,
            "five": 5,
            "six": 6,
            "seven": 7,
            "eight": 8,
            "nine": 9,
            "ten": 10,
        }
        ordinal = re.search(
            r"\b(first|second|third|fourth|fifth|sixth)[ -](oldest|youngest)\b", lower
        )
        group = re.search(
            r"among (?:the )?(\d+|one|two|three|four|five|six|seven|eight|nine|ten) (oldest|youngest)",
            lower,
        )
        count = (
            int(group[1])
            if group and group[1].isdigit()
            else words[group[1]]
            if group
            else words[ordinal[1]]
            if ordinal
            else 1
        )
        oldest = "oldest" in lower
        for record in execution.sorts:
            keys = record["keys"]
            if (
                not keys
                or not ages
                or not all(isinstance(value, (int, float)) for value in keys)
            ):
                continue
            sign = (
                1
                if sorted(keys) == sorted(ages)
                else -1
                if sorted(keys) == sorted(-age for age in ages)
                else 0
            )
            if not sign:
                continue
            descending_age = bool(record["reverse"]) == (sign == 1)
            expected = (
                list(range(count))
                if descending_age == oldest
                else list(range(len(keys) - count, len(keys)))
            )
            expected_index = (
                count - 1 if descending_age == oldest else len(keys) - count
            )
            for selection in record["selected"]:
                if (
                    isinstance(selection, int)
                    and not group
                    and selection != expected_index
                ):
                    return [
                        "Elapsed-age sort direction/index contradicts the requested rank. Oldest means LARGEST elapsed days; youngest means SMALLEST. Correct reverse and the selected index."
                    ]
                if (
                    isinstance(selection, list)
                    and group
                    and sorted(selection) != expected
                ):
                    return [
                        "Selected age slice does not contain the requested youngest/oldest group. Correct the elapsed-age sort direction and slice."
                    ]
    return []


def label_errors(
    execution: ExecutionResult, query: str, workbook: dict[str, Any] | None = None
) -> list[str]:
    requested = re.search(
        r"answer with only (?:the )?(.*(?:name|category))\b", query, re.IGNORECASE
    )
    if isinstance(execution.result, (int, float)) and requested and workbook:

        def tokens(text: str) -> set[str]:
            text = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", text)
            return {
                word.rstrip("s") for word in re.findall(r"[a-z]+", text.casefold())
            } - {"of", "the", "only", "with", "answer"}

        identity = tokens(requested[1])
        if "name" in identity:
            identity = {"name"}
        for key, entry in execution.reads.items():
            if entry["kind"] != "number" or entry["value"] != execution.result:
                continue
            name, address = key.rsplit("!", 1)
            sheet = workbook[name]
            row, column = coordinate_to_tuple(address)
            headers = source_headers(sheet, address)
            own = headers.get(get_column_letter(column), [])
            if own and identity & tokens(own[-1]):
                return []  # Numeric identifiers/categories are legitimate labels.
            for position, text in row_label_index(sheet).get(row, []):
                path = headers.get(get_column_letter(position), [])
                if position < column and path and identity & tokens(path[-1]):
                    return [
                        "The question requests a name/category, but the result is a numeric measure. Return the complete corresponding identity cell, not its score or percentage."
                    ]
    if (
        not isinstance(execution.result, str)
        or not re.search(
            r"answer with only (?:the )?.*(?:name|category)\b", query, re.IGNORECASE
        )
        or "full name" in query.lower()
    ):
        return []
    labels = {
        canonical_label(entry["value"])
        for entry in execution.reads.values()
        if entry["kind"] == "string"
    }
    if canonical_label(execution.result) not in labels:
        return [
            "Final label does not match a whole selected workbook label. Do not index a scalar string to obtain a name; use cell(), or a complete element from values()."
        ]
    if re.search(r"\bhour category\b", query, re.IGNORECASE) and re.match(
        r"^(?:Total\b|All\b)", execution.result, re.IGNORECASE
    ):
        return [
            "A Total/All row aggregates hour categories; the question asks for a specific hour category. Exclude the grand total from candidate rows."
        ]
    return []


def complete_category_contract(
    workbook: dict[str, Any], execution: ExecutionResult, query: str
) -> str | None:
    """Audit an unrestricted category extremum against its observed eligible rows."""
    if (
        not isinstance(execution.result, str)
        or not re.search(
            r"\bwhich\b.*\bcategory\b.*\b(highest|largest|lowest|smallest)\b",
            query,
            re.IGNORECASE,
        )
        or re.search(
            r"\b(?:among|only|excluding|except|leaf|leaves|mutually exclusive|before|after)\b|\d+\s*(?:to|[-–])\s*\d+",
            query.split("?")[0],
            re.IGNORECASE,
        )
    ):
        return None
    smallest = bool(re.search(r"\b(?:lowest|smallest)\b", query, re.IGNORECASE))
    candidates = []
    for name, sheet in workbook.items():
        observed = [
            (key.rsplit("!", 1)[1], entry)
            for key, entry in execution.reads.items()
            if key.startswith(name + "!") and entry["kind"] == "number"
        ]
        columns = {coordinate_to_tuple(address)[1] for address, _ in observed}
        areas = {
            region["range"]
            if (region := cell_region(sheet, *coordinate_to_tuple(address)))
            else None
            for address, _ in observed
        }
        if len(columns) != 1 or len(areas) != 1 or None in areas:
            continue
        # An unrestricted extremum cannot silently skip an intermediate category
        # within the observed table span. Only fill existing numeric cells in the
        # same table/column; explicit restricted queries were excluded above.
        column = next(iter(columns))
        low, high = (
            min(coordinate_to_tuple(address)[0] for address, _ in observed),
            max(coordinate_to_tuple(address)[0] for address, _ in observed),
        )
        region = cell_region(sheet, *coordinate_to_tuple(observed[0][0]))
        observed = [
            (address, entry)
            for address, entry in sheet["cells"].items()
            if entry["kind"] == "number"
            and (position := coordinate_to_tuple(address))[1] == column
            and low <= position[0] <= high
            and cell_region(sheet, *position) == region
        ]
        labels = row_label_index(sheet)
        label_counts: dict[str, int] = {}
        for row_labels in labels.values():
            for _, text in row_labels:
                label_counts[text] = label_counts.get(text, 0) + 1
        choices = []
        for address, entry in observed:
            row, col = coordinate_to_tuple(address)
            leading = [
                (position, text)
                for position, text in labels.get(row, [])
                if position < col
            ]
            if len(leading) != 1:
                continue
            label_col, text = leading[0]
            if re.match(r"^(?:Total|All)\b", text, re.IGNORECASE):
                continue
            # A declared section is required when labels repeat across groups.
            if label_counts[text] != 1:
                return None
            choices.append(
                (entry["value"], f"{get_column_letter(label_col)}{row}", text, address)
            )
        if len(choices) < 2:
            continue
        best = (min if smallest else max)(value for value, _, _, _ in choices)
        winners = [
            (address, text) for value, address, text, _ in choices if value == best
        ]
        if len(winners) != 1:
            continue  # Tie policy belongs to the question; never invent one.
        address, text = winners[0]
        if canonical_label(execution.result) != text:
            candidates.append(
                (name, [(measure, label) for _, label, _, measure in choices])
            )
    if len(candidates) != 1:
        return None
    name, coordinates = candidates[0]
    operation = "min" if smallest else "max"
    return (
        f"category_values=[(num({name!r}, measure), cell({name!r}, label)) "
        f"for measure, label in {coordinates!r}]\n"
        f"result={operation}(category_values,key=lambda pair:pair[0])[1]"
    )


def entity_label_errors(
    workbook: dict[str, Any], execution: ExecutionResult, query: str
) -> list[str]:
    """Named row entities must come from their identity field, not a measure header."""
    match = re.search(r"answer with only (?:the )?(\w+) name", query, re.IGNORECASE)
    if not match or not isinstance(execution.result, (str, int, float)):
        return []
    term = match[1].casefold()
    for name, sheet in workbook.items():
        selected = [
            key[len(name) + 1 :]
            for key, entry in execution.reads.items()
            if key.startswith(name + "!") and entry["kind"] != "type_probe"
        ]
        if not selected:
            continue
        headers = read_header_paths(sheet, selected)
        selected_areas = {
            table["range"]
            if (table := cell_region(sheet, *coordinate_to_tuple(address)))
            else None
            for address in selected
        }
        columns = {
            column
            for column, path in headers.items()
            if any(
                re.search(rf"\b{re.escape(term)}\b", label, re.IGNORECASE)
                and len(label.split()) <= 3
                and not re.search(r"[,/]|\band\b", label, re.IGNORECASE)
                for label in path
            )
        }
        if not columns:
            continue
        number_rows = {
            coordinate_to_tuple(address)[0]
            for address, entry in sheet["cells"].items()
            if entry["kind"] == "number"
        }
        labels = {
            canonical_label(entry["value"])
            for address, entry in sheet["cells"].items()
            if entry["kind"] == "string"
            and get_column_letter(coordinate_to_tuple(address)[1]) in columns
            and coordinate_to_tuple(address)[0] in number_rows
            and (
                not sheet.get("tables")
                or (
                    (
                        table["range"]
                        if (table := cell_region(sheet, *coordinate_to_tuple(address)))
                        else None
                    )
                    in selected_areas
                )
            )
        }
        if labels and canonical_label(execution.result) not in labels:
            return [
                "Requested entity name must come from the identity column's data rows, not a numeric measure header or table title."
            ]
    return []


def aggregation_overlap_errors(
    workbook: dict[str, Any], reads: dict[str, Any], query: str
) -> list[str]:
    """Across-all sums must not combine a published parent with its children."""
    if not re.search(
        r"\b(?:sum|total|combined)\b", query, re.IGNORECASE
    ) or not re.search(r"\bacross all\b|\ball age groups\b", query, re.IGNORECASE):
        return []
    if re.search(
        r"\b(?:average|mean|difference|change|compare|percentage|ratio|share)\b|yes or no",
        query,
        re.IGNORECASE,
    ):
        return []
    for name, sheet in workbook.items():
        coordinates = [
            coordinate_to_tuple(key[len(name) + 1 :])
            for key, entry in reads.items()
            if key.startswith(name + "!") and entry["kind"] == "number"
        ]
        number_rows = {
            coordinate_to_tuple(address)[0]
            for address, entry in sheet["cells"].items()
            if entry["kind"] == "number"
        }
        row_labels = row_label_index(sheet)
        intervals = []
        for row, col in coordinates:
            for text in labels_before(sheet, row, col, row_labels):
                interval = age_interval(text)
                if interval is not None:
                    intervals.append((row, col, interval))
        for index, (row, col, interval) in enumerate(intervals):
            for other_row, other_col, other in intervals[index + 1 :]:
                if (
                    row != other_row
                    and col == other_col
                    and cell_region(sheet, row, col)
                    == cell_region(sheet, other_row, col)
                    and max(interval[0], other[0]) <= min(interval[1], other[1])
                ):
                    return [
                        (
                            "Across-all age aggregation includes overlapping age intervals. "
                            "Use the matching published total, or a disjoint partition; "
                            "a broad age band and its narrower bands cannot both be summed."
                        )
                    ]
        boundaries = set()
        for address, entry in sheet["cells"].items():
            if entry["kind"] == "string" and entry["value"]:
                row, _ = coordinate_to_tuple(address)
                if row not in number_rows:
                    boundaries.add(row)
        for row, col in coordinates:
            if any(
                re.match(r"^(?:Total\b|All\b)", label, re.IGNORECASE)
                for label in labels_before(sheet, row, col, row_labels)
            ):
                children = any(
                    other_col == col
                    and other > row
                    and (
                        not sheet.get("tables")
                        or cell_region(sheet, row, col)
                        == cell_region(sheet, other, col)
                    )
                    and not any(
                        re.match(r"^(?:Total\b|All\b)", label, re.IGNORECASE)
                        for label in labels_before(sheet, other, col, row_labels)
                    )
                    and not any(row < boundary < other for boundary in boundaries)
                    for other, other_col in coordinates
                )
                if children:
                    return [
                        "Across-all aggregation reads both a published Total/All parent and rows within its group. Use the matching published total directly or only disjoint children; do not sum both."
                    ]
    return []


def age_interval(label: str) -> tuple[int, float] | None:
    """Parse labelled age bands only; identifiers and hour categories are not ages."""
    if not re.search(r"\byears?\b", label, re.IGNORECASE):
        return None
    bounded = re.search(
        r"\b(\d+)\s*(?:to|[-–])\s*(\d+)\s*years?\b", label, re.IGNORECASE
    )
    if bounded:
        return int(bounded[1]), float(bounded[2])
    open_ended = re.search(
        r"\b(\d+)\s*years?\s*(?:and|or)\s*(?:over|older|more)\b", label, re.IGNORECASE
    )
    return (int(open_ended[1]), float("inf")) if open_ended else None


def mean_scope_cells(sheet: dict[str, Any], row: int, query: str) -> list[str] | None:
    """Only infer a mean's columns inside one evidenced table/header branch."""
    labels = [
        (coordinate_to_tuple(address)[1], canonical_label(entry["value"]))
        for address, entry in sheet["cells"].items()
        if coordinate_to_tuple(address)[0] == row
        and entry["kind"] == "string"
        and entry["value"]
        and canonical_label(entry["value"]).casefold() in query.casefold()
    ]
    if len(labels) != 1:
        return None
    label_col = labels[0][0]
    regions = [
        table
        for table in sheet.get("tables", [])
        if (bounds := range_boundaries(table["range"]))[0] <= label_col <= bounds[2]
        and bounds[1] <= row <= bounds[3]
    ]
    if sheet.get("tables") and len(regions) != 1:
        return None
    region = regions[0] if regions else None
    bounds = range_boundaries(region["range"]) if region else None
    headers = (region.get("headers") if region else None) or sheet.get("headers", {})
    available = [
        address
        for address, entry in sheet["cells"].items()
        if entry["kind"] == "number"
        and (position := coordinate_to_tuple(address))[0] == row
        and position[1] > label_col
        and (bounds is None or bounds[0] <= position[1] <= bounds[2])
    ]
    if not available:
        return None
    paths = {
        address: headers.get(get_column_letter(coordinate_to_tuple(address)[1]), [])
        for address in available
    }
    # Unlabelled columns cannot establish that every measure belongs in the mean.
    if any(not path for path in paths.values()):
        return None
    years = set(re.findall(r"\b(?:19|20)\d{2}\b", query))
    if years:
        available = [
            address
            for address in available
            if years & set(re.findall(r"\b(?:19|20)\d{2}\b", " ".join(paths[address])))
        ]
        if not available:
            return None
    # A named header branch constrains columns; a year alone is not such a branch.
    matched = []
    for address in available:
        for label in paths[address]:
            text = re.sub(r"\(\d+\)", "", canonical_label(label)).strip()
            if (
                not re.fullmatch(r"(?:19|20)\d{2}", text)
                and len(text) >= 3
                and re.search(rf"(?<!\w){re.escape(text)}(?!\w)", query, re.IGNORECASE)
            ):
                matched.append(address)
                break
    if matched:
        available = matched
    # Mixed measures and identity/weight columns need explicit semantic selection.
    if any(
        re.search(
            r"\b(?:id|identifier|code|weight|denominator)\b|sample size",
            " ".join(paths[address]),
            re.IGNORECASE,
        )
        and not re.search(
            r"\b(?:id|identifier|code|weight|denominator)\b|sample size",
            query,
            re.IGNORECASE,
        )
        for address in available
    ):
        return None
    # Competing multi-column header branches must be resolved by the model, not
    # automatically merged just because they share a metric row.
    first_labels = {
        paths[address][0] for address in available if len(paths[address]) > 1
    }
    if len(first_labels) > 1 and not years and not matched:
        return None
    return sorted(available, key=coordinate_to_tuple)


def complete_mean_contract(
    workbook: dict[str, Any], execution: ExecutionResult, query: str
) -> str | None:
    """Materialize an unambiguous ALL-column mean on a model-selected metric row."""
    if not isinstance(execution.result, (int, float)) or isinstance(
        execution.result, bool
    ):
        return None
    if not re.search(
        r"\b(?:average|mean)\b.*\bacross all\b", query, re.IGNORECASE
    ) or re.search(
        r"\b(?:leaf|mutually exclusive|weighted|highest|lowest|rank)\b",
        query,
        re.IGNORECASE,
    ):
        return None
    import json

    candidates = []
    for name, sheet in workbook.items():
        selected_rows = {
            coordinate_to_tuple(key[len(name) + 1 :])[0]
            for key, entry in execution.reads.items()
            if key.startswith(name + "!") and entry["kind"] == "number"
        }
        for row in selected_rows:
            matching_label = any(
                entry["kind"] == "string"
                and entry["value"]
                and coordinate_to_tuple(address)[0] == row
                and canonical_label(entry["value"]).casefold() in query.casefold()
                for address, entry in sheet["cells"].items()
            )
            if not matching_label:
                continue
            available = mean_scope_cells(sheet, row, query)
            observed = {
                key[len(name) + 1 :]
                for key, entry in execution.reads.items()
                if key.startswith(name + "!") and entry["kind"] == "number"
            }
            if available and set(available) - observed:
                candidates.append((name, available))
    if len(candidates) != 1:
        return None
    name, addresses = candidates[0]
    return f"values_to_average=[num({json.dumps(name)}, address) for address in {addresses!r}]\nresult=sum(values_to_average)/len(values_to_average)"


def row_context_features(
    sheet: dict[str, Any], matches: list[int], label: str
) -> dict[int, set[str]]:
    """Index row dimensions once per repeated label, not once per numeric read."""
    features: dict[int, set[str]] = {}
    matched_rows = set(matches)
    for address, entry in sheet["cells"].items():
        position, column = coordinate_to_tuple(address)
        if position not in matched_rows:
            continue
        text = canonical_label(entry["value"]).casefold()
        if (
            entry["kind"] == "string"
            and len(text) >= 2
            and text != label.casefold()
            and text
            not in {
                "of",
                "at",
                "in",
                "on",
                "an",
                "as",
                "is",
                "to",
                "or",
                "and",
                "the",
                "for",
            }
        ):
            features.setdefault(position, set()).add(text)
        elif entry["kind"] == "number" and re.fullmatch(r"\d{4}", text):
            path = source_headers(sheet, address).get(get_column_letter(column), [])
            if any(
                re.search(r"\b(?:year|period)\b", value, re.IGNORECASE)
                for value in path
            ):
                features.setdefault(position, set()).add(text)
    return features


def row_context_resolves_label(
    features: dict[int, set[str]], row: int, matches: list[int], query: str
) -> bool:
    """Only a unique row supported by explicit question qualifiers resolves scope."""
    named = {
        text
        for text in features.get(row, set())
        if re.search(rf"(?<!\w){re.escape(text)}(?!\w)", query, re.IGNORECASE)
    }
    return bool(named) and {
        other for other in matches if named <= features.get(other, set())
    } == {row}


def scope_risks(
    workbook: dict[str, Any], execution: ExecutionResult, query: str
) -> list[str]:
    """Application-triggered review for ambiguity, rather than model confidence."""
    risks = []
    for name, sheet in workbook.items():
        selected_rows = {
            coordinate_to_tuple(key[len(name) + 1 :])[0]
            for key, entry in execution.reads.items()
            if key.startswith(name + "!") and entry["kind"] == "number"
        }
        targets: dict[str, set[int]] = {}
        for row, labels_in_row in row_label_index(sheet).items():
            for _, label in labels_in_row:
                if re.search(
                    rf"\bfor\s+(?:the\s+)?{re.escape(label)}\s+only\b",
                    query,
                    re.IGNORECASE,
                ):
                    targets.setdefault(label, set()).add(row)
        for label, matching in targets.items():
            if selected_rows and selected_rows - matching:
                risks.append(
                    f"{name}: question is explicitly restricted to {label!r}, but numeric reads include other rows; independently verify filtering and joins rather than summing all entities."
                )
        if re.search(
            r"\b(?:average|mean)\b.*\bacross all\b", query, re.IGNORECASE
        ) and any(mean_scope_cells(sheet, row, query) is None for row in selected_rows):
            risks.append(
                f"{name}: the mean's table/header branch is not uniquely evidenced; check measures, years, identities and weights."
            )
        labels: dict[str, list[int]] = {}
        context_rows: list[tuple[int, str]] = []
        by_row: dict[int, list[dict[str, Any]]] = {}
        for address, entry in sheet["cells"].items():
            row, _ = coordinate_to_tuple(address)
            by_row.setdefault(row, []).append(entry)
            if entry["kind"] == "string" and entry["value"]:
                label = canonical_label(entry["value"])
                if re.search(rf"(?<!\w){re.escape(label)}(?!\w)", query, re.IGNORECASE):
                    labels.setdefault(label, []).append(row)
        for row, entries in by_row.items():
            text = [
                canonical_label(e["value"])
                for e in entries
                if e["kind"] == "string" and e["value"]
            ]
            if len(text) == 1 and not any(
                e["kind"] in {"number", "datetime"} for e in entries
            ):
                context_rows.append((row, text[0]))
        for label, matches in labels.items():
            if len(set(matches)) <= 1 or not selected_rows.intersection(matches):
                continue
            justified = True
            features = row_context_features(sheet, matches, label)
            specific_context = any(
                not re.fullmatch(r"(?:Total|All)", text, re.IGNORECASE)
                and re.search(rf"(?<!\w){re.escape(text)}(?!\w)", query, re.IGNORECASE)
                for _, text in context_rows
            )
            for row in selected_rows.intersection(matches):
                if row_context_resolves_label(features, row, matches, query):
                    continue
                prior = [text for before, text in sorted(context_rows) if before < row]
                context = prior[-1] if prior else ""
                if (
                    re.fullmatch(r"(?:Total|All)", context, re.IGNORECASE)
                    and not specific_context
                ):
                    # An explicit overall section is evidence for an unqualified
                    # question; repeated labels in subgroups alone need no extra call.
                    continue
                words = re.findall(r"\w+", context)
                prefix = " ".join(words[:2]) if len(words) >= 2 else context
                if not prefix or not re.search(
                    rf"(?<!\w){re.escape(prefix)}(?!\w)", query, re.IGNORECASE
                ):
                    justified = False
                    break
            if not justified:
                risks.append(
                    f"{name}: label {label!r} occurs in several sections; independently verify table, group and year."
                )
                break
    return risks


def evidence_errors(
    workbook: dict[str, Any], reads: dict[str, Any], selections: list[dict[str, Any]]
) -> list[str]:
    """Validate claimed source regions against observed reads, not reference answers."""
    if not selections:
        return []
    areas: dict[str, list[tuple[int, int, int, int]]] = {}
    for selection in selections:
        name, area = selection["sheet"], selection["range"]
        if name not in workbook:
            return [f"Evidence sheet {name!r} does not exist"]
        if not re.fullmatch(r"[A-Z]+[1-9]\d*(?::[A-Z]+[1-9]\d*)?", area):
            return ["Evidence needs an explicit Excel cell/range"]
        bounds = range_boundaries(area)
        if bounds[0] > bounds[2] or bounds[1] > bounds[3]:
            return ["Evidence range is reversed"]
        areas.setdefault(name, []).append(bounds)
    for key, entry in reads.items():
        if entry["kind"] in {"type_probe", "lookup_probe"}:
            continue
        # Sheet names can contain !; the final delimiter precedes the address.
        name, address = key.rsplit("!", 1)
        row, col = coordinate_to_tuple(address)
        if not any(
            left <= col <= right and top <= row <= bottom
            for left, top, right, bottom in areas.get(name, [])
        ):
            return [
                f"Read {key} lies outside the declared evidence; include all input/header regions"
            ]
    return []


def complete_total_contract(
    workbook: dict[str, Any], execution: ExecutionResult, query: str
) -> str | None:
    """Use the sole selected published parent for an across-all total contract."""
    if re.search(
        r"\b(?:leaf|leaves)\b|\b(?:exclude|excluding|ignore)\b[^.?!]{0,40}\b(?:total|parent)s?\b",
        query,
        re.IGNORECASE,
    ):
        return None
    overlap = aggregation_overlap_errors(workbook, execution.reads, query)
    if not overlap or re.search(r"\b\d+\s*(?:to|[-–])\s*\d+\b", query):
        return None
    candidates = []
    for name, sheet in workbook.items():
        row_labels = row_label_index(sheet)
        selected = [
            key[len(name) + 1 :]
            for key, entry in execution.reads.items()
            if key.startswith(name + "!") and entry["kind"] == "number"
        ]
        # An omitted parent is recoverable only inside the same evidenced table
        # and single measure column, for an explicit all-age aggregation.
        addresses = selected
        if re.search(r"\ball age groups\b|\bacross all ages\b", query, re.IGNORECASE):
            positions = [coordinate_to_tuple(address) for address in selected]
            regions = {
                region["range"] if (region := cell_region(sheet, row, col)) else None
                for row, col in positions
            }
            columns = {col for _, col in positions}
            if len(columns) == 1 and len(regions) == 1 and None not in regions:
                column = next(iter(columns))
                region = cell_region(sheet, *positions[0])
                addresses = [
                    address
                    for address, entry in sheet["cells"].items()
                    if entry["kind"] == "number"
                    and coordinate_to_tuple(address)[1] == column
                    and cell_region(sheet, *coordinate_to_tuple(address)) == region
                ]
        for address in addresses:
            row, col = coordinate_to_tuple(address)
            if any(
                re.match(
                    r"^(?:Total\b|All\b)",
                    text,
                    re.IGNORECASE,
                )
                for text in labels_before(sheet, row, col, row_labels)
            ):
                candidates.append((name, address))
    if len(candidates) != 1:
        return None
    name, address = candidates[0]
    return f"result=num({name!r}, {address!r})"
