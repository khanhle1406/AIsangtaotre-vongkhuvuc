"""Observable scope/ranking invariants, independent of reference answers."""

import datetime
import re
from typing import Any

from openpyxl.utils import coordinate_to_tuple, get_column_letter

from .executor import ExecutionResult


def canonical_label(value: Any) -> str:
    return " ".join(re.sub(r"(?<=[A-Za-z])-\s*\n\s*(?=[a-z])", "", str(value)).split())


def age_reference_columns(sheet: dict[str, Any], query: str) -> tuple[str, str] | None:
    """An explicitly named reference date wins over the Admission default."""
    if not re.search(r"\b(oldest|youngest)\b", query, re.IGNORECASE):
        return None
    if re.search(
        r"\b(?:as of|on)\s+\d{4}[-/]\d|\bby (?:the )?(?:birth date|date of birth)|\bbased on (?:the )?(?:birth date|date of birth)",
        query,
        re.IGNORECASE,
    ):
        return None
    date_columns = {
        get_column_letter(coordinate_to_tuple(address)[1])
        for address, entry in sheet["cells"].items()
        if entry["kind"] == "datetime"
    }
    headers = sheet.get("headers", {})
    birth = [
        col
        for col, path in headers.items()
        if col in date_columns
        and any(re.search(r"birth|dob", value, re.IGNORECASE) for value in path)
    ]
    if len(birth) != 1:
        return None
    references = {}
    for col in date_columns - {birth[0]}:
        label = " ".join(headers.get(col, []))
        label = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", label).lower()
        tokens = set(re.findall(r"[a-z]+", label)) - {
            "date",
            "of",
            "on",
            "at",
            "the",
            "time",
            "year",
        }
        if tokens:
            references[col] = tokens
    contextual = [
        col
        for col, tokens in references.items()
        if any(
            re.search(
                rf"\b(?:at|on|relative to|with respect to)\b[^.?!]{{0,45}}\b{re.escape(token)}\b",
                query,
                re.IGNORECASE,
            )
            for token in tokens
        )
    ]
    mentioned = [
        col
        for col, tokens in references.items()
        if any(
            re.search(rf"\b{re.escape(token)}\b", query, re.IGNORECASE)
            for token in tokens
        )
    ]
    choices = contextual or mentioned
    if len(choices) == 1:
        return birth[0], choices[0]
    if choices:
        return None
    admission = [col for col, tokens in references.items() if "admission" in tokens]
    return (birth[0], admission[0]) if len(admission) == 1 else None


def row_scope_errors(
    workbook: dict[str, Any], reads: dict[str, Any], query: str
) -> list[str]:
    errors = []
    ranges = re.findall(r"\b(\d+)\s*(?:to|[-–])\s*(\d+)\b", query.lower())
    for name, sheet in workbook.items():
        selected = [
            (coordinate_to_tuple(key[len(name) + 1 :]), entry)
            for key, entry in reads.items()
            if key.startswith(name + "!")
            and re.fullmatch(r"[A-Z]+[1-9]\d*", key[len(name) + 1 :])
        ]
        numeric_rows = {
            position[0] for position, entry in selected if entry["kind"] == "number"
        }
        numeric_columns = {
            get_column_letter(position[1])
            for position, entry in selected
            if entry["kind"] == "number"
        }
        data_rows = {
            coordinate_to_tuple(address)[0]
            for address, entry in sheet["cells"].items()
            if entry["kind"] == "number"
        }
        for lower, upper in ranges:
            pattern = rf"\b{lower}\s*(?:to|[-–])\s*{upper}\b"
            matching = {
                coordinate_to_tuple(address)[0]
                for address, entry in sheet["cells"].items()
                if entry["kind"] == "string"
                and coordinate_to_tuple(address)[0] in data_rows
                and re.search(pattern, str(entry["value"]), re.IGNORECASE)
            }
            if matching and numeric_rows and not numeric_rows & matching:
                errors.append(
                    f"Numeric reads omit the requested row range {lower} to {upper} in {name}; matching rows: {sorted(matching)}"
                )

        headers = sheet.get("headers", {})
        numeric_columns_by_row: dict[int, list[int]] = {}
        for address, entry in sheet["cells"].items():
            if entry["kind"] == "number":
                row, col = coordinate_to_tuple(address)
                numeric_columns_by_row.setdefault(row, []).append(col)
        leading_labels = []
        for address, entry in sheet["cells"].items():
            if entry["kind"] == "string" and entry["value"]:
                row, col = coordinate_to_tuple(address)
                if any(other > col for other in numeric_columns_by_row.get(row, [])):
                    leading_labels.append(((row, col), canonical_label(entry["value"])))
        # Across ALL displayed categories is a column-coverage contract when the
        # metric is a row label. Explicit leaf/weighted qualifiers override it.
        if re.search(
            r"\b(?:average|mean)\b.*\bacross all\b", query, re.IGNORECASE
        ) and not re.search(
            r"\b(?:leaf|mutually exclusive|weighted)\b", query, re.IGNORECASE
        ):
            years = set(re.findall(r"\b(?:19|20)\d{2}\b", query))
            for row in numeric_rows:
                if any(
                    position[0] == row and label.casefold() in query.casefold()
                    for position, label in leading_labels
                ):
                    available = {
                        address
                        for address, entry in sheet["cells"].items()
                        if entry["kind"] == "number"
                        and coordinate_to_tuple(address)[0] == row
                        and (
                            not years
                            or years
                            & set(
                                headers.get(
                                    get_column_letter(coordinate_to_tuple(address)[1]),
                                    [],
                                )
                            )
                        )
                    }
                    observed = {
                        key[len(name) + 1 :]
                        for key, entry in reads.items()
                        if key.startswith(name + "!") and entry["kind"] == "number"
                    }
                    if available - observed:
                        errors.append(
                            f"Average across all displayed categories omitted cells {sorted(available - observed)}. Include all displayed data columns, including parent Totals, unless the question explicitly restricts to leaves/weighted averaging."
                        )
        options = re.search(
            r"rank the following [^:]+:\s*(.*?)\.\s*answer", query, re.IGNORECASE
        )
        if options:
            for option in options[1].split(","):
                text = re.sub(
                    r"^and\s+", "", option.strip(), flags=re.IGNORECASE
                ).casefold()
                matching = {
                    position[0]
                    for position, label in leading_labels
                    if label.casefold() == text
                    or label.casefold().startswith(text + " ")
                }
                if matching and not numeric_rows & matching:
                    errors.append(
                        f"Ranking reads omit the requested option {option.strip()!r}; use find_rows() to locate each label separately, not adjacent row assumptions."
                    )
        headers = sheet.get("headers", {})
        if (
            "across all industries" in query.lower()
            and len(numeric_columns) == 1
            and len(headers) >= 4
        ):
            paths = [
                " ".join(headers.get(column, [])).lower() for column in numeric_columns
            ]
            if paths and not any(
                "all industries" in path or "total" in path for path in paths
            ):
                errors.append(
                    "The selected column is a single industry, not an overall total. Aggregate disjoint industry columns when no total column exists."
                )
    return errors


def age_sort_errors(
    workbook: dict[str, Any], execution: ExecutionResult, query: str
) -> list[str]:
    """Audit elapsed-age order/index without computing a reference answer."""
    lower = query.lower()
    for name, sheet in workbook.items():
        reference = age_reference_columns(sheet, query)
        if reference is None:
            continue
        birth_col, reference_col = reference
        ages = []
        for key, entry in execution.reads.items():
            if key.startswith(name + "!") and entry["kind"] == "datetime":
                row, col = coordinate_to_tuple(key[len(name) + 1 :])
                if get_column_letter(col) != birth_col:
                    continue
                other = execution.reads.get(f"{name}!{reference_col}{row}")
                if other and other["kind"] == "datetime":
                    ages.append(
                        (
                            datetime.datetime.fromisoformat(other["value"])
                            - datetime.datetime.fromisoformat(entry["value"])
                        ).days
                    )
        words = {
            "first": 1,
            "second": 2,
            "third": 3,
            "fourth": 4,
            "fifth": 5,
            "sixth": 6,
            "one": 1,
            "two": 2,
            "three": 3,
            "four": 4,
            "five": 5,
            "six": 6,
            "seven": 7,
            "eight": 8,
            "nine": 9,
            "ten": 10,
        }
        ordinal = re.search(
            r"\b(first|second|third|fourth|fifth|sixth)[ -](oldest|youngest)\b", lower
        )
        group = re.search(
            r"among (?:the )?(\d+|one|two|three|four|five|six|seven|eight|nine|ten) (oldest|youngest)",
            lower,
        )
        count = (
            int(group[1])
            if group and group[1].isdigit()
            else words[group[1]]
            if group
            else words[ordinal[1]]
            if ordinal
            else 1
        )
        oldest = "oldest" in lower
        for record in execution.sorts:
            keys = record["keys"]
            if (
                not keys
                or not ages
                or not all(isinstance(value, (int, float)) for value in keys)
            ):
                continue
            sign = (
                1
                if sorted(keys) == sorted(ages)
                else -1
                if sorted(keys) == sorted(-age for age in ages)
                else 0
            )
            if not sign:
                continue
            descending_age = bool(record["reverse"]) == (sign == 1)
            expected = (
                list(range(count))
                if descending_age == oldest
                else list(range(len(keys) - count, len(keys)))
            )
            expected_index = (
                count - 1 if descending_age == oldest else len(keys) - count
            )
            for selection in record["selected"]:
                if (
                    isinstance(selection, int)
                    and not group
                    and selection != expected_index
                ):
                    return [
                        "Elapsed-age sort direction/index contradicts the requested rank. Oldest means LARGEST elapsed days; youngest means SMALLEST. Correct reverse and the selected index."
                    ]
                if (
                    isinstance(selection, list)
                    and group
                    and sorted(selection) != expected
                ):
                    return [
                        "Selected age slice does not contain the requested youngest/oldest group. Correct the elapsed-age sort direction and slice."
                    ]
    return []


def label_errors(execution: ExecutionResult, query: str) -> list[str]:
    if (
        not isinstance(execution.result, str)
        or not re.search(
            r"answer with only (?:the )?.*(?:name|category)\b", query, re.IGNORECASE
        )
        or "full name" in query.lower()
    ):
        return []
    labels = {
        canonical_label(entry["value"])
        for entry in execution.reads.values()
        if entry["kind"] == "string"
    }
    if canonical_label(execution.result) not in labels:
        return [
            "Final label does not match a whole selected workbook label. Do not index a scalar string to obtain a name; use cell(), or a complete element from values()."
        ]
    if re.search(r"\bhour category\b", query, re.IGNORECASE) and re.match(
        r"^(?:Total\b|All\b)", execution.result, re.IGNORECASE
    ):
        return [
            "A Total/All row aggregates hour categories; the question asks for a specific hour category. Exclude the grand total from candidate rows."
        ]
    return []


def entity_label_errors(
    workbook: dict[str, Any], execution: ExecutionResult, query: str
) -> list[str]:
    """Named row entities must come from their identity field, not a measure header."""
    match = re.search(r"answer with only (?:the )?(\w+) name", query, re.IGNORECASE)
    if not match or not isinstance(execution.result, (str, int, float)):
        return []
    term = match[1].casefold()
    for sheet in workbook.values():
        columns = {
            column
            for column, path in sheet.get("headers", {}).items()
            if any(
                re.search(rf"\b{re.escape(term)}\b", label, re.IGNORECASE)
                and len(label.split()) <= 3
                and not re.search(r"[,/]|\band\b", label, re.IGNORECASE)
                for label in path
            )
        }
        if not columns:
            continue
        number_rows = {
            coordinate_to_tuple(address)[0]
            for address, entry in sheet["cells"].items()
            if entry["kind"] == "number"
        }
        labels = {
            canonical_label(entry["value"])
            for address, entry in sheet["cells"].items()
            if entry["kind"] == "string"
            and get_column_letter(coordinate_to_tuple(address)[1]) in columns
            and coordinate_to_tuple(address)[0] in number_rows
        }
        if labels and canonical_label(execution.result) not in labels:
            return [
                "Requested entity name must come from the identity column's data rows, not a numeric measure header or table title."
            ]
    return []


def aggregation_overlap_errors(
    workbook: dict[str, Any], reads: dict[str, Any], query: str
) -> list[str]:
    """Across-all sums must not combine a published parent with its children."""
    if not re.search(
        r"\b(?:sum|total|combined)\b", query, re.IGNORECASE
    ) or not re.search(r"\bacross all\b|\ball age groups\b", query, re.IGNORECASE):
        return []
    if re.search(
        r"\b(?:average|mean|difference|change|compare)\b|yes or no",
        query,
        re.IGNORECASE,
    ):
        return []
    for name, sheet in workbook.items():
        coordinates = [
            coordinate_to_tuple(key[len(name) + 1 :])
            for key, entry in reads.items()
            if key.startswith(name + "!") and entry["kind"] == "number"
        ]
        number_rows = {
            coordinate_to_tuple(address)[0]
            for address, entry in sheet["cells"].items()
            if entry["kind"] == "number"
        }
        row_labels: dict[int, list[str]] = {}
        boundaries = set()
        for address, entry in sheet["cells"].items():
            if entry["kind"] == "string" and entry["value"]:
                row, _ = coordinate_to_tuple(address)
                row_labels.setdefault(row, []).append(canonical_label(entry["value"]))
                if row not in number_rows:
                    boundaries.add(row)
        for row, col in coordinates:
            if any(
                re.match(r"^(?:Total\b|All\b)", label, re.IGNORECASE)
                for label in row_labels.get(row, [])
            ):
                children = [
                    other
                    for other, other_col in coordinates
                    if other_col == col
                    and other > row
                    and not any(
                        re.match(r"^(?:Total\b|All\b)", label, re.IGNORECASE)
                        for label in row_labels.get(other, [])
                    )
                    and not any(row < boundary < other for boundary in boundaries)
                ]
                if children:
                    return [
                        "Across-all aggregation reads both a published Total/All parent and rows within its group. Use the matching published total directly or only disjoint children; do not sum both."
                    ]
    return []


def complete_mean_contract(
    workbook: dict[str, Any], execution: ExecutionResult, query: str
) -> str | None:
    """Materialize an unambiguous ALL-column mean on a model-selected metric row."""
    if not isinstance(execution.result, (int, float)) or isinstance(
        execution.result, bool
    ):
        return None
    if not re.search(
        r"\b(?:average|mean)\b.*\bacross all\b", query, re.IGNORECASE
    ) or re.search(
        r"\b(?:leaf|mutually exclusive|weighted|highest|lowest|rank)\b",
        query,
        re.IGNORECASE,
    ):
        return None
    import json

    candidates = []
    years = set(re.findall(r"\b(?:19|20)\d{2}\b", query))
    for name, sheet in workbook.items():
        selected_rows = {
            coordinate_to_tuple(key[len(name) + 1 :])[0]
            for key, entry in execution.reads.items()
            if key.startswith(name + "!") and entry["kind"] == "number"
        }
        for row in selected_rows:
            matching_label = any(
                entry["kind"] == "string"
                and entry["value"]
                and coordinate_to_tuple(address)[0] == row
                and canonical_label(entry["value"]).casefold() in query.casefold()
                for address, entry in sheet["cells"].items()
            )
            if not matching_label:
                continue
            available = [
                address
                for address, entry in sheet["cells"].items()
                if entry["kind"] == "number"
                and coordinate_to_tuple(address)[0] == row
                and (
                    not years
                    or years
                    & set(
                        sheet.get("headers", {}).get(
                            get_column_letter(coordinate_to_tuple(address)[1]), []
                        )
                    )
                )
            ]
            observed = {
                key[len(name) + 1 :]
                for key, entry in execution.reads.items()
                if key.startswith(name + "!") and entry["kind"] == "number"
            }
            if available and set(available) - observed:
                candidates.append((name, available))
    if len(candidates) != 1:
        return None
    name, addresses = candidates[0]
    return f"values_to_average=[num({json.dumps(name)}, address) for address in {addresses!r}]\nresult=sum(values_to_average)/len(values_to_average)"


def complete_total_contract(
    workbook: dict[str, Any], execution: ExecutionResult, query: str
) -> str | None:
    """Use the sole selected published parent for an across-all total contract."""
    if not aggregation_overlap_errors(workbook, execution.reads, query) or re.search(
        r"\b\d+\s*(?:to|[-–])\s*\d+\b", query
    ):
        return None
    candidates = []
    for name, sheet in workbook.items():
        for key, entry in execution.reads.items():
            if not key.startswith(name + "!") or entry["kind"] != "number":
                continue
            address = key[len(name) + 1 :]
            row, _ = coordinate_to_tuple(address)
            if any(
                value["kind"] == "string"
                and value["value"]
                and coordinate_to_tuple(coord)[0] == row
                and re.match(
                    r"^(?:Total\b|All\b)",
                    canonical_label(value["value"]),
                    re.IGNORECASE,
                )
                for coord, value in sheet["cells"].items()
            ):
                candidates.append((name, address))
    if len(candidates) != 1:
        return None
    name, address = candidates[0]
    return f"result=num({name!r}, {address!r})"
