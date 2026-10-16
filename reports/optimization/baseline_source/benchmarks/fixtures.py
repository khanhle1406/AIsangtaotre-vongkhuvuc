"""Independent oracle workbooks: no competition answers or solver imports."""

import csv
import datetime as dt
import json
import random
from pathlib import Path
from typing import Any

import openpyxl


def generate_fixtures(root: Path, seed: int = 8117) -> tuple[Path, Path]:
    rng = random.Random(seed)
    files = root / "xlsx_files"
    files.mkdir(parents=True, exist_ok=True)
    queries: list[dict[str, Any]] = []
    references: list[dict[str, str]] = []

    def add(file: str, question: str, answer: str) -> None:
        key = f"generated-{len(queries) + 1:03d}"
        queries.append({"id": key, "file_name": file, "query": question})
        references.append({"id": key, "answer": answer})

    # Unfamiliar names, reordered measure columns, two sheets, distractor table, long tail.
    for variant in range(2):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Invoices {rng.randrange(100, 999)}"
        columns = [2, 4, 6]
        rng.shuffle(columns)
        name_col, sold_col, returned_col = columns
        header = 3 if variant == 0 else 307
        ws.cell(header, 1, "Units, current period")
        for col, label in [
            (name_col, "Company"),
            (sold_col, "Units sold"),
            (returned_col, "Units returned"),
        ]:
            ws.cell(header + 1, col, label)
        names = ["Delta Co.", "Orion Ltd.", "Zenith Inc."]
        sold = [rng.randrange(80, 200) for _ in names]
        returned = [rng.randrange(2, 20) for _ in names]
        for row, name, sale, ret in zip(
            range(header + 2, header + 5), names, sold, returned, strict=True
        ):
            for col, value in [(name_col, name), (sold_col, sale), (returned_col, ret)]:
                ws.cell(row, col, value)
        distractor = wb.create_sheet("Previous period")
        distractor.append(["Company", "Units sold"])
        distractor.append(["Delta Co.", 9999])
        ws.cell(1, 10, "Unrelated note: counts below are current period only.")
        file = f"renamed_invoice_{variant}.xlsx"
        wb.save(files / file)
        wb.close()
        add(
            file,
            "In the current period, how many units did Delta Co. sell? Answer with only the integer.",
            str(sold[0]),
        )
        add(
            file,
            "In the current period, which company sold the most units? Answer with only the company name.",
            names[sold.index(max(sold))],
        )
        add(
            file,
            "What is the combined current-period units sold for Orion Ltd. and Zenith Inc.? Answer with only the integer.",
            str(sold[1] + sold[2]),
        )
        add(
            file,
            "For Delta Co., calculate current-period units sold minus units returned. Answer with only the integer.",
            str(sold[0] - returned[0]),
        )
        add(
            file,
            "Did Delta Co. sell more units than Orion Ltd. in the current period? Answer with only Yes or No.",
            "Yes" if sold[0] > sold[1] else "No",
        )

    # Store true Excel dates/fractions; event-relative ranking differs from birth order.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registrations"
    ws.append(["Last name", "Birth date", "Admission date", "Score"])
    people = [
        (
            "Vela",
            dt.datetime(1990, 1, 1),
            dt.datetime(2010, 1, 1),
            rng.randrange(40, 99),
        ),
        (
            "Nova",
            dt.datetime(1992, 1, 1),
            dt.datetime(2025, 1, 1),
            rng.randrange(40, 99),
        ),
        (
            "Mira",
            dt.datetime(1994, 1, 1),
            dt.datetime(2024, 1, 1),
            rng.randrange(40, 99),
        ),
        (
            "Lyra",
            dt.datetime(2000, 1, 1),
            dt.datetime(2021, 1, 1),
            rng.randrange(40, 99),
        ),
    ]
    for row in people:
        ws.append(row)
    wb.save(files / "registrations.xlsx")
    wb.close()
    ranked = sorted(people, key=lambda row: (row[2] - row[1]).days, reverse=True)
    add(
        "registrations.xlsx",
        "Who was the oldest at their admission date? Answer with only the last name.",
        ranked[0][0],
    )
    add(
        "registrations.xlsx",
        "What score did the second-oldest person on their admission date receive? Answer with only the integer.",
        str(ranked[1][3]),
    )
    add(
        "registrations.xlsx",
        "Among the two youngest on their admission dates, what was the highest score? Answer with only the integer.",
        str(max(row[3] for row in ranked[-2:])),
    )
    add(
        "registrations.xlsx",
        "What is the admission date for Mira? Answer with only the ISO date.",
        "2024-01-01",
    )

    # Nested headers and rounding-discrepant parent totals; percent vs percent points.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Survey"
    ws["A1"] = "Counts and rates"
    ws["B1"] = "2025"
    ws.merge_cells("B1:E1")
    ws["A2"] = "Group"
    ws["B2"] = "Published workers"
    ws["C2"] = "Completion rate"
    ws["D2"] = "Change, percentage points"
    ws["E2"] = "Balance"
    ws.append(["All departments", 101, 0.625, 2.5, -30])
    ws.append(["East", 40, 0.5, 1.5, -10])
    ws.append(["West", 60, 0.75, 3.5, -20])
    for address in ["C3", "C4", "C5"]:
        ws[address].number_format = "0.0%"
    ws["A8"] = "Audited adjustment"
    ws["B8"] = "=B3-B4-B5"
    wb.save(files / "survey.xlsx")
    wb.close()
    add(
        "survey.xlsx",
        "How many workers were there across all departments in 2025? Use the published total. Answer with only the integer.",
        "101",
    )
    add(
        "survey.xlsx",
        "What was the completion rate for All departments in percent? Answer as a number rounded to two decimal places, without the percent sign.",
        "62.50",
    )
    add(
        "survey.xlsx",
        "What was the change in percentage points for East? Answer rounded to two decimal places, without units.",
        "1.50",
    )
    add(
        "survey.xlsx",
        "What is the sum of the balances for East and West? Answer with only the integer.",
        "-30",
    )
    add(
        "survey.xlsx",
        "Calculate the audited adjustment in the workbook. Answer with only the integer.",
        "1",
    )
    add(
        "survey.xlsx",
        "Return the published workers for East and West as two integers, in that order, separated by a comma and one space.",
        "40, 60",
    )

    # A new structural family: two tables on one sheet with repeated labels,
    # multi-year header branches and fractional rates.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Annual and pilot"
    ws["A1"] = "Annual survey (employees are persons)"
    ws["B2"], ws["D2"] = "2024", "2025"
    ws.merge_cells("B2:C2")
    ws.merge_cells("D2:E2")
    ws.append(
        ["Region", "Employees", "Completion rate", "Employees", "Completion rate"]
    )
    east, west = rng.randrange(100, 500), rng.randrange(100, 500)
    ws.append(["East", 80, 0.25, east, 0.6])
    ws.append(["West", 90, 0.35, west, 0.8])
    for address in ["C4", "C5", "E4", "E5"]:
        ws[address].number_format = "0.0%"
    ws["A20"] = "Pilot survey (separate population)"
    ws["A21"], ws["B21"] = "East", 99999
    wb.save(files / "two_tables.xlsx")
    wb.close()
    add(
        "two_tables.xlsx",
        "In the Annual survey, how many employees were in East in 2025? Answer with only the integer.",
        str(east),
    )
    add(
        "two_tables.xlsx",
        "In the Annual survey, what is the total employees for East and West in 2025? Answer with only the integer.",
        str(east + west),
    )
    add(
        "two_tables.xlsx",
        "In the Annual survey, what is the East completion rate for 2025 in percent? Answer rounded to two decimal places, without units.",
        "60.00",
    )
    add(
        "two_tables.xlsx",
        "In the Annual survey, which region had more employees in 2025? Answer with only the region name.",
        "East" if east > west else "West",
    )

    queries_path = root / "queries.json"
    answers_path = root / "answers.csv"
    queries_path.write_text(
        json.dumps(queries, ensure_ascii=False, indent=2), encoding="utf8"
    )
    with answers_path.open("w", encoding="utf8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["id", "answer"])
        writer.writeheader()
        writer.writerows(references)
    return queries_path, answers_path


if __name__ == "__main__":
    generate_fixtures(Path("reports/holdout"))
