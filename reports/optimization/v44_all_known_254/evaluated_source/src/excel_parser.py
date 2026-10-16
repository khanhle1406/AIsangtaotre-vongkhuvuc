"""Lossless, sparse Excel input; sheet XML is streamed instead of densified."""

import datetime as dt
import json
import posixpath
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, from_excel
from openpyxl.worksheet.worksheet import Worksheet

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


@dataclass
class CellData:
    value: Any
    kind: str
    number_format: str = "General"
    formula: str | None = None
    indent: int = 0
    bold: bool = False

    def wire(self) -> dict[str, Any]:
        value = self.value
        if isinstance(value, (dt.datetime, dt.date, dt.time)):
            value = value.isoformat()
        return {
            "value": value,
            "kind": self.kind,
            "number_format": self.number_format,
            "formula": self.formula,
            "indent": self.indent,
            "bold": self.bold,
        }


@dataclass
class SheetData:
    name: str
    cells: dict[str, CellData] = field(default_factory=dict)
    merges: list[str] = field(default_factory=list)
    hidden_rows: list[int] = field(default_factory=list)
    hidden_columns: list[str] = field(default_factory=list)
    state: str = "visible"

    def wire(self) -> dict[str, Any]:
        return {
            "cells": {address: cell.wire() for address, cell in self.cells.items()},
            "merges": self.merges,
            "hidden_rows": self.hidden_rows,
            "hidden_columns": self.hidden_columns,
            "state": self.state,
        }


@dataclass
class WorkbookData:
    sheets: dict[str, SheetData]
    estimated_bytes: int = 0

    def wire(self) -> dict[str, Any]:
        return {
            name: {
                **sheet.wire(),
                "headers": self.header_paths(name),
                "tables": self.table_regions(name),
            }
            for name, sheet in self.sheets.items()
        }

    def table_regions(self, name: str) -> list[dict[str, Any]]:
        """Advisory populated bands; original coordinates remain authoritative."""
        sheet = self.sheets[name]
        populated: dict[int, set[int]] = {}
        entries_by_row: dict[int, list[tuple[str, CellData]]] = {}
        for address, entry in sheet.cells.items():
            row, col = coordinate_to_tuple(address)
            populated.setdefault(row, set()).add(col)
            entries_by_row.setdefault(row, []).append((address, entry))
        bands: list[list[int]] = []
        for row in sorted(populated):
            if not bands or row - bands[-1][-1] > 1:
                bands.append([])
            bands[-1].append(row)
        regions = []
        if len(bands) > 128:
            # Hints are optional. Never create quadratic metadata for sparse data.
            return []
        for band in bands:
            entries = [entry for row in band for entry in entries_by_row[row]]
            mixed = []
            for row in band:
                text_columns = [
                    coordinate_to_tuple(address)[1]
                    for address, entry in entries_by_row[row]
                    if entry.kind == "string" and entry.value
                ]
                number_columns = [
                    coordinate_to_tuple(address)[1]
                    for address, entry in entries_by_row[row]
                    if entry.kind in {"number", "datetime", "uncached_formula"}
                ]
                if text_columns and number_columns:
                    mixed.append(
                        (
                            max(min(text_columns), min(number_columns)),
                            min(max(text_columns), max(number_columns)),
                        )
                    )
            left_threshold = min((left for left, _ in mixed), default=sys.maxsize)
            right_threshold = max((right for _, right in mixed), default=0)
            columns = sorted(set().union(*(populated[row] for row in band)))
            groups: list[list[int]] = []
            for col in columns:
                # Split a blank column only when BOTH sides contain labelled data.
                # This preserves a single table with spaced-out measure columns.
                if not groups or (
                    col - groups[-1][-1] > 1
                    and groups[-1][-1] >= left_threshold
                    and col <= right_threshold
                ):
                    groups.append([])
                groups[-1].append(col)
            for group in groups:
                left, right = group[0], group[-1]
                top, bottom = band[0], band[-1]
                numeric = [
                    coordinate_to_tuple(address)
                    for address, entry in entries
                    if entry.kind in {"number", "datetime", "uncached_formula"}
                    and top <= coordinate_to_tuple(address)[0] <= bottom
                    and left <= coordinate_to_tuple(address)[1] <= right
                ]
                if not numeric:
                    continue
                local = SheetData(
                    name,
                    cells={
                        address: entry
                        for address, entry in entries
                        if left <= coordinate_to_tuple(address)[1] <= right
                    },
                    merges=sheet.merges,
                )
                regions.append(
                    {
                        "range": f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}",
                        "headers": WorkbookData({name: local}).header_paths(name),
                    }
                )
        return regions

    def header_paths(
        self, name: str, bounds: tuple[int, int, int, int] | None = None
    ) -> dict[str, list[str]]:
        """Conservative header hint; original cells remain the source of truth."""
        sheet = self.sheets[name]
        rows: dict[int, list[tuple[str, CellData]]] = {}
        for address, cell in sheet.cells.items():
            row, col = coordinate_to_tuple(address)
            if bounds and not (
                bounds[0] <= col <= bounds[2] and bounds[1] <= row <= bounds[3]
            ):
                continue
            rows.setdefault(row, []).append((address, cell))
        first_data = None
        captioned_columns: set[int] = set()
        for row, entries in sorted(rows.items()):
            numbers = [cell.value for _, cell in entries if cell.kind == "number"]
            # ISO text remains a string in the wire, but can mark the first data
            # row for advisory headers on an otherwise entirely textual table.
            dates = [
                cell
                for _, cell in entries
                if cell.kind == "datetime"
                or cell.kind == "string"
                and re.fullmatch(
                    r"\d{4}-\d{2}-\d{2}(?:[ T].+)?", str(cell.value).strip()
                )
            ]
            numeric_columns = {
                coordinate_to_tuple(address)[1]
                for address, cell in entries
                if cell.kind == "number"
            }
            years_only = (
                numbers
                and not numeric_columns <= captioned_columns
                and all(
                    float(value).is_integer() and 1900 <= value <= 2100
                    for value in numbers
                )
            )
            leading_label = any(cell.kind == "string" for _, cell in entries)
            if (
                dates
                or (numbers and leading_label and not years_only)
                or (len(numbers) >= 2 and not years_only)
            ):
                first_data = row
                break
            captioned_columns.update(
                coordinate_to_tuple(address)[1]
                for address, cell in entries
                if cell.kind == "string" and cell.value
            )
        if first_data is None:
            return {}
        max_col = (
            bounds[2]
            if bounds
            else max(coordinate_to_tuple(address)[1] for address in sheet.cells)
        )
        min_col = min(
            coordinate_to_tuple(address)[1]
            for entries in rows.values()
            for address, _ in entries
        )
        paths: dict[str, list[str]] = {}
        merged = {area.split(":")[0]: range_boundaries(area) for area in sheet.merges}
        for row, entries in sorted(rows.items()):
            if row >= first_data:
                break
            for address, cell in entries:
                _, col = coordinate_to_tuple(address)
                value = cell.value
                if value is None or not str(value).strip():
                    continue
                bounds = merged.get(address)
                # Broad text titles/footnotes are context, not a column-specific header.
                if (
                    bounds
                    and isinstance(value, str)
                    and bounds[2] - bounds[0] + 1 >= max_col - min_col + 1
                ):
                    continue
                text = (
                    re.sub(r"(?<=[A-Za-z])-\s*\n\s*(?=[a-z])", "", str(value))
                    .replace("\n", " ")
                    .strip()
                )
                columns = range(bounds[0], bounds[2] + 1) if bounds else (col,)
                for index in columns:
                    key = get_column_letter(index)
                    if text not in paths.setdefault(key, []):
                        paths[key].append(text)
        return paths

    def to_text(
        self, compact: bool = False, *, metadata: dict[str, Any] | None = None
    ) -> str:
        """Keep every populated cell, including titles, units and footnotes."""
        blocks = []
        for name, sheet in self.sheets.items():
            lines = [f"SHEET {json.dumps(name, ensure_ascii=False)} [{sheet.state}]"]
            headers = (
                metadata[name]["headers"]
                if metadata is not None
                else self.header_paths(name)
            )
            if headers:
                lines.append(
                    "HEADER HINTS (verify against original rows): "
                    + json.dumps(headers, ensure_ascii=False)
                )
            if compact:
                lines.append(
                    "CELL VALUES use JSON types (strings quoted, numbers/booleans literal); "
                    "tags d=datetime,e=error,u=uncached_formula. JSON null is blank. "
                    "All populated cells follow; no rows omitted."
                )
            tables = (
                metadata[name]["tables"]
                if metadata is not None
                else self.table_regions(name)
            )
            if len(tables) > 1:
                lines.append(
                    "TABLE HINTS (advisory): " + json.dumps(tables, ensure_ascii=False)
                )
            if sheet.merges:
                lines.append(
                    "MERGES (only top-left stores a value): " + ", ".join(sheet.merges)
                )
            if sheet.hidden_rows or sheet.hidden_columns:
                lines.append(
                    f"HIDDEN rows={sheet.hidden_rows}, columns={sheet.hidden_columns}"
                )
            rows: dict[int, list[str]] = {}
            for address, cell in sheet.cells.items():
                row, _ = coordinate_to_tuple(address)
                flags = []
                if cell.indent:
                    flags.append(f"indent={cell.indent}")
                if cell.bold:
                    flags.append("bold")
                if cell.number_format not in {
                    "General",
                    "0",
                    "0.0",
                    "0.00",
                    "#,##0",
                    "#,##0.0",
                    "#,##0.00",
                }:
                    flags.append(f"format={cell.number_format}")
                if cell.formula:
                    flags.append(f"formula={cell.formula}")
                kind = (
                    {
                        "string": "s",
                        "number": "n",
                        "datetime": "d",
                        "boolean": "b",
                        "error": "e",
                        "uncached_formula": "u",
                    }.get(cell.kind, cell.kind)
                    if compact
                    else cell.kind
                )
                label = ";".join([kind, *flags])
                value = json.dumps(cell.wire()["value"], ensure_ascii=False)
                tag = (
                    ""
                    if compact
                    and not flags
                    and cell.kind in {"number", "string", "boolean"}
                    else f"<{label}>"
                )
                rows.setdefault(row, []).append(f"{address}{tag}={value}")
            lines.extend(
                ("" if compact else f"ROW {row}: ")
                + (" ".join(cells) if compact else " | ".join(cells))
                for row, cells in rows.items()
            )
            blocks.append("\n".join(lines))
        return "\n\n".join(blocks)


def _read_styles(archive: ZipFile) -> list[tuple[str, int, bool]]:
    if "xl/styles.xml" not in archive.namelist():
        return [("General", 0, False)]
    root = ET.fromstring(archive.read("xl/styles.xml"))
    formats = dict(BUILTIN_FORMATS)
    formats.update(
        (int(node.attrib["numFmtId"]), node.attrib["formatCode"])
        for node in root.findall(f"{NS}numFmts/{NS}numFmt")
    )
    fonts = root.findall(f"{NS}fonts/{NS}font")
    styles = []
    for node in root.findall(f"{NS}cellXfs/{NS}xf"):
        align = node.find(f"{NS}alignment")
        font_id = int(node.get("fontId", "0"))
        bold_node = fonts[font_id].find(f"{NS}b") if font_id < len(fonts) else None
        styles.append(
            (
                formats.get(int(node.get("numFmtId", "0")), "General"),
                int(float(align.get("indent", "0"))) if align is not None else 0,
                bold_node is not None
                and bold_node.get("val", "1") not in {"0", "false"},
            )
        )
    return styles or [("General", 0, False)]


def load_workbook_data(
    filepath: Path | str, max_bytes: int = 128 * 1024**2
) -> WorkbookData:
    """Stream actual XML cells; reject excessive inputs explicitly, never truncate."""
    size = 0
    sheets: dict[str, SheetData] = {}
    with ZipFile(filepath) as archive:
        metadata = ET.fromstring(archive.read("xl/workbook.xml"))
        properties = metadata.find(f"{NS}workbookPr")
        epoch = (
            CALENDAR_MAC_1904
            if properties is not None and properties.get("date1904") in {"1", "true"}
            else CALENDAR_WINDOWS_1900
        )
        relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        paths = {
            node.attrib["Id"]: posixpath.normpath("xl/" + node.attrib["Target"])
            if not node.attrib["Target"].startswith("/")
            else node.attrib["Target"].lstrip("/")
            for node in relations
            if node.get("TargetMode") != "External"
        }
        styles = _read_styles(archive)
        strings: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            with archive.open("xl/sharedStrings.xml") as stream:
                for _, element in ET.iterparse(stream, events=("end",)):
                    if element.tag == f"{NS}si":
                        value = "".join(n.text or "" for n in element.iter(f"{NS}t"))
                        strings.append(value)
                        size += sys.getsizeof(value) + 16
                        element.clear()
                        if size > max_bytes:
                            raise MemoryError(
                                "Shared strings exceed workbook memory budget"
                            )
        for node in metadata.findall(f"{NS}sheets/{NS}sheet"):
            name = node.attrib["name"]
            sheet = SheetData(name, state=node.get("state", "visible"))
            sheets[name] = sheet
            with archive.open(paths[node.attrib[f"{REL_NS}id"]]) as stream:
                for _, element in ET.iterparse(stream, events=("end",)):
                    if element.tag == f"{NS}c":
                        address = element.attrib["r"]
                        v = element.find(f"{NS}v")
                        f = element.find(f"{NS}f")
                        cell_type = element.get("t", "n")
                        fmt, indent, bold = styles[int(element.get("s", "0"))]
                        value: Any = v.text if v is not None else None
                        kind = "string"
                        if cell_type == "inlineStr":
                            value = "".join(
                                n.text or "" for n in element.iter(f"{NS}t")
                            )
                        elif cell_type == "s" and value is not None:
                            value = strings[int(value)]
                        elif cell_type == "b" and value is not None:
                            value, kind = value == "1", "boolean"
                        elif cell_type == "n" and value is not None:
                            value = float(value)
                            if is_date_format(fmt):
                                value = from_excel(value, epoch)
                                kind = (
                                    "time" if isinstance(value, dt.time) else "datetime"
                                )
                            else:
                                value = int(value) if value.is_integer() else value
                                kind = "number"
                        elif cell_type == "d" and value is not None:
                            value, kind = dt.datetime.fromisoformat(value), "datetime"
                        elif cell_type == "e":
                            kind = "error"
                        formula = "=" + (f.text or "") if f is not None else None
                        if formula and value is None:
                            kind = "uncached_formula"
                        if value is not None or formula:
                            sheet.cells[address] = CellData(
                                value, kind, fmt, formula, indent, bold
                            )
                            size += sys.getsizeof(value) + sys.getsizeof(address) + 512
                            if size > max_bytes:
                                raise MemoryError(
                                    "Workbook exceeds sparse cache budget; no rows were truncated"
                                )
                        element.clear()
                    elif element.tag == f"{NS}row":
                        if element.get("hidden") in {"1", "true"}:
                            sheet.hidden_rows.append(int(element.attrib["r"]))
                        element.clear()
                    elif element.tag == f"{NS}mergeCell":
                        sheet.merges.append(element.attrib["ref"])
                        element.clear()
                    elif element.tag == f"{NS}col":
                        if element.get("hidden") in {"1", "true"}:
                            sheet.hidden_columns.append(
                                f"{element.attrib['min']}:{element.attrib['max']}"
                            )
                        element.clear()
    return WorkbookData(sheets, size)


def unmerge_sheet(ws: Worksheet) -> list[list[Any]]:
    """Compatibility view: inherit merged text labels, never duplicate measures."""
    grid = [[cell.value for cell in row] for row in ws]
    for merged in ws.merged_cells.ranges:
        left, top, right, bottom = range_boundaries(str(merged))
        value = ws.cell(top, left).value
        if isinstance(value, str):
            for row in range(top - 1, bottom):
                for col in range(left - 1, right):
                    grid[row][col] = value
    return grid


def grid_to_clean_text(grid: list[list[Any]], max_rows: int | None = None) -> str:
    """Format a 2D grid without silently dropping rows."""
    if max_rows is not None and len(grid) > max_rows:
        raise ValueError("Explicit row limit would truncate workbook")
    return "\n".join(
        f"Row {idx:02d}: "
        + " | ".join(
            f"{get_column_letter(col)}={value}"
            for col, value in enumerate(row, 1)
            if value is not None
        )
        for idx, row in enumerate(grid)
        if any(value is not None for value in row)
    )


def parse_excel_file(filepath: Path | str) -> str:
    """Parse an entire Excel workbook into structured text representation."""
    return load_workbook_data(filepath).to_text()
