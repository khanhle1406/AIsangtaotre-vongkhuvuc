"""Compile a bounded spreadsheet expression language to the existing grounded worker.

This compiler has no answer data and never computes measures in the application
process. Every source value is read inside the regular restricted executor.
"""

import math
import re
from dataclasses import dataclass, field
from typing import Any

from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries

from .executor import MAX_ITEMS
from .llm_agent import Program
from .verification import age_reference_columns, canonical_label, source_headers


def fields(
    value: Any, required: set[str], optional: set[str] | None = None
) -> dict[str, Any]:
    """Reject ignored/ambiguous keys, including instructions disguised as metadata."""
    if (
        not isinstance(value, dict)
        or not required <= set(value)
        or set(value) - required - (optional or set())
    ):
        raise ValueError(
            f"Expected keys {sorted(required)}, optional {sorted(optional or set())}"
        )
    return value


@dataclass
class PlanCompiler:
    workbook: dict[str, Any]
    query: str
    statements: list[str] = field(default_factory=list)
    nodes: int = 0
    sequence: int = 0
    source_count: int = 0

    def ranking_evidence(self, expression: Any) -> None:
        """A guessed cell is not an executable proof of a requested ranking.

        This is a capability check, not an answer calculation. Published rank
        fields/unsupported algorithms can still use the normal Python fallback.
        """
        if not re.search(
            r"\b(?:highest|largest|lowest|smallest|oldest|youngest|rank|most|least)\b",
            self.query,
            re.IGNORECASE,
        ):
            return

        def objects(value: Any) -> list[dict[str, Any]]:
            if isinstance(value, dict):
                return [
                    value,
                    *[item for child in value.values() for item in objects(child)],
                ]
            if isinstance(value, list):
                return [item for child in value for item in objects(child)]
            return []

        nodes = objects(expression)
        scans = [node for node in nodes if node.get("op") == "scan" and "sort" in node]
        if not scans and not any(node.get("op") in {"max", "min"} for node in nodes):
            raise ValueError(
                "Ranking needs an explicit scan/sort or min/max over workbook candidates; a guessed output cell is insufficient"
            )
        if re.search(r"\b(?:oldest|youngest)\b", self.query, re.IGNORECASE):
            for sheet, data in self.workbook.items():
                reference = age_reference_columns(data, self.query)
                if reference and not any(
                    scan.get("sheet") == sheet
                    and set(reference)
                    <= {
                        node.get("col")
                        for node in objects(scan["sort"])
                        if isinstance(node.get("col"), str)
                    }
                    for scan in scans
                ):
                    raise ValueError(
                        "Elapsed-age rank needs a scan sorted from each person's birth AND reference date; a guessed score/name or birth-date order is insufficient"
                    )

    def variable(self) -> str:
        self.sequence += 1
        return f"plan_{self.sequence}"

    def constant(self, value: Any) -> str:
        if (
            type(value) not in {int, float, str, bool}
            or isinstance(value, float)
            and not math.isfinite(value)
        ):
            raise ValueError("Plan constants must be finite scalars")
        if isinstance(value, str):
            if len(value) > 500:
                raise ValueError("Excessive string constant")
        elif type(value) in {int, float} and value not in {0, 1, 100}:
            tokens = re.findall(r"[-+]?\d+(?:\.\d+)?", self.query.replace(",", ""))
            if value not in {float(token) for token in tokens}:
                raise ValueError("Numeric constants must be stated in the question")
        return repr(value)

    def source(self, node: dict[str, Any], sheet: str, address: str) -> str:
        self.source_count += 1
        if self.source_count > MAX_ITEMS:
            raise ValueError("Excessive total source expansion")
        if (
            sheet not in self.workbook
            or not isinstance(address, str)
            or not re.fullmatch(r"[A-Z]{1,3}[1-9]\d{0,6}", address)
        ):
            raise ValueError("Invalid or foreign source coordinate")
        row, col = coordinate_to_tuple(address)
        if row > 1048576 or col > 16384:
            raise ValueError("Source outside Excel bounds")
        kind = node.get("kind", "number")
        if kind not in {"number", "text", "date", "raw"}:
            raise ValueError("Unknown plan source kind")
        if "header" in node:
            header = node["header"]
            actual = source_headers(self.workbook[sheet], address).get(
                get_column_letter(col), []
            )
            if (
                not isinstance(header, list)
                or not header
                or any(not isinstance(item, str) for item in header)
                or [canonical_label(item) for item in header]
                != [canonical_label(item) for item in actual]
            ):
                raise ValueError(
                    f"Header path disagrees with original source {sheet}!{address}: {actual!r}"
                )
        helper = {"number": "num", "text": "label", "date": "cell", "raw": "cell"}[kind]
        return helper

    def quoted_metric(self, metric: Any, sheet: str, row: int) -> None:
        """For a direct ranking metric, verify a unique quoted header branch.

        Derived formulas are not mistaken for a reported column. A citation alone
        is insufficient: inspect the actual selected table's header paths.
        """
        if (
            isinstance(metric, dict)
            and metric.get("op") == "abs"
            and len(metric.get("args", [])) == 1
        ):
            metric = metric["args"][0]
        if not isinstance(metric, dict) or "col" not in metric:
            return
        column = metric["col"]
        if not isinstance(column, str) or not re.fullmatch(r"[A-Z]{1,3}", column):
            return  # The expression validator reports malformed columns.
        headers = source_headers(self.workbook[sheet], f"{column}{row}")
        numeric_columns = {
            get_column_letter(coordinate_to_tuple(address)[1])
            for address, entry in self.workbook[sheet]["cells"].items()
            if entry["kind"] in {"number", "uncached_formula"}
        }
        for quote in re.findall(r'["“]([^"”]+)["”]', self.query):
            quote = canonical_label(quote).casefold()
            targets = {
                col
                for col, path in headers.items()
                if quote in {canonical_label(label).casefold() for label in path}
            }
            targets &= numeric_columns
            if not targets or column in targets:
                continue
            # A caption repeated under competing ancestors does not identify a branch.
            contexts = {
                tuple(path[: index + 1])
                for path in headers.values()
                for index, label in enumerate(path)
                if canonical_label(label).casefold() == quote
            }
            if len(contexts) == 1:
                raise ValueError(
                    f"Ranking metric {column} contradicts quoted header {quote!r}; matching columns {sorted(targets)}"
                )

    def expression(
        self, node: Any, scope: tuple[str, str, list[int]] | None = None, depth: int = 0
    ) -> tuple[str, bool]:
        self.nodes += 1
        if self.nodes > 256 or depth > 16:
            raise ValueError("Excessive plan complexity")
        if not isinstance(node, dict):
            raise TypeError("Expression must be an object")
        if "const" in node:
            fields(node, {"const"})
            return self.constant(node["const"]), False
        if "cell" in node:
            fields(node, {"cell"}, {"kind", "header"})
            value = node["cell"]
            if (
                not isinstance(value, list)
                or len(value) != 2
                or any(not isinstance(item, str) for item in value)
            ):
                raise ValueError("cell requires [sheet,address]")
            sheet, address = value
            helper = self.source(node, sheet, address)
            return f"{helper}({sheet!r},{address!r})", True
        if "range" in node:
            fields(node, {"range"})
            value = node["range"]
            if (
                not isinstance(value, list)
                or len(value) != 2
                or any(not isinstance(item, str) for item in value)
            ):
                raise ValueError("range requires [sheet,original_area]")
            sheet, area = value
            if not re.fullmatch(
                r"[A-Z]{1,3}[1-9]\d{0,6}(?::[A-Z]{1,3}[1-9]\d{0,6})?", area
            ):
                raise ValueError("Invalid numeric source range")
            left, top, right, bottom = range_boundaries(area)
            count = (right - left + 1) * (bottom - top + 1)
            if (
                right < left
                or bottom < top
                or count > MAX_ITEMS
                or self.source_count + count > MAX_ITEMS
            ):
                raise ValueError("Invalid or excessive numeric source range")
            self.source({}, sheet, f"{get_column_letter(left)}{top}")
            self.source({}, sheet, f"{get_column_letter(right)}{bottom}")
            self.source_count += max(0, count - 2)
            return (
                f"num({sheet!r},{area!r})"
                if ":" in area
                else f"[num({sheet!r},{area!r})]",
                True,
            )
        if "col" in node:
            fields(node, {"col"}, {"kind", "header"})
            if (
                scope is None
                or not isinstance(node["col"], str)
                or not re.fullmatch(r"[A-Z]{1,3}", node["col"])
            ):
                raise ValueError(
                    "col requires an enclosing scan and original column letters"
                )
            sheet, variable, rows = scope
            helpers = {self.source(node, sheet, f"{node['col']}{row}") for row in rows}
            return f"{helpers.pop()}({sheet!r},{node['col']!r}+str({variable}))", True
        op = node.get("op")
        if op == "scan":
            if scope is not None:
                raise ValueError("Nested scans are unsupported; defer to Python")
            fields(
                node,
                {"op", "sheet", "rows", "value"},
                {"where", "sort", "descending", "take"},
            )
            sheet, bounds = node["sheet"], node["rows"]
            if not isinstance(sheet, str) or sheet not in self.workbook:
                raise ValueError("Scan references foreign sheet")
            if (
                not isinstance(bounds, list)
                or len(bounds) != 2
                or any(type(item) is not int for item in bounds)
                or not 1 <= bounds[0] <= bounds[1] <= 1048576
                or bounds[1] - bounds[0] + 1 > MAX_ITEMS
            ):
                raise ValueError("Scan requires bounded inclusive [first_row,last_row]")
            rows = list(range(bounds[0], bounds[1] + 1))
            variable, selected, result = (
                self.variable(),
                self.variable(),
                self.variable(),
            )
            row_scope = (sheet, variable, rows)
            predicate, grounded = (
                self.expression(node["where"], row_scope, depth + 1)
                if "where" in node
                else ("True", False)
            )
            self.statements.append(
                f"{selected}=[{variable} for {variable} in range({bounds[0]},{bounds[1] + 1}) if {predicate}]"
            )
            if "sort" in node:
                self.quoted_metric(node["sort"], sheet, bounds[0])
                metric, reads = self.expression(node["sort"], row_scope, depth + 1)
                grounded |= reads
                descending = node.get("descending", False)
                if type(descending) is not bool:
                    raise ValueError("descending must be boolean")
                self.statements.append(
                    f"{selected}=sorted({selected},key=lambda {variable}:{metric},reverse={descending})"
                )
            elif "descending" in node:
                raise ValueError("descending requires sort")
            if "take" in node:
                take = node["take"]
                if (
                    not isinstance(take, list)
                    or len(take) != 2
                    or any(type(item) is not int for item in take)
                    or min(take) < 0
                    or take[1] < 1
                    or take[0] + take[1] > MAX_ITEMS
                ):
                    raise ValueError("take requires [zero_based_offset,positive_count]")
                self.statements.append(f"{selected}={selected}[{take[0]}:{sum(take)}]")
            value, reads = self.expression(node["value"], row_scope, depth + 1)
            grounded |= reads
            if not reads:
                raise ValueError("Each scan value must depend on its workbook row")
            self.statements.append(f"{result}=[{value} for {variable} in {selected}]")
            return result, grounded
        fields(node, {"op", "args"})
        args = node["args"]
        if not isinstance(args, list) or not 1 <= len(args) <= 64:
            raise ValueError("Operator requires 1..64 expression arguments")
        expressions = [self.expression(item, scope, depth + 1) for item in args]
        values = [item[0] for item in expressions]
        grounded = any(item[1] for item in expressions)
        if op == "list":
            return "[" + ",".join(values) + "]", grounded
        binary = {
            "add": "+",
            "sub": "-",
            "mul": "*",
            "div": "/",
            "gt": ">",
            "ge": ">=",
            "lt": "<",
            "le": "<=",
            "eq": "==",
            "ne": "!=",
            "and": "and",
            "or": "or",
        }
        if op in binary and len(values) == 2:
            return f"({values[0]} {binary[op]} {values[1]})", grounded
        if len(values) != 1:
            raise ValueError("Unary/reduction operators take exactly one argument")
        value = values[0]
        unary = {
            "abs": f"abs({value})",
            "sum": f"sum({value})",
            "mean": f"statistics.mean({value})",
            "min": f"min({value})",
            "max": f"max({value})",
            "count": f"len({value})",
            "first": f"{value}[0]",
            "last": f"{value}[-1]",
            "clean": f"clean_label({value})",
            "not": f"(not {value})",
            "days": f"({value}).days",
        }
        if op not in unary:
            raise ValueError(
                f"Unsupported calculation operator {op!r}; defer to Python"
            )
        return unary[op], grounded

    def compile(self, plan: Any) -> Program:
        fields(plan, {"expr"}, {"format", "assumptions", "needs_verification"})
        expression, grounded = self.expression(plan["expr"])
        self.ranking_evidence(plan["expr"])
        if not grounded:
            raise ValueError("Result must depend on workbook sources")
        code = "import statistics\n" + "\n".join(
            [*self.statements, f"result={expression}"]
        )
        return Program.model_validate(
            {
                "code": code,
                **{
                    key: plan[key]
                    for key in ("format", "assumptions", "needs_verification")
                    if key in plan
                },
            }
        )


def compile_plan(plan: Any, workbook: dict[str, Any], query: str) -> str:
    """Return the regular Program envelope; execution/formatting/guards stay unchanged."""
    return PlanCompiler(workbook, query).compile(plan).model_dump_json()
