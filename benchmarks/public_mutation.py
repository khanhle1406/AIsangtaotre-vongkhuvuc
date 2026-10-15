"""Whole-public layout invariance: no IDs, paths or original coordinates survive."""

import csv
import json
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

from evaluate import read_answers
from src.excel_parser import load_workbook_data
from src.executor import WorkbookAccess


def generate(
    root: Path, row_offset: int = 217, col_offset: int = 2
) -> tuple[Path, Path]:
    if root.exists():
        raise ValueError("Preserve previous mutations")
    files = root / "xlsx_files"
    files.mkdir(parents=True)
    source = json.loads(Path("data/queries.json").read_text())
    gold = read_answers(Path("data/sample_submission.csv"))
    mapping: dict[str, str] = {}
    for item in source:
        name = item["file_name"]
        if name in mapping:
            continue
        original = openpyxl.load_workbook(
            Path("data/xlsx_files") / name, data_only=False
        )
        access = WorkbookAccess(
            load_workbook_data(Path("data/xlsx_files") / name).wire()
        )
        altered = openpyxl.Workbook()
        altered.remove(altered.active)
        for index, ws in enumerate(original.worksheets):
            target = altered.create_sheet(f"Ledger {len(mapping) + 31}-{index + 7}")
            for row in ws:
                for cell in row:
                    if cell.value is None:
                        continue
                    value = (
                        access.cell(ws.title, cell.coordinate)
                        if cell.data_type == "f"
                        else cell.value
                    )
                    updated = target.cell(
                        cell.row + row_offset, cell.column + col_offset, value
                    )
                    updated.number_format = cell.number_format
                    updated.font = copy(cell.font)
                    updated.alignment = copy(cell.alignment)
            for merged in ws.merged_cells.ranges:
                left, top, right, bottom = range_boundaries(str(merged))
                target.merge_cells(
                    f"{get_column_letter(left + col_offset)}{top + row_offset}:{get_column_letter(right + col_offset)}{bottom + row_offset}"
                )
        new = f"opaque_{len(mapping) * 13 + 401}.xlsx"
        altered.save(files / new)
        altered.close()
        original.close()
        mapping[name] = new
    queries = []
    references = []
    for index, item in enumerate(source):
        key = f"moved-{index * 31 + 1401}"
        queries.append(
            {"id": key, "file_name": mapping[item["file_name"]], "query": item["query"]}
        )
        references.append({"id": key, "answer": gold[item["id"]]})
    qpath, rpath = root / "queries.json", root / "references.csv"
    qpath.write_text(json.dumps(queries, indent=2))
    with rpath.open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["id", "answer"])
        writer.writeheader()
        writer.writerows(references)
    return qpath, rpath
