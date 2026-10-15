"""Independent unseen task families; freeze before candidate selection, evaluate once."""

import csv
import hashlib
import json
import random
from pathlib import Path

import openpyxl


def generate(root: Path, seed: int = 719403) -> tuple[Path, Path]:
    if root.exists():
        raise ValueError("Use a new directory; never overwrite a sealed dataset")
    files = root / "xlsx_files"
    files.mkdir(parents=True)
    rng = random.Random(seed)
    queries = []
    answers = []

    def add(file: str, family: str, query: str, answer: str) -> None:
        key = f"blind-{len(queries) + 1:03d}"
        queries.append({"id": key, "file_name": file, "query": query, "family": family})
        answers.append({"id": key, "answer": answer})

    # Duplicate identities and missing values: record counts and person counts differ.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Membership"
    ws.append(["Person ID", "Team", "Status", "Score"])
    records = []
    for i in range(12):
        row = (
            f"P{i + 41}",
            ["East", "West", "Central"][i % 3],
            "Completed" if i % 4 else "Pending",
            None if i % 5 == 0 else rng.randrange(20, 91),
        )
        records.append(row)
        ws.append(row)
        if i % 3 == 0:
            records.append(row)
            ws.append(row)
    wb.create_sheet("Prior snapshot").append(["Completed", 99999])
    file = "cohort_duplicate.xlsx"
    wb.save(files / file)
    wb.close()
    done = [row for row in records if row[2] == "Completed"]
    ids = {r[0] for r in done}
    add(
        file,
        "membership_dedup",
        "In Membership, how many distinct Person IDs have Completed status? Count each person once. Answer only the integer.",
        str(len(ids)),
    )
    add(
        file,
        "membership_dedup",
        "In Membership, how many records have Completed status, including duplicate records? Answer only the integer.",
        str(len(done)),
    )
    east = {r[0] for r in done if r[1] == "East"}
    add(
        file,
        "membership_dedup",
        "How many distinct Completed people belong to East in Membership? Answer only the integer.",
        str(len(east)),
    )
    scores = [r[3] for r in done if r[3] is not None]
    add(
        file,
        "membership_dedup",
        "In Membership, what is the average Score of Completed records? Include duplicate records; exclude blank scores. Answer rounded to two decimal places.",
        f"{sum(scores) / len(scores):.2f}",
    )
    missing = sum(r[3] is None for r in records)
    add(
        file,
        "membership_dedup",
        "How many records in Membership have a blank Score, regardless of Status? Answer only the integer.",
        str(missing),
    )
    teams = sorted({r[1] for r in records})
    counts = {team: len({r[0] for r in done if r[1] == team}) for team in teams}
    add(
        file,
        "membership_dedup",
        "Which Team has the most distinct Completed people in Membership? Break ties alphabetically. Answer only the Team name.",
        min(teams, key=lambda t: (-counts[t], t)),
    )

    # A transposed lookup matrix and a separate scenario select coordinates by labels.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Service matrix"
    ws.append(["Measure", "North", "South", "Central"])
    markets = ["North", "South", "Central"]
    numerator = [rng.randrange(20, 70) for _ in markets]
    denominator = [rng.randrange(90, 150) for _ in markets]
    fees = [rng.randrange(5, 20) for _ in markets]
    ws.append(["Resolved tickets", *numerator])
    ws.append(["Received tickets", *denominator])
    ws.append(["Fee, USD per ticket", *fees])
    scenario = wb.create_sheet("Scenario")
    chosen = markets[rng.randrange(3)]
    scenario.append(["Requested market", chosen])
    scenario.append(["Reporting currency", "USD"])
    noise = wb.create_sheet("Plan")
    noise.append(["Measure", *markets])
    noise.append(["Resolved tickets", 999, 999, 999])
    file = "matrix_scenario.xlsx"
    wb.save(files / file)
    wb.close()
    index = markets.index(chosen)
    add(
        file,
        "transposed_scenario",
        "For the Requested market named in Scenario, how many Resolved tickets are in Service matrix? Answer only the integer.",
        str(numerator[index]),
    )
    add(
        file,
        "transposed_scenario",
        "For the Requested market named in Scenario, what percentage of Received tickets were Resolved in Service matrix? Answer to two decimal places without percent sign.",
        f"{100 * numerator[index] / denominator[index]:.2f}",
    )
    add(
        file,
        "transposed_scenario",
        "In Service matrix, which market has the highest ratio of Resolved tickets to Received tickets? Answer only the market name.",
        markets[max(range(3), key=lambda i: numerator[i] / denominator[i])],
    )
    add(
        file,
        "transposed_scenario",
        "In Service matrix, how many Received tickets remain unresolved across all three markets? Answer only the integer.",
        str(sum(denominator) - sum(numerator)),
    )
    add(
        file,
        "transposed_scenario",
        "Using Service matrix, what is the total fee in USD for Resolved tickets across all markets? Multiply each market count by its own Fee. Answer only the integer.",
        str(sum(n * f for n, f in zip(numerator, fees, strict=True))),
    )
    add(
        file,
        "transposed_scenario",
        "Using Service matrix, what is the overall resolution percentage across all markets? Use total Resolved divided by total Received, not an unweighted mean of rates. Answer to two decimal places without percent sign.",
        f"{100 * sum(numerator) / sum(denominator):.2f}",
    )

    # Opening stock is separate from signed movements; price weighting is by receipts.
    wb = openpyxl.Workbook()
    opening = wb.active
    opening.title = "Opening stock"
    opening.append(["SKU", "Opening units"])
    skus = ["Jade", "Opal", "Onyx"]
    stock = {sku: rng.randrange(60, 100) for sku in skus}
    for sku in skus:
        opening.append([sku, stock[sku]])
    ws = wb.create_sheet("Movements")
    ws.append(["SKU", "Movement", "Units", "Unit cost USD"])
    moves = []
    for sku in skus:
        for kind, quantity in [
            ("Receipt", rng.randrange(10, 30)),
            ("Receipt", rng.randrange(10, 30)),
            ("Shipment", -rng.randrange(15, 45)),
            ("Loss", -rng.randrange(1, 8)),
        ]:
            row = (
                sku,
                kind,
                quantity,
                rng.randrange(4, 12) if kind == "Receipt" else None,
            )
            moves.append(row)
            ws.append(row)
    file = "stock_reconciliation.xlsx"
    wb.save(files / file)
    wb.close()
    balance = {
        sku: stock[sku] + sum(r[2] for r in moves if r[0] == sku) for sku in skus
    }
    add(
        file,
        "stock_reconciliation",
        "Using Opening stock plus all signed Movements, what is the closing stock for Jade? Answer only the integer.",
        str(balance["Jade"]),
    )
    add(
        file,
        "stock_reconciliation",
        "Using Opening stock plus all signed Movements, what is the combined closing stock for all SKUs? Answer only the integer.",
        str(sum(balance.values())),
    )
    add(
        file,
        "stock_reconciliation",
        "Which SKU has the largest closing stock, calculated from Opening stock plus all signed Movements? Answer only the SKU.",
        max(skus, key=lambda sku: balance[sku]),
    )
    receipts = [r for r in moves if r[1] == "Receipt"]
    cost = sum(r[2] * r[3] for r in receipts)
    units = sum(r[2] for r in receipts)
    add(
        file,
        "stock_reconciliation",
        "In Movements, what is the weighted average Unit cost USD for Receipt rows? Weight by Units; exclude Shipment and Loss. Answer to two decimal places.",
        f"{cost / units:.2f}",
    )
    add(
        file,
        "stock_reconciliation",
        "In Movements, how many units were lost across all SKUs? Report the positive magnitude of the Loss rows only. Answer only the integer.",
        str(-sum(r[2] for r in moves if r[1] == "Loss")),
    )
    add(
        file,
        "stock_reconciliation",
        "In Movements, what is the net unit change for Opal across all movement kinds? Keep its sign. Answer only the integer.",
        str(sum(r[2] for r in moves if r[0] == "Opal")),
    )

    # Quarters and months share a table but are not a disjoint set; scale is explicit.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fiscal sales"
    ws.append(["Period", "Revenue, thousand USD", "Cost, thousand USD"])
    months = [rng.randrange(20, 90) for _ in range(6)]
    costs = [rng.randrange(5, 18) for _ in range(6)]
    rows = [("H1", sum(months), sum(costs)), ("Q1", sum(months[:3]), sum(costs[:3]))]
    rows += [
        (name, months[i], costs[i])
        for i, name in enumerate(["January", "February", "March"])
    ]
    rows += [("Q2", sum(months[3:]), sum(costs[3:]))]
    rows += [
        (name, months[i], costs[i])
        for i, name in enumerate(["April", "May", "June"], start=3)
    ]
    for row in rows:
        ws.append(row)
    ws["A13"] = (
        "Q1 = January–March; Q2 = April–June; H1 = Q1 + Q2. Values are thousand USD."
    )
    file = "fiscal_partition.xlsx"
    wb.save(files / file)
    wb.close()
    add(
        file,
        "fiscal_partition",
        "What is H1 Revenue in Fiscal sales? Retain the table unit of thousand USD. Answer only the integer.",
        str(sum(months)),
    )
    add(
        file,
        "fiscal_partition",
        "What is H1 Revenue in actual USD, rather than thousand USD, in Fiscal sales? Answer only the integer.",
        str(sum(months) * 1000),
    )
    add(
        file,
        "fiscal_partition",
        "Using Fiscal sales, what is Q2 Revenue minus Q1 Revenue, in thousand USD? Keep the sign. Answer only the integer.",
        str(sum(months[3:]) - sum(months[:3])),
    )
    add(
        file,
        "fiscal_partition",
        "What is total Revenue for February and April only in Fiscal sales, in thousand USD? Answer only the integer.",
        str(months[1] + months[3]),
    )
    add(
        file,
        "fiscal_partition",
        "What is H1 profit in Fiscal sales, defined as Revenue minus Cost, in thousand USD? Answer only the integer.",
        str(sum(months) - sum(costs)),
    )
    add(
        file,
        "fiscal_partition",
        "Using Fiscal sales, what is H1 profit margin as a percent: (Revenue minus Cost) divided by Revenue? Answer to two decimal places without percent sign.",
        f"{100 * (sum(months) - sum(costs)) / sum(months):.2f}",
    )
    qpath, rpath = root / "queries.json", root / "references.csv"
    qpath.write_text(json.dumps(queries, indent=2))
    with rpath.open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["id", "answer"])
        writer.writeheader()
        writer.writerows(answers)
    manifest = {
        "seed": seed,
        "families": sorted({q["family"] for q in queries}),
        "questions": len(queries),
        "sha256": {
            p.relative_to(root).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in [qpath, rpath, *files.glob("*.xlsx")]
        },
        "status": "sealed before candidate selection; references never go to inference",
    }
    (root / "manifest.json").write_text(json.dumps(manifest, indent=2))
    return qpath, rpath
