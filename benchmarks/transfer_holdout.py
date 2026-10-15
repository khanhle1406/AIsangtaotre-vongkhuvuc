"""Independent progressive policy, anti-join and interval-union oracle workbooks."""

import csv
import hashlib
import json
import random
from pathlib import Path
from typing import Any

import openpyxl


def generate(root: Path, seed: int = 19620483) -> None:
    """Create once; answers are computed without any solver/model imports."""
    if root.exists():
        raise ValueError("Preserve the frozen transfer holdout")
    files = root / "xlsx_files"
    files.mkdir(parents=True)
    rng = random.Random(seed)
    queries: list[dict[str, Any]] = []
    answers: list[dict[str, str]] = []

    def save(
        family: str,
        tables: list[tuple[str, list[list[Any]]]],
        cases: list[tuple[str, str]],
    ) -> None:
        book = openpyxl.Workbook()
        book.remove(book.active)
        for name, rows in tables:
            sheet = book.create_sheet(name)
            row0 = rng.choice([19, 287, 503])
            column0 = rng.choice([2, 4, 6])
            for r, row in enumerate(rows, row0):
                for c, value in enumerate(row, column0):
                    sheet.cell(r, c, value)
        book.create_sheet("Unapproved archive").append(
            ["Do not use archive values", 999999]
        )
        filename = f"audit_{rng.randrange(100000, 999999)}.xlsx"
        book.save(files / filename)
        book.close()
        for query, answer in cases:
            key = f"transfer-{len(queries) + 1:03d}"
            queries.append(
                {"id": key, "file_name": filename, "family": family, "query": query}
            )
            answers.append({"id": key, "answer": answer})

    bands = [
        (0, 25000, 0.10),
        (25000, 70000, 0.20),
        (70000, 130000, 0.30),
        (130000, 400000, 0.40),
    ]
    people = ["Mara", "Nico", "Odel", "Pia", "Rami", "Tova"]
    incomes = [18000, 55000, 97000, 180000, 25000, 70000]

    def tax(income: float) -> float:
        return sum(
            max(0, min(income, upper) - lower) * rate for lower, upper, rate in bands
        )

    taxes = [tax(income) for income in incomes]
    net = [income - paid for income, paid in zip(incomes, taxes, strict=True)]
    save(
        "progressive_policy",
        [
            (
                "Policy",
                [
                    ["Lower USD inclusive", "Upper USD exclusive", "Rate fraction"],
                    *[list(row) for row in bands],
                ],
            ),
            (
                "Incomes",
                [
                    ["Person", "Income USD"],
                    *[
                        [name, value]
                        for name, value in zip(people, incomes, strict=True)
                    ],
                ],
            ),
        ],
        [
            (
                "Using Policy and Incomes, calculate progressive tax for Nico: only the portion of income within each band is taxed at that band rate. Answer to two decimal places, without currency symbol.",
                f"{taxes[1]:.2f}",
            ),
            (
                "Using Policy and Incomes, what is Nico effective tax rate: progressive tax divided by Income USD, expressed in percent? Answer to two decimal places without percent sign.",
                f"{100 * taxes[1] / incomes[1]:.2f}",
            ),
            (
                "Using Policy and Incomes, what marginal rate applies to Odel at their current income? Express Rate fraction as percent, answer to two decimal places without percent sign.",
                f"{100 * next(rate for lower, upper, rate in bands if lower <= incomes[2] < upper):.2f}",
            ),
            (
                "Using Policy and Incomes, sum progressive tax for all people, applying each rate only to the income portion in its band. Answer to two decimal places.",
                f"{sum(taxes):.2f}",
            ),
            (
                "Using Policy and Incomes, which Person has highest income after progressive tax? Break ties alphabetically; answer with only the Person name.",
                min(people, key=lambda name: (-net[people.index(name)], name)),
            ),
            (
                "Using Policy and Incomes, how much additional progressive tax would Tova pay if Income USD increased by exactly 1000? Answer to two decimal places.",
                f"{tax(incomes[5] + 1000) - taxes[5]:.2f}",
            ),
            (
                "Using Policy and Incomes, how many people owe progressive tax strictly greater than 10000 USD? Answer only the integer.",
                str(sum(paid > 10000 for paid in taxes)),
            ),
            (
                "Using Policy and Incomes, what is Odel income after progressive tax? Answer to two decimal places.",
                f"{net[2]:.2f}",
            ),
        ],
    )

    ids = [f"Order-{i + 41}" for i in range(7)]
    ordered = [rng.randrange(30, 91) for _ in ids]
    shipments = [
        (ids[0], ordered[0]),
        (ids[1], 10),
        (ids[1], 13),
        (ids[2], ordered[2] + 7),
        (ids[3], 9),
        (ids[3], 8),
        (ids[5], ordered[5] - 4),
        ("Unknown-99", 6),
    ]
    shipped = {
        key: sum(quantity for identity, quantity in shipments if identity == key)
        for key in ids
    }
    remaining = {key: max(0, ordered[ids.index(key)] - shipped[key]) for key in ids}
    save(
        "reconciliation_anti_join",
        [
            (
                "Orders",
                [
                    ["Order ID", "Ordered units"],
                    *[[key, value] for key, value in zip(ids, ordered, strict=True)],
                ],
            ),
            (
                "Shipments",
                [["Order ID", "Shipped units"], *[list(row) for row in shipments]],
            ),
        ],
        [
            (
                "In Orders, what is total Ordered units? Answer only the integer.",
                str(sum(ordered)),
            ),
            (
                "Using Orders and Shipments, sum Shipped units only for Order IDs present in Orders. Include multiple shipment records for an order. Answer only the integer.",
                str(sum(shipped.values())),
            ),
            (
                "Using Orders and Shipments, how many shipment records reference an Order ID absent from Orders? Answer only the integer.",
                str(sum(key not in ids for key, _ in shipments)),
            ),
            (
                "Using Orders and Shipments, what is total positive unshipped units across orders? For each order use max(0, Ordered units minus its summed Shipped units); ignore unknown Order IDs. Answer only the integer.",
                str(sum(remaining.values())),
            ),
            (
                "Using Orders and Shipments, how many orders have summed Shipped units at least Ordered units? Ignore unknown Order IDs. Answer only the integer.",
                str(sum(shipped[key] >= ordered[ids.index(key)] for key in ids)),
            ),
            (
                "Using Orders and Shipments, list Order IDs with no shipment record at all, sorted alphabetically and separated by comma and space.",
                ", ".join(
                    sorted(
                        key
                        for key in ids
                        if not any(identity == key for identity, _ in shipments)
                    )
                ),
            ),
            (
                "Using Orders and Shipments, which Order ID has largest positive unshipped units after summing shipment records per order? Break ties alphabetically, answer only the Order ID.",
                min(ids, key=lambda key: (-remaining[key], key)),
            ),
            (
                "Using Orders and Shipments, sum positive overshipped units per known order, defined as max(0, summed Shipped units minus Ordered units). Ignore unknown Order IDs. Answer only the integer.",
                str(sum(max(0, shipped[key] - ordered[ids.index(key)]) for key in ids)),
            ),
        ],
    )

    intervals = [
        (15, 75),
        (50, 120),
        (190, 240),
        (230, 310),
        (400, 450),
        (445, 480),
        (55, 95),
    ]
    occupancy = [
        sum(start <= minute < end for start, end in intervals) for minute in range(600)
    ]
    occupied = sum(count > 0 for count in occupancy)
    longest = run = 0
    for count in occupancy:
        run = 0 if count else run + 1
        longest = max(longest, run)
    segments = sum(
        count > 0 and (minute == 0 or occupancy[minute - 1] == 0)
        for minute, count in enumerate(occupancy)
    )
    save(
        "interval_union",
        [
            (
                "Reservations",
                [
                    ["Reservation", "Start minute inclusive", "End minute exclusive"],
                    *[
                        [f"Block-{i + 21}", start, end]
                        for i, (start, end) in enumerate(intervals)
                    ],
                ],
            )
        ],
        [
            (
                "In Reservations, within minute 0 inclusive to 600 exclusive, how many distinct minutes are covered by at least one reservation? Count overlapping minutes once. Answer only the integer.",
                str(occupied),
            ),
            (
                "In Reservations, sum all reservation durations in minutes, counting overlap separately for each reservation. Answer only the integer.",
                str(sum(end - start for start, end in intervals)),
            ),
            (
                "In Reservations, within minute 0 inclusive to 600 exclusive, how many minutes have no reservation? Answer only the integer.",
                str(600 - occupied),
            ),
            (
                "In Reservations, within minute 0 inclusive to 600 exclusive, what is the longest consecutive gap with no reservation? Include gaps at the beginning and end. Answer only the integer number of minutes.",
                str(longest),
            ),
            (
                "In Reservations, what is the maximum number of simultaneously active reservations? Start is inclusive and End exclusive. Answer only the integer.",
                str(max(occupancy)),
            ),
            (
                "In Reservations, what is the earliest minute when the maximum number of simultaneous reservations is active? Start is inclusive and End exclusive. Answer only the integer minute index.",
                str(occupancy.index(max(occupancy))),
            ),
            (
                "In Reservations, how many disjoint busy segments remain after merging overlapping or directly touching reservation intervals? Answer only the integer.",
                str(segments),
            ),
            (
                "In Reservations, how many distinct minutes are covered by at least two reservations? Count each minute once. Answer only the integer.",
                str(sum(count >= 2 for count in occupancy)),
            ),
        ],
    )
    (root / "queries.json").write_text(json.dumps(queries, indent=2))
    with (root / "references.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["id", "answer"])
        writer.writeheader()
        writer.writerows(answers)
    paths = [
        root / "queries.json",
        root / "references.csv",
        *sorted(files.glob("*.xlsx")),
    ]
    (root / "manifest.json").write_text(
        json.dumps(
            {
                "seed": seed,
                "questions": len(queries),
                "families": sorted({q["family"] for q in queries}),
                "hashes": {
                    str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest()
                    for p in paths
                },
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    generate(Path("reports/optimization/efficiency_v5_holdout"))
