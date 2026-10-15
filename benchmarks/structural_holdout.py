"""Independent FIFO allocation, interpolation/integration and nested bill-of-materials."""

import argparse
import csv
import datetime as dt
import hashlib
import itertools
import json
import random
from collections import Counter
from decimal import Decimal
from fractions import Fraction
from pathlib import Path
from typing import Any

import openpyxl


def generate(root: Path, seed: int = 548031) -> None:
    """Freeze once; pure independent oracles never call solver/executor helpers."""
    if root.exists():
        raise ValueError("Preserve frozen structural test data")
    files = root / "xlsx_files"
    files.mkdir(parents=True)
    rng = random.Random(seed)
    queries: list[dict[str, Any]] = []
    answers: list[dict[str, str]] = []

    def save(
        family: str,
        tables: list[tuple[str, list[list[Any]]]],
        cases: list[tuple[str, str]],
    ) -> None:
        book = openpyxl.Workbook()
        book.remove(book.active)
        for name, rows in tables:
            sheet = book.create_sheet(name)
            top, left = rng.choice([29, 331, 701]), rng.choice([2, 5, 8])
            for row, values in enumerate(rows, top):
                for col, value in enumerate(values, left):
                    sheet.cell(row, col, value)
        book.create_sheet("Superseded archive").append(["Excluded archive", 999999])
        filename = f"record_{rng.randrange(100000, 999999)}.xlsx"
        book.save(files / filename)
        book.close()
        for query, answer in cases:
            key = f"structural-{len(queries) + 1:03d}"
            queries.append(
                {"id": key, "file_name": filename, "family": family, "query": query}
            )
            answers.append({"id": key, "answer": answer})

    available = 137
    orders = [
        (f"Order-{chr(65 + i)}", dt.datetime(2032, 6, 1 + i), qty)
        for i, qty in enumerate([33, 51, 28, 42, 19, 36])
    ]

    def allocation(stock: int) -> dict[str, int]:
        # Closed-form cumulative-prefix oracle, rather than a mutable stock loop.
        return {
            key: min(qty, max(0, stock - sum(prior[2] for prior in orders[:i])))
            for i, (key, _, qty) in enumerate(orders)
        }

    allocated = allocation(available)
    prefix = "Use Stock settings and Requests only. Allocate Available units FIFO by Created timestamp ascending, ties by Order ID alphabetically; partial fills are allowed and no request receives more than Requested units. "
    cases = [
        (
            prefix
            + "How many units are allocated to Order-D? Answer only the integer.",
            str(allocated["Order-D"]),
        ),
        (
            prefix
            + "Which Order ID is partially filled, meaning allocated units are positive but smaller than Requested units? Answer only the Order ID.",
            next(key for key, _, qty in orders if 0 < allocated[key] < qty),
        ),
        (
            prefix
            + "What is total unfilled demand after allocation? Answer only the integer.",
            str(sum(qty - allocated[key] for key, _, qty in orders)),
        ),
        (
            prefix + "How many requests are fully filled? Answer only the integer.",
            str(sum(qty == allocated[key] for key, _, qty in orders)),
        ),
        (
            prefix
            + "Which is the earliest chronological Order ID receiving zero units? Answer only the Order ID.",
            next(key for key, _, _ in orders if allocated[key] == 0),
        ),
        (
            prefix
            + "If Available units increased by 20, how many units would Order-E receive? Answer only the integer.",
            str(allocation(available + 20)["Order-E"]),
        ),
    ]
    shuffled = [list(row) for row in orders]
    rng.shuffle(shuffled)
    save(
        "fifo_stock_allocation",
        [
            (
                "Stock settings",
                [["Parameter", "Value"], ["Available units", available]],
            ),
            (
                "Requests",
                [["Order ID", "Created timestamp", "Requested units"], *shuffled],
            ),
        ],
        cases,
    )

    knots = [(0, 4), (3, 16), (8, 6), (14, 23), (21, 11)]

    def interpolate(x: int) -> Fraction:
        for (left, y0), (right, y1) in itertools.pairwise(knots):
            if left <= x <= right:
                return Fraction(y0) + Fraction((y1 - y0) * (x - left), right - left)
        raise ValueError("Oracle does not extrapolate")

    def area(lower: int, upper: int) -> Fraction:
        # Exact antiderivative per clipped segment, independently of model floats.
        total = Fraction()
        for (left, y0), (right, y1) in itertools.pairwise(knots):
            a, b = max(lower, left), min(upper, right)
            if a < b:
                slope = Fraction(y1 - y0, right - left)
                total += y0 * (b - a) + slope * Fraction(
                    (b - left) ** 2 - (a - left) ** 2, 2
                )
        return total

    def decimal(value: Fraction) -> str:
        return f"{Decimal(value.numerator) / Decimal(value.denominator):.2f}"

    prefix = "Use Knot measurements only. Sort by X and linearly interpolate Y between neighboring knots; do not extrapolate. "
    cases = [
        (
            prefix + f"What is Y at X={x}? Answer to two decimal places.",
            decimal(interpolate(x)),
        )
        for x in [5, 11, 18]
    ]
    cases.extend(
        [
            (
                prefix
                + "What is the integral of this piecewise-linear Y from X=2 to X=18? Include partial end segments. Answer to two decimal places.",
                decimal(area(2, 18)),
            ),
            (
                prefix
                + "What is the integral of this piecewise-linear Y from X=0 to X=21? Answer to two decimal places.",
                decimal(area(0, 21)),
            ),
            (
                prefix
                + "What is the continuous average of Y over X=0 to X=21, defined as the integral divided by interval width? Answer to two decimal places.",
                decimal(area(0, 21) / 21),
            ),
        ]
    )
    shuffled = [list(row) for row in knots]
    rng.shuffle(shuffled)
    save(
        "piecewise_linear_integration",
        [("Knot measurements", [["X", "Y"], *shuffled])],
        cases,
    )

    rules = [
        ("Beacon", "Armature", 2),
        ("Beacon", "Core", 1),
        ("Armature", "Rod", 3),
        ("Armature", "Bolt", 2),
        ("Core", "Rod", 1),
        ("Core", "Wire", 4),
        ("Core", "Bolt", 1),
        ("Rod", "Alloy", 2),
        ("Crate", "Beacon", 1),
        ("Crate", "Core", 2),
    ]
    costs = {"Alloy": Decimal("3.25"), "Bolt": Decimal("0.80"), "Wire": Decimal("1.15")}

    def leaves(product: str) -> Counter[str]:
        result: Counter[str] = Counter()
        pending = [(product, 1)]
        while pending:
            item, quantity = pending.pop()
            if item in costs:
                result[item] += quantity
            else:
                pending.extend(
                    (child, quantity * multiplier)
                    for parent, child, multiplier in rules
                    if parent == item
                )
        return result

    def material_cost(product: str) -> Decimal:
        return sum(
            (costs[part] * quantity for part, quantity in leaves(product).items()),
            Decimal(),
        )

    beacon, crate = leaves("Beacon"), leaves("Crate")
    prefix = "Use Rules and Leaf costs only. Rules are acyclic; Units is component multiplicity for one parent. Recursively expand assemblies until every part is listed in Leaf costs, counting repeated paths separately. "
    cases = [
        (
            prefix
            + "What are total leaf quantities for Alloy, Bolt, Wire in one Beacon? Answer three integers in that order, separated by comma and one space.",
            ", ".join(str(beacon[part]) for part in ["Alloy", "Bolt", "Wire"]),
        ),
        (
            prefix
            + "What is total leaf material cost for one Beacon? Answer to two decimal places.",
            f"{material_cost('Beacon'):.2f}",
        ),
        (
            prefix
            + "What is total Wire quantity for one Crate? Answer only the integer.",
            str(crate["Wire"]),
        ),
        (
            prefix
            + "What is total leaf material cost for one Crate? Answer to two decimal places.",
            f"{material_cost('Crate'):.2f}",
        ),
        (
            prefix
            + "How much more is total material cost of one Crate than one Beacon? Answer to two decimal places.",
            f"{material_cost('Crate') - material_cost('Beacon'):.2f}",
        ),
        (
            prefix
            + "Which leaf Part contributes the largest total material cost to one Beacon? Break ties alphabetically; answer only the Part.",
            min(costs, key=lambda part: (-beacon[part] * costs[part], part)),
        ),
    ]
    shuffled = [list(row) for row in rules]
    rng.shuffle(shuffled)
    save(
        "recursive_bill_of_materials",
        [
            ("Rules", [["Parent", "Component", "Units"], *shuffled]),
            (
                "Leaf costs",
                [
                    ["Part", "Unit cost"],
                    *[[part, float(cost)] for part, cost in costs.items()],
                ],
            ),
        ],
        cases,
    )
    (root / "queries.json").write_text(
        json.dumps(queries, ensure_ascii=False, indent=2)
    )
    with (root / "references.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["id", "answer"])
        writer.writeheader()
        writer.writerows(answers)
    paths = [
        root / "queries.json",
        root / "references.csv",
        *sorted(files.glob("*.xlsx")),
    ]
    (root / "manifest.json").write_text(
        json.dumps(
            {
                "seed": seed,
                "questions": len(queries),
                "families": sorted({query["family"] for query in queries}),
                "hashes": {
                    str(path.relative_to(root)): hashlib.sha256(
                        path.read_bytes()
                    ).hexdigest()
                    for path in paths
                },
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, required=True)
    generate(parser.parse_args().root)
