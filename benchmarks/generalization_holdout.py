"""Freeze independent assignment, weighted regression and Markov workbook problems."""

import argparse
import csv
import hashlib
import itertools
import json
import random
from decimal import Decimal
from fractions import Fraction
from pathlib import Path
from typing import Any

import openpyxl


def generate(root: Path, seed: int = 601927) -> None:
    """Use exhaustive enumeration and exact rational oracles, without solver imports."""
    if root.exists():
        raise ValueError("Preserve frozen generalization data")
    files = root / "xlsx_files"
    files.mkdir(parents=True)
    rng = random.Random(seed)
    queries: list[dict[str, Any]] = []
    references: list[dict[str, str]] = []

    def fixed(value: Fraction) -> str:
        return f"{Decimal(value.numerator) / Decimal(value.denominator):.6f}"

    def save(
        family: str, tables: dict[str, list[list[Any]]], cases: list[tuple[str, str]]
    ) -> None:
        book = openpyxl.Workbook()
        book.remove(book.active)
        for name, rows in tables.items():
            sheet = book.create_sheet(name)
            top, left = rng.choice([41, 389, 811]), rng.choice([3, 7, 12])
            order = list(range(len(rows[0])))
            rng.shuffle(order)
            for r, row in enumerate(rows, top):
                for c, original in enumerate(order, left):
                    sheet.cell(r, c, row[original])
        book.create_sheet("Historical unrelated data").append(["Excluded", 987654])
        name = f"source_{rng.randrange(100000, 999999)}.xlsx"
        book.save(files / name)
        book.close()
        for question, answer in cases:
            key = f"generalization-{len(queries) + 1:03d}"
            queries.append(
                {"id": key, "file_name": name, "family": family, "query": question}
            )
            references.append({"id": key, "answer": answer})

    workers = ["Aster", "Birch", "Cedar", "Dune"]
    tasks = ["Polish", "Pack", "Inspect"]
    costs = {
        ("Aster", "Polish"): 7,
        ("Aster", "Pack"): 4,
        ("Aster", "Inspect"): 9,
        ("Birch", "Polish"): 3,
        ("Birch", "Pack"): 8,
        ("Cedar", "Pack"): 2,
        ("Cedar", "Inspect"): 5,
        ("Dune", "Polish"): 6,
        ("Dune", "Inspect"): 1,
    }

    def assignments(
        requested: list[str], excluded: str | None = None
    ) -> list[tuple[int, tuple[str, ...]]]:
        return sorted(
            (
                sum(
                    costs[worker, task]
                    for worker, task in zip(chosen, requested, strict=True)
                ),
                chosen,
            )
            for chosen in itertools.permutations(
                [w for w in workers if w != excluded], len(requested)
            )
            if all(
                (worker, task) in costs
                for worker, task in zip(chosen, requested, strict=True)
            )
        )

    best_cost, chosen = assignments(tasks)[0]
    assert best_cost < assignments(tasks)[1][0], "Oracle requires a unique optimum"
    prefix = "Use Allowed assignments only. Each worker can perform at most one task, each task needs one worker, and missing worker-task pairs are forbidden. "
    edges = [[worker, task, cost] for (worker, task), cost in costs.items()]
    rng.shuffle(edges)
    save(
        "minimum_cost_assignment",
        {"Allowed assignments": [["Worker", "Task", "Cost"], *edges]},
        [
            (
                prefix
                + "What is the maximum number of different tasks that can be covered? Answer only the integer.",
                str(
                    max(
                        n
                        for n in range(1, len(tasks) + 1)
                        if any(
                            assignments(list(subset))
                            for subset in itertools.combinations(tasks, n)
                        )
                    )
                ),
            ),
            (
                prefix
                + "What is the minimum total Cost covering all three tasks? Answer only the integer.",
                str(best_cost),
            ),
            (
                prefix
                + "In the unique minimum-Cost assignment covering all tasks, which Worker performs Pack? Answer only the name.",
                chosen[tasks.index("Pack")],
            ),
            (
                prefix
                + "In that minimum-Cost assignment covering all tasks, which Worker is unused? Answer only the name.",
                next(w for w in workers if w not in chosen),
            ),
            (
                prefix
                + "If Dune is unavailable, what is the minimum total Cost covering all tasks? Answer only the integer.",
                str(assignments(tasks, "Dune")[0][0]),
            ),
            (
                prefix
                + "If only Polish and Inspect need covering, what is the minimum total Cost? Answer only the integer.",
                str(assignments(["Polish", "Inspect"])[0][0]),
            ),
        ],
    )

    observations = [(0, 1, 1), (1, 4, 2), (2, 8, 1), (4, 11, 3)]
    w = sum(weight for x, y, weight in observations)
    sx = sum(x * weight for x, y, weight in observations)
    sy = sum(y * weight for x, y, weight in observations)
    sxx = sum(x * x * weight for x, y, weight in observations)
    sxy = sum(x * y * weight for x, y, weight in observations)
    syy = sum(y * y * weight for x, y, weight in observations)
    slope = Fraction(w * sxy - sx * sy, w * sxx - sx * sx)
    intercept = (sy - slope * sx) / w
    residual = syy - intercept * sy - slope * sxy
    prefix = "Use Observations only. Fit Y=a+b*X by minimizing sum(Weight*(Y-a-b*X)^2), with an intercept. Do not round intermediate calculations. "
    rows = [list(row) for row in observations]
    rng.shuffle(rows)
    save(
        "weighted_least_squares",
        {"Observations": [["X", "Y", "Weight"], *rows]},
        [
            (prefix + "What is slope b? Answer to six decimal places.", fixed(slope)),
            (
                prefix + "What is intercept a? Answer to six decimal places.",
                fixed(intercept),
            ),
            (
                prefix
                + "What Y does the fitted line predict at X=3? Answer to six decimal places.",
                fixed(intercept + 3 * slope),
            ),
            (
                prefix
                + "What is the minimized weighted sum of squared residuals? Answer to six decimal places.",
                fixed(residual),
            ),
            (
                prefix
                + "What is the Weight-weighted mean of observed Y? Answer to six decimal places.",
                fixed(Fraction(sy, w)),
            ),
            (
                prefix
                + "What is weighted R-squared, defined as 1 minus weighted residual sum of squares divided by sum(Weight*(Y-weighted_mean_Y)^2)? Answer to six decimal places.",
                fixed(1 - residual / (syy - Fraction(sy * sy, w))),
            ),
        ],
    )

    states = ["Idle", "Active", "Done"]
    transition = [
        [Fraction(1, 2), Fraction(1, 2), Fraction(0)],
        [Fraction(1, 4), Fraction(1, 2), Fraction(1, 4)],
        [Fraction(0), Fraction(0), Fraction(1)],
    ]
    distributions = [[Fraction(1), Fraction(0), Fraction(0)]]
    for _ in range(4):
        distributions.append(
            [
                sum(distributions[-1][i] * transition[i][j] for i in range(3))
                for j in range(3)
            ]
        )
    rows = [
        [states[i], states[j], float(transition[i][j])]
        for i in range(3)
        for j in range(3)
    ]
    rng.shuffle(rows)
    prefix = "Use Transition probabilities only. This is a time-homogeneous Markov chain with From state and To state; at step 0 it is certainly Idle. Probabilities are fractions from 0 to 1. "
    save(
        "markov_chain",
        {
            "Transition probabilities": [
                ["From state", "To state", "Probability"],
                *rows,
            ]
        },
        [
            (
                prefix
                + "What is the probability of Done at step 2? Answer to six decimal places.",
                fixed(distributions[2][2]),
            ),
            (
                prefix
                + "What is the probability of Active at step 3? Answer to six decimal places.",
                fixed(distributions[3][1]),
            ),
            (
                prefix
                + "What is the probability of Idle at step 4? Answer to six decimal places.",
                fixed(distributions[4][0]),
            ),
            (
                prefix
                + "What is the expected number of visits to Active at steps 1, 2 and 3 combined? Answer to six decimal places.",
                fixed(sum(distributions[t][1] for t in range(1, 4))),
            ),
            (
                prefix
                + "What is the probability of having reached Done by step 4 inclusive? Done is absorbing. Answer to six decimal places.",
                fixed(distributions[4][2]),
            ),
            (
                prefix
                + "A one-time reward at step 3 is 0 for Idle, 2 for Active, and 7 for Done. What is its expected value? Answer to six decimal places.",
                fixed(2 * distributions[3][1] + 7 * distributions[3][2]),
            ),
        ],
    )
    (root / "queries.json").write_text(
        json.dumps(queries, ensure_ascii=False, indent=2)
    )
    with (root / "references.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["id", "answer"])
        writer.writeheader()
        writer.writerows(references)
    paths = [
        root / "queries.json",
        root / "references.csv",
        *sorted(files.glob("*.xlsx")),
    ]
    (root / "manifest.json").write_text(
        json.dumps(
            {
                "seed": seed,
                "questions": len(queries),
                "families": sorted({q["family"] for q in queries}),
                "hashes": {
                    str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest()
                    for p in paths
                },
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, required=True)
    generate(parser.parse_args().root)
