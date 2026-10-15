"""Unseen critical paths, tied rank correlation and relational division, frozen once."""

import argparse
import csv
import hashlib
import itertools
import json
import random
import sqlite3
from contextlib import closing
from decimal import Decimal
from fractions import Fraction
from pathlib import Path
from typing import Any

import openpyxl


def generate(root: Path, seed: int = 685193) -> None:
    """Independent exhaustive-path, pairwise-rank and SQL oracles; no solver imports."""
    if root.exists():
        raise ValueError("Preserve frozen compositional data")
    files = root / "xlsx_files"
    files.mkdir(parents=True)
    rng = random.Random(seed)
    queries: list[dict[str, Any]] = []
    references: list[dict[str, str]] = []

    def fixed(value: Fraction, decimals: int = 6) -> str:
        return f"{Decimal(value.numerator) / Decimal(value.denominator):.{decimals}f}"

    def save(
        family: str, tables: dict[str, list[list[Any]]], cases: list[tuple[str, str]]
    ) -> None:
        book = openpyxl.Workbook()
        book.remove(book.active)
        for name, rows in tables.items():
            sheet = book.create_sheet(name)
            first, left = rng.choice([37, 419, 907]), rng.choice([2, 6, 11])
            order = list(range(len(rows[0])))
            rng.shuffle(order)
            for row, values in enumerate(rows, first):
                for column, original in enumerate(order, left):
                    sheet.cell(row, column, values[original])
        book.create_sheet("Excluded archive").append(
            ["Unrelated historical total", 765432]
        )
        filename = f"work_{rng.randrange(100000, 999999)}.xlsx"
        book.save(files / filename)
        book.close()
        for question, answer in cases:
            key = f"compositional-{len(queries) + 1:03d}"
            queries.append(
                {"id": key, "file_name": filename, "family": family, "query": question}
            )
            references.append({"id": key, "answer": answer})

    durations = {
        "Survey": 4,
        "Permit": 3,
        "Design": 6,
        "Prep": 2,
        "Review": 5,
        "Launch": 1,
    }
    edges = {
        ("Survey", "Design"),
        ("Permit", "Design"),
        ("Survey", "Prep"),
        ("Design", "Review"),
        ("Prep", "Launch"),
        ("Review", "Launch"),
    }
    roots = set(durations) - {right for left, right in edges}
    terminals = set(durations) - {left for left, right in edges}
    paths = [
        path
        for length in range(1, len(durations) + 1)
        for path in itertools.permutations(durations, length)
        if all(pair in edges for pair in itertools.pairwise(path))
    ]
    complete = [path for path in paths if path[0] in roots and path[-1] in terminals]
    duration = max(sum(durations[node] for node in path) for path in complete)
    critical = [
        path for path in complete if sum(durations[node] for node in path) == duration
    ]
    assert len(critical) == 1
    earliest = {
        node: max(
            sum(durations[item] for item in path[:-1])
            for path in paths
            if path[0] in roots and path[-1] == node
        )
        for node in durations
    }
    latest = {
        node: duration
        - max(
            sum(durations[item] for item in path)
            for path in paths
            if path[0] == node and path[-1] in terminals
        )
        for node in durations
    }
    tasks = [[task, days] for task, days in durations.items()]
    dependencies = [list(edge) for edge in sorted(edges)]
    rng.shuffle(tasks)
    rng.shuffle(dependencies)
    prefix = "Use Tasks and Dependencies only. Tasks may run in parallel without resource limits; each starts only after all its predecessors finish. Time starts at 0. Durations are days. "
    save(
        "critical_path_dag",
        {
            "Tasks": [["Task", "Duration"], *tasks],
            "Dependencies": [["Predecessor", "Successor"], *dependencies],
        },
        [
            (
                prefix
                + "What is the earliest project completion time? Answer only the integer.",
                str(duration),
            ),
            (
                prefix
                + "List the unique critical path from start to finish, as Task names separated by a comma and one space.",
                ", ".join(critical[0]),
            ),
            (
                prefix
                + "What is the earliest start time of Design? Answer only the integer.",
                str(earliest["Design"]),
            ),
            (
                prefix
                + "What is the earliest finish time of Review? Answer only the integer.",
                str(earliest["Review"] + durations["Review"]),
            ),
            (
                prefix
                + "What is total slack of Permit, defined as latest start without delaying the project minus earliest start? Answer only the integer.",
                str(latest["Permit"] - earliest["Permit"]),
            ),
            (
                prefix
                + "What is total slack of Prep using the same definition? Answer only the integer.",
                str(latest["Prep"] - earliest["Prep"]),
            ),
        ],
    )

    people = ["Nora", "Pax", "Quin", "Rui", "Sora", "Tavi"]
    x = [10, 20, 20, 40, 50, 60]
    y = [7, 4, 4, 9, 12, 10]

    def ranks(values: list[int]) -> list[Fraction]:
        return [
            Fraction(1 + sum(other < value for other in values))
            + Fraction(sum(other == value for other in values) - 1, 2)
            for value in values
        ]

    rx, ry = ranks(x), ranks(y)
    mean = Fraction(len(people) + 1, 2)
    covariance_numerator = sum(
        (a - mean) * (b - mean) for a, b in zip(rx, ry, strict=True)
    )
    vx, vy = (sum((rank - mean) ** 2 for rank in vector) for vector in [rx, ry])
    assert vx == vy
    largest_gap = max(range(len(people)), key=lambda i: abs(rx[i] - ry[i]))
    rows = [[name, a, b] for name, a, b in zip(people, x, y, strict=True)]
    rng.shuffle(rows)
    prefix = "Use Paired measurements only. Rank ascending starting at 1, giving tied values the average of their occupied ranks. Spearman correlation means Pearson correlation of these rank vectors, including ties. "
    save(
        "tied_rank_correlation",
        {"Paired measurements": [["Person", "X", "Y"], *rows]},
        [
            (
                prefix
                + "What is Spearman correlation of X and Y? Answer to six decimal places.",
                fixed(covariance_numerator / vx),
            ),
            (
                prefix + "What is the X rank of Pax? Answer to six decimal places.",
                fixed(rx[people.index("Pax")]),
            ),
            (
                prefix + "What is the Y rank of Nora? Answer to six decimal places.",
                fixed(ry[people.index("Nora")]),
            ),
            (
                prefix
                + "What is sum((X_rank-Y_rank)^2) over all people? Answer only the integer.",
                str(int(sum((a - b) ** 2 for a, b in zip(rx, ry, strict=True)))),
            ),
            (
                prefix
                + "What is population covariance of the two rank vectors, dividing by the number of people? Answer to six decimal places.",
                fixed(covariance_numerator / len(people)),
            ),
            (
                prefix
                + "Which Person has the largest absolute difference between X rank and Y rank? Answer only the name.",
                people[largest_gap],
            ),
        ],
    )

    required = ["Ceramics", "Glass", "Metal"]
    purchases = [
        ["ACME Inc.", "Ceramics", 10],
        ["ACME Inc.", "Glass", 20],
        ["ACME Inc.", "Metal", 30],
        ["ACME Inc.", "Glass", 5],
        ["Blue Co.", "Ceramics", 12],
        ["Blue Co.", "Glass", 3],
        ["Blue Co.", "Metal", 1],
        ["Delta Works", "Ceramics", 20],
        ["Delta Works", "Glass", 10],
        ["Echo Ltd.", "Ceramics", 4],
        ["Echo Ltd.", "Glass", 5],
        ["Echo Ltd.", "Metal", 6],
        ["Echo Ltd.", "Wood", 100],
        ["Fox Group", "Wood", 2],
    ]
    with closing(sqlite3.connect(":memory:")) as database:
        database.execute("CREATE TABLE required(category TEXT PRIMARY KEY)")
        database.execute(
            "CREATE TABLE purchases(customer TEXT, category TEXT, amount INTEGER)"
        )
        database.executemany(
            "INSERT INTO required VALUES (?)", [(value,) for value in required]
        )
        database.executemany("INSERT INTO purchases VALUES (?,?,?)", purchases)
        qualifying = database.execute(
            "SELECT p.customer, SUM(p.amount) FROM purchases p JOIN required r ON p.category=r.category GROUP BY p.customer HAVING COUNT(DISTINCT p.category)=(SELECT COUNT(*) FROM required) ORDER BY p.customer"
        ).fetchall()
        all_customers = database.execute(
            "SELECT COUNT(DISTINCT customer) FROM purchases"
        ).fetchone()[0]
    total = sum(amount for name, amount in qualifying)
    rng.shuffle(purchases)
    prefix = "Use Required categories and Purchases only. A Customer qualifies if they purchased every distinct required Category at least once. Duplicate purchases do not add categories. Spend counts Amount in required categories only; ignore other categories. "
    save(
        "relational_division",
        {
            "Required categories": [["Category"], *[[value] for value in required]],
            "Purchases": [["Customer", "Category", "Amount"], *purchases],
        },
        [
            (
                prefix + "How many Customers qualify? Answer only the integer.",
                str(len(qualifying)),
            ),
            (
                prefix
                + "List qualifying Customer names alphabetically, preserving punctuation, separated by a comma and one space.",
                ", ".join(name for name, amount in qualifying),
            ),
            (
                prefix
                + "What is total required-category Spend of qualifying Customers? Answer only the integer.",
                str(total),
            ),
            (
                prefix
                + "Which qualifying Customer has the largest required-category Spend? Answer only the name.",
                max(qualifying, key=lambda item: item[1])[0],
            ),
            (
                prefix
                + "What is the arithmetic mean required-category Spend per qualifying Customer? Answer to two decimal places.",
                fixed(Fraction(total, len(qualifying)), 2),
            ),
            (
                prefix
                + "How many distinct Customers do not qualify? Answer only the integer.",
                str(all_customers - len(qualifying)),
            ),
        ],
    )
    (root / "queries.json").write_text(
        json.dumps(queries, ensure_ascii=False, indent=2)
    )
    with (root / "references.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["id", "answer"])
        writer.writeheader()
        writer.writerows(references)
    paths_to_hash = [
        root / "queries.json",
        root / "references.csv",
        *sorted(files.glob("*.xlsx")),
    ]
    (root / "manifest.json").write_text(
        json.dumps(
            {
                "seed": seed,
                "questions": len(queries),
                "families": sorted({query["family"] for query in queries}),
                "hashes": {
                    str(path.relative_to(root)): hashlib.sha256(
                        path.read_bytes()
                    ).hexdigest()
                    for path in paths_to_hash
                },
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, required=True)
    generate(parser.parse_args().root)
