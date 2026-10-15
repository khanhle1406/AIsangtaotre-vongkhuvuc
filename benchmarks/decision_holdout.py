"""Freeze effective-dated joins, set cover and Bayesian evidence with independent oracles."""

import argparse
import csv
import datetime as dt
import hashlib
import itertools
import json
import math
import random
import sqlite3
from contextlib import closing
from decimal import Decimal
from fractions import Fraction
from pathlib import Path
from typing import Any

import openpyxl


def covers(
    costs: dict[str, int],
    coverage: dict[str, set[str]],
    required: set[str],
    excluded: frozenset[str] = frozenset(),
) -> list[tuple[int, tuple[str, ...]]]:
    """Enumerate the entire feasible power set; no solver/runtime imports."""
    names = sorted(costs.keys() - excluded)
    return sorted(
        (sum(costs[name] for name in subset), subset)
        for length in range(len(names) + 1)
        for subset in itertools.combinations(names, length)
        if required <= set().union(*(coverage[name] for name in subset))
    )


def posterior(
    priors: dict[str, Fraction],
    likelihoods: dict[str, Fraction],
    outcomes: list[bool],
) -> dict[str, Fraction]:
    """Exact joint-probability enumeration of conditionally independent evidence."""
    joint = {
        name: prior
        * math.prod(
            likelihoods[name] if positive else 1 - likelihoods[name]
            for positive in outcomes
        )
        for name, prior in priors.items()
    }
    return {name: value / sum(joint.values()) for name, value in joint.items()}


def effective_amounts(
    transactions: list[tuple[str, dt.datetime, int]],
    rates: list[tuple[str, dt.datetime, dt.datetime, Fraction]],
) -> list[Fraction]:
    """SQLite performs the interval join; exact fractions compute converted amounts."""
    with closing(sqlite3.connect(":memory:")) as connection:
        connection.execute(
            "CREATE TABLE tx (id INTEGER, currency TEXT, day TEXT, amount INTEGER)"
        )
        connection.execute(
            "CREATE TABLE fx (currency TEXT, start TEXT, end TEXT, numerator INTEGER, denominator INTEGER)"
        )
        connection.executemany(
            "INSERT INTO tx VALUES (?,?,?,?)",
            [
                (i, c, day.isoformat(), amount)
                for i, (c, day, amount) in enumerate(transactions)
            ],
        )
        connection.executemany(
            "INSERT INTO fx VALUES (?,?,?,?,?)",
            [
                (
                    c,
                    start.isoformat(),
                    end.isoformat(),
                    rate.numerator,
                    rate.denominator,
                )
                for c, start, end, rate in rates
            ],
        )
        joined = connection.execute(
            "SELECT tx.id, tx.amount, fx.numerator, fx.denominator FROM tx JOIN fx ON tx.currency=fx.currency AND tx.day>=fx.start AND tx.day<fx.end ORDER BY tx.id"
        ).fetchall()
    if [row[0] for row in joined] != list(range(len(transactions))):
        raise ValueError(
            "Oracle requires exactly one applicable interval per transaction"
        )
    return [
        Fraction(amount * numerator, denominator)
        for _, amount, numerator, denominator in joined
    ]


def generate(root: Path, seed: int = 914783) -> dict[str, Any]:
    """Write all questions/references and hashes before any model evaluation."""
    if root.exists():
        raise ValueError("Preserve frozen decision data")
    files = root / "xlsx_files"
    files.mkdir(parents=True)
    rng = random.Random(seed)
    queries: list[dict[str, str]] = []
    references: list[dict[str, str]] = []

    def fixed(value: Fraction, decimals: int = 6) -> str:
        return f"{Decimal(value.numerator) / Decimal(value.denominator):.{decimals}f}"

    def save(
        family: str, tables: dict[str, list[list[Any]]], cases: list[tuple[str, str]]
    ) -> None:
        book = openpyxl.Workbook()
        book.remove(book.active)
        for name, rows in tables.items():
            sheet = book.create_sheet(name)
            order = list(range(len(rows[0])))
            rng.shuffle(order)
            first, left = rng.choice([41, 409, 803]), rng.choice([3, 7, 12])
            for row, values in enumerate(rows, first):
                for column, original in enumerate(order, left):
                    cell = sheet.cell(row, column, values[original])
                    if isinstance(values[original], dt.datetime):
                        cell.number_format = "yyyy-mm-dd"
        book.create_sheet("Excluded archive").append(
            ["Historical unrelated total", 876543]
        )
        filename = f"decision_{rng.randrange(100000, 999999)}.xlsx"
        book.save(files / filename)
        book.close()
        for query, answer in cases:
            key = f"decision-{len(queries) + 1:03d}"
            queries.append(
                {"id": key, "file_name": filename, "family": family, "query": query}
            )
            references.append({"id": key, "answer": answer})

    costs = {"Aster": 8, "Birch": 5, "Cedar": 5, "Dune": 4, "Elm": 7, "Fir": 9}
    coverage = {
        "Aster": {"R", "S"},
        "Birch": {"R"},
        "Cedar": {"S", "T"},
        "Dune": {"T", "U"},
        "Elm": {"U"},
        "Fir": {"S", "U"},
    }
    required = {"R", "S", "T", "U"}
    feasible = covers(costs, coverage, required)
    optimum = feasible[0][0]
    best = [subset for cost, subset in feasible if cost == optimum]
    assert len(best) == 1
    prefix = "Use Providers, Coverage and Required only. Select each provider at most once and pay its Cost once; every required Area must be covered by at least one selected provider. Overlap is allowed and does not add a charge. "
    save(
        "minimum_weight_set_cover",
        {
            "Providers": [
                ["Provider", "Cost"],
                *[list(item) for item in costs.items()],
            ],
            "Coverage": [
                ["Provider", "Area"],
                *[[name, area] for name in costs for area in sorted(coverage[name])],
            ],
            "Required": [["Area"], *[[area] for area in sorted(required)]],
        },
        [
            (
                prefix
                + "What is the minimum total Cost covering all required Areas? Answer only the integer.",
                str(optimum),
            ),
            (
                prefix
                + "List the providers in the unique minimum-cost solution alphabetically, separated by a comma and one space.",
                ", ".join(best[0]),
            ),
            (
                prefix
                + "If Aster is unavailable, what is the minimum total Cost? Answer only the integer.",
                str(covers(costs, coverage, required, frozenset({"Aster"}))[0][0]),
            ),
            (
                prefix
                + "If only R and T require covering, what is the minimum total Cost? Answer only the integer.",
                str(covers(costs, coverage, {"R", "T"})[0][0]),
            ),
            (
                prefix
                + "Ignoring cost, what is the minimum number of providers covering all required Areas? Answer only the integer.",
                str(min(len(subset) for _, subset in feasible)),
            ),
            (
                prefix
                + "How many distinct provider sets attain the minimum total Cost for all required Areas? Answer only the integer.",
                str(len(best)),
            ),
        ],
    )

    priors = {
        "Amber source": Fraction(1, 2),
        "Blue source": Fraction(3, 10),
        "Clear source": Fraction(1, 5),
    }
    likelihoods = {
        "Amber source": Fraction(4, 5),
        "Blue source": Fraction(2, 5),
        "Clear source": Fraction(1, 10),
    }
    prefix = "Use Priors and Likelihoods only. Prior is the hypothesis probability; Positive is P(positive observation | hypothesis), and negative has probability 1-Positive. Observations are independent conditional on the hypothesis. "
    save(
        "conditional_bayesian_evidence",
        {
            "Priors": [
                ["Hypothesis", "Prior"],
                *[[name, float(value)] for name, value in priors.items()],
            ],
            "Likelihoods": [
                ["Hypothesis", "Positive"],
                *[[name, float(value)] for name, value in likelihoods.items()],
            ],
        },
        [
            (
                prefix
                + "What is P(Amber source | one positive observation)? Answer to six decimal places.",
                fixed(posterior(priors, likelihoods, [True])["Amber source"]),
            ),
            (
                prefix
                + "What is P(Blue source | one positive and one negative observation)? Answer to six decimal places.",
                fixed(posterior(priors, likelihoods, [True, False])["Blue source"]),
            ),
            (
                prefix
                + "What is P(Clear source | two negative observations)? Answer to six decimal places.",
                fixed(posterior(priors, likelihoods, [False, False])["Clear source"]),
            ),
            (
                prefix
                + "What is the marginal probability of one positive observation? Answer to six decimal places.",
                fixed(sum(priors[name] * likelihoods[name] for name in priors)),
            ),
            (
                prefix
                + "What is the marginal probability of two positive observations? Answer to six decimal places.",
                fixed(sum(priors[name] * likelihoods[name] ** 2 for name in priors)),
            ),
            (
                prefix
                + "Which hypothesis has the highest posterior probability after two negative observations? Answer with only its full name.",
                max(
                    posterior(priors, likelihoods, [False, False]),
                    key=posterior(priors, likelihoods, [False, False]).get,
                ),
            ),
        ],
    )

    def day(month: int, date: int) -> dt.datetime:
        return dt.datetime(2031, month, date)

    rates = [
        ("EUR", day(1, 1), day(1, 16), Fraction(11, 10)),
        ("EUR", day(1, 16), day(2, 1), Fraction(5, 4)),
        ("GBP", day(1, 1), day(1, 16), Fraction(13, 10)),
        ("GBP", day(1, 16), day(2, 1), Fraction(3, 2)),
    ]
    transactions = [
        ("EUR", day(1, 2), 10),
        ("EUR", day(1, 16), 20),
        ("EUR", day(1, 31), 12),
        ("GBP", day(1, 15), 30),
        ("GBP", day(1, 16), 8),
        ("GBP", day(1, 20), 6),
    ]
    amounts = effective_amounts(transactions, rates)
    prefix = "Use Transactions and Effective rates only. Match Currency and the interval Start <= Event date < End; Rate is USD per unit of source currency. Convert each Source amount at its own applicable rate before aggregation. "
    save(
        "effective_dated_interval_join",
        {
            "Transactions": [
                ["Currency", "Event date", "Source amount"],
                *[list(row) for row in transactions],
            ],
            "Effective rates": [
                ["Currency", "Start", "End", "Rate"],
                *[[c, start, end, float(rate)] for c, start, end, rate in rates],
            ],
        },
        [
            (
                prefix + "What is total USD amount? Answer to two decimal places.",
                fixed(sum(amounts), 2),
            ),
            (
                prefix
                + "What is total USD amount for EUR? Answer to two decimal places.",
                fixed(
                    sum(
                        amount
                        for amount, row in zip(amounts, transactions, strict=True)
                        if row[0] == "EUR"
                    ),
                    2,
                ),
            ),
            (
                prefix
                + "What is total USD amount for GBP? Answer to two decimal places.",
                fixed(
                    sum(
                        amount
                        for amount, row in zip(amounts, transactions, strict=True)
                        if row[0] == "GBP"
                    ),
                    2,
                ),
            ),
            (
                prefix
                + "What is total USD amount for events before 2031-01-16? Answer to two decimal places.",
                fixed(
                    sum(
                        amount
                        for amount, row in zip(amounts, transactions, strict=True)
                        if row[1] < day(1, 16)
                    ),
                    2,
                ),
            ),
            (
                prefix
                + "What is total USD amount for events on 2031-01-16? Answer to two decimal places.",
                fixed(
                    sum(
                        amount
                        for amount, row in zip(amounts, transactions, strict=True)
                        if row[1] == day(1, 16)
                    ),
                    2,
                ),
            ),
            (
                prefix
                + "How many transactions use a Rate strictly greater than 1.25? Answer only the integer.",
                str(
                    sum(
                        amount / Fraction(row[2]) > Fraction(5, 4)
                        for amount, row in zip(amounts, transactions, strict=True)
                    )
                ),
            ),
        ],
    )
    (root / "queries.json").write_text(json.dumps(queries, indent=2))
    with (root / "references.csv").open("w", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=["id", "answer"])
        writer.writeheader()
        writer.writerows(references)
    paths = [
        root / "queries.json",
        root / "references.csv",
        *sorted(files.glob("*.xlsx")),
    ]
    manifest = {
        "seed": seed,
        "questions": len(queries),
        "families": sorted({q["family"] for q in queries}),
        "oracle": "Power-set enumeration; exact joint probabilities; SQLite interval join plus Fraction arithmetic",
        "hashes": {
            str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in paths
        },
    }
    (root / "manifest.json").write_text(json.dumps(manifest, indent=2))
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, required=True)
    generate(parser.parse_args().root)
