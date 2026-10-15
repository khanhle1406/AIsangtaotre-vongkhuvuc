"""Fresh validation families and independent oracles, frozen before live evaluation."""

import csv
import datetime as dt
import json
import random
from pathlib import Path
from typing import Any

import openpyxl


def generate(root: Path, seed: int = 153791) -> tuple[Path, Path]:
    if root.exists():
        raise ValueError("Preserve holdout: choose a new directory")
    files = root / "xlsx_files"
    files.mkdir(parents=True)
    rng = random.Random(seed)
    queries: list[dict[str, Any]] = []
    references: list[dict[str, str]] = []

    def add(file: str, family: str, query: str, answer: str) -> None:
        key = f"unseen-{len(queries) + 1:03d}"
        queries.append({"id": key, "file_name": file, "query": query, "family": family})
        references.append({"id": key, "answer": answer})

    # Long-form events: repeated entities, filtering by period, signed adjustments.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Current ledger"
    ws.append(["Date", "Region", "Product", "Net units", "Unit price"])
    events = []
    for i in range(24):
        date = dt.datetime(2025, 1 + i % 4, 1 + i % 20)
        region = ["North", "South", "West"][i % 3]
        product = ["Orchid", "Cedar"][i % 2]
        units = rng.randrange(-8, 35)
        price = [7, 11][i % 2]
        events.append((date, region, product, units, price))
        ws.append(events[-1])
    notes = wb.create_sheet("Historical totals")
    notes.append(["Region", "Units"])
    notes.append(["North", 999999])
    file = "events_renamed.xlsx"
    wb.save(files / file)
    wb.close()
    add(
        file,
        "long_form_events",
        "In Current ledger, what are total net units for North? Include negative adjustments. Answer with only the integer.",
        str(sum(e[3] for e in events if e[1] == "North")),
    )
    add(
        file,
        "long_form_events",
        "In Current ledger, how many transactions had negative net units? Answer with only the integer.",
        str(sum(e[3] < 0 for e in events)),
    )
    add(
        file,
        "long_form_events",
        "In Current ledger, calculate total net units dated in February 2025. Answer with only the integer.",
        str(sum(e[3] for e in events if e[0].month == 2)),
    )
    add(
        file,
        "long_form_events",
        "In Current ledger, calculate total revenue for Orchid: sum each transaction's net units times its unit price. Answer with only the integer.",
        str(sum(e[3] * e[4] for e in events if e[2] == "Orchid")),
    )
    totals = {
        r: sum(e[3] for e in events if e[1] == r) for r in ("North", "South", "West")
    }
    winner = min(totals, key=lambda r: (-totals[r], r))
    add(
        file,
        "long_form_events",
        "Which region has the greatest total net units in Current ledger? Break any tie alphabetically. Answer with only the region name.",
        winner,
    )
    add(
        file,
        "long_form_events",
        "In Current ledger, what is the combined net units for South and West dated before 2025-03-01? Answer with only the integer.",
        str(
            sum(
                e[3]
                for e in events
                if e[1] in {"South", "West"} and e[0] < dt.datetime(2025, 3, 1)
            )
        ),
    )

    # Cross-sheet keyed join: sheets deliberately use different row/column orders.
    wb = openpyxl.Workbook()
    budget = wb.active
    budget.title = "Approved"
    budget.append(["Project code", "Budget USD"])
    codes = ["R-41", "K-08", "M-72", "P-19"]
    budgets = {code: rng.randrange(80, 150) * 100 for code in codes}
    actuals = {code: budgets[code] + rng.randrange(-25, 50) * 100 for code in codes}
    for code in codes:
        budget.append([code, budgets[code]])
    actual = wb.create_sheet("Actual spending")
    actual.append(["Spent USD", "Project code"])
    order = codes[:]
    rng.shuffle(order)
    for code in order:
        actual.append([actuals[code], code])
    summary = wb.create_sheet("Checks")
    summary["A1"] = "Calculated spending sum"
    summary["B1"] = "=SUM('Actual spending'!A2:A5)"
    file = "joined_projects.xlsx"
    wb.save(files / file)
    wb.close()
    deltas = {code: actuals[code] - budgets[code] for code in codes}
    highest = min(codes, key=lambda code: (-deltas[code], code))
    add(
        file,
        "keyed_cross_sheet_join",
        "Using project codes to join Approved and Actual spending, which project has the greatest actual spending minus budget? Break ties alphabetically. Answer with only the project code.",
        highest,
    )
    add(
        file,
        "keyed_cross_sheet_join",
        "How many projects spent more than their approved budget? Join by project code. Answer with only the integer.",
        str(sum(d > 0 for d in deltas.values())),
    )
    add(
        file,
        "keyed_cross_sheet_join",
        "What is actual spending minus approved budget for project K-08? Join by project code. Answer with only the integer.",
        str(deltas["K-08"]),
    )
    add(
        file,
        "keyed_cross_sheet_join",
        "Calculate the combined actual spending for projects R-41 and M-72. Answer with only the integer.",
        str(actuals["R-41"] + actuals["M-72"]),
    )
    add(
        file,
        "keyed_cross_sheet_join",
        "What value does the formula in Checks!B1 calculate? Answer with only the integer.",
        str(sum(actuals.values())),
    )
    add(
        file,
        "keyed_cross_sheet_join",
        "Calculate total actual spending divided by total approved budget, expressed as a percentage. Answer rounded to two decimal places, without the percent sign.",
        f"{100 * sum(actuals.values()) / sum(budgets.values()):.2f}",
    )

    # String dates and an explicit competing reference; no Admission-age default.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Audit people"
    ws.append(
        [
            "Surname",
            "DOB (ISO text)",
            "Admission date",
            "Audit date (ISO text)",
            "Score",
        ]
    )
    people = [
        ("Iris", "1988-03-04", "2011-01-02", "2020-07-09", rng.randrange(30, 100)),
        ("Beryl", "1991-05-06", "2025-01-02", "2026-07-09", rng.randrange(30, 100)),
        ("Cora", "1995-07-08", "2025-01-02", "2024-07-09", rng.randrange(30, 100)),
        ("Dara", "2001-09-10", "2025-01-02", "2023-07-09", rng.randrange(30, 100)),
    ]
    for person in people:
        ws.append(person)
    file = "audit_records.xlsx"
    wb.save(files / file)
    wb.close()
    ranked = sorted(
        people,
        key=lambda p: (dt.date.fromisoformat(p[3]) - dt.date.fromisoformat(p[1])).days,
        reverse=True,
    )
    add(
        file,
        "string_dates_reference",
        "Who was oldest on their own audit date, using DOB and Audit date rather than Admission date? Answer with only the surname.",
        ranked[0][0],
    )
    add(
        file,
        "string_dates_reference",
        "What score did the second-oldest person on their own audit date receive? Use DOB and Audit date. Answer with only the integer.",
        str(ranked[1][4]),
    )
    add(
        file,
        "string_dates_reference",
        "Among the two youngest people on their own audit dates, what was the highest score? Answer with only the integer.",
        str(max(p[4] for p in ranked[-2:])),
    )
    add(
        file,
        "string_dates_reference",
        "What is Beryl's audit date? Answer with only the ISO date.",
        people[1][3],
    )
    add(
        file,
        "string_dates_reference",
        "Was Iris older than Beryl on their respective audit dates? Answer with only Yes or No.",
        "Yes"
        if (
            dt.date.fromisoformat(people[0][3]) - dt.date.fromisoformat(people[0][1])
        ).days
        > (
            dt.date.fromisoformat(people[1][3]) - dt.date.fromisoformat(people[1][1])
        ).days
        else "No",
    )
    add(
        file,
        "string_dates_reference",
        "How many days elapsed from Cora's DOB to her Audit date? Answer with only the integer.",
        str(
            (
                dt.date.fromisoformat(people[2][3])
                - dt.date.fromisoformat(people[2][1])
            ).days
        ),
    )
    query_path, reference_path = root / "queries.json", root / "references.csv"
    query_path.write_text(json.dumps(queries, indent=2), encoding="utf8")
    with reference_path.open("w", newline="", encoding="utf8") as stream:
        writer = csv.DictWriter(stream, fieldnames=["id", "answer"])
        writer.writeheader()
        writer.writerows(references)
    return query_path, reference_path
