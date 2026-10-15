"""Freeze new panel, continuous-allocation and working-calendar families before inference."""

import argparse
import bisect
import csv
import datetime as dt
import hashlib
import itertools
import json
import random
from decimal import Decimal
from fractions import Fraction
from pathlib import Path
from typing import Any

import openpyxl


class NonUniqueOptimum(ValueError):
    """A random instance can be redrawn before its values are frozen."""


def panel_contrasts(
    records: list[tuple[str, str, str, int]],
) -> tuple[Fraction, Fraction, Fraction]:
    """Use paired unit changes; exclude units without both observations."""
    pairs: dict[tuple[str, str], dict[str, int]] = {}
    for unit, group, period, value in records:
        pairs.setdefault((unit, group), {})[period] = value
    complete = [
        (group, values)
        for (_, group), values in pairs.items()
        if set(values) == {"Pre", "Post"}
    ]
    means = {
        (group, period): sum(
            Fraction(values[period]) for cohort, values in complete if cohort == group
        )
        / sum(cohort == group for cohort, _ in complete)
        for group in ("Intervention", "Reference")
        for period in ("Pre", "Post")
    }
    changes = {
        group: sum(
            Fraction(values["Post"] - values["Pre"])
            for cohort, values in complete
            if cohort == group
        )
        / sum(cohort == group for cohort, _ in complete)
        for group in ("Intervention", "Reference")
    }
    return (
        means["Intervention", "Pre"] - means["Reference", "Pre"],
        means["Intervention", "Post"] - means["Reference", "Post"],
        changes["Intervention"] - changes["Reference"],
    )


def linear_optimum(
    constraints: list[tuple[int, int, int]], profit: tuple[int, int]
) -> tuple[Fraction, Fraction, Fraction]:
    """Enumerate exact rational vertices, including nonnegativity boundaries."""
    boundaries = [*constraints, (1, 0, 0), (0, 1, 0)]
    candidates = set()
    for (a, b, c), (d, e, f) in itertools.combinations(boundaries, 2):
        det = a * e - b * d
        if not det:
            continue
        x, y = Fraction(c * e - b * f, det), Fraction(a * f - c * d, det)
        if x >= 0 and y >= 0 and all(u * x + v * y <= cap for u, v, cap in constraints):
            candidates.add((profit[0] * x + profit[1] * y, x, y))
    best = max(value for value, _, _ in candidates)
    optimal = [(x, y, value) for value, x, y in candidates if value == best]
    if len(optimal) != 1:
        raise NonUniqueOptimum("Oracle requires a unique continuous optimum")
    return optimal[0]


def service_dates(
    open_dates: list[dt.date], cases: list[tuple[str, dt.date, int]]
) -> dict[str, dt.date]:
    """Rank-select strictly subsequent open dates, independently of model loops."""
    days = sorted(open_dates)
    return {
        name: days[bisect.bisect_right(days, opened) + lead - 1]
        for name, opened, lead in cases
    }


def fixed(value: Fraction) -> str:
    return f"{Decimal(value.numerator) / Decimal(value.denominator):.2f}"


def generate(root: Path, seed: int = 16062026) -> dict[str, Any]:
    """Randomize numerical values as well as layout; never overwrite a frozen set."""
    if root.exists():
        raise ValueError("Preserve prospective holdout")
    files = root / "xlsx_files"
    files.mkdir(parents=True)
    rng = random.Random(seed)
    queries: list[dict[str, str]] = []
    references: list[dict[str, str]] = []

    def save(
        family: str, tables: dict[str, list[list[Any]]], cases: list[tuple[str, str]]
    ) -> None:
        book = openpyxl.Workbook()
        book.remove(book.active)
        for name, rows in tables.items():
            sheet = book.create_sheet(name)
            order = list(range(len(rows[0])))
            rng.shuffle(order)
            first, left = rng.choice([37, 305, 707]), rng.choice([2, 6, 10])
            for row, values in enumerate(rows, first):
                for column, original in enumerate(order, left):
                    value = values[original]
                    if isinstance(value, dt.date):
                        value = dt.datetime.combine(value, dt.time())
                    cell = sheet.cell(row, column, value)
                    if isinstance(value, dt.datetime):
                        cell.number_format = "yyyy-mm-dd"
        book.create_sheet("Excluded archive").append(["Historical total", 987654])
        filename = (
            f"prospective_{len(queries) // 3}_{rng.randrange(100000, 999999)}.xlsx"
        )
        book.save(files / filename)
        book.close()
        for question, answer in cases:
            key = f"prospective-{len(queries) + 1:03d}"
            queries.append(
                {"id": key, "file_name": filename, "family": family, "query": question}
            )
            references.append({"id": key, "answer": answer})

    for _ in range(2):
        panel = []
        for group in ["Intervention", "Reference"]:
            for i in range(3):
                base = rng.randint(12, 65)
                panel += [
                    (f"{group}-{i}", group, "Pre", base),
                    (f"{group}-{i}", group, "Post", base + rng.randint(-9, 18)),
                ]
        panel.append(("Incomplete", "Reference", "Pre", rng.randint(70, 95)))
        rng.shuffle(panel)
        pre, post, did = panel_contrasts(panel)
        prefix = "Use Panel only. Include only units having both Pre and Post observations; use equal weight per complete unit. A group mean averages Outcome within that group and period. "
        save(
            "balanced_panel_difference_in_differences",
            {
                "Panel": [
                    ["Unit", "Group", "Period", "Outcome"],
                    *[list(row) for row in panel],
                ]
            },
            [
                (
                    prefix
                    + "What is Intervention minus Reference in the Pre period? Answer to two decimal places.",
                    fixed(pre),
                ),
                (
                    prefix
                    + "What is Intervention minus Reference in the Post period? Answer to two decimal places.",
                    fixed(post),
                ),
                (
                    prefix
                    + "Compute difference-in-differences: Intervention Post-minus-Pre change minus Reference Post-minus-Pre change. Answer to two decimal places.",
                    fixed(did),
                ),
            ],
        )

    for _ in range(2):
        for _attempt in range(100):
            constraints = [
                (rng.randint(1, 5), rng.randint(1, 5), rng.randint(30, 90))
                for _ in range(2)
            ]
            constraints.append((1, 1, max(row[2] for row in constraints) + 20))
            profits = (rng.randint(2, 11), rng.randint(2, 11))
            try:
                x, y, value = linear_optimum(constraints, profits)
                break
            except NonUniqueOptimum:
                continue  # Regenerate only nonunique mathematical optima, before inference.
        else:
            raise ValueError("Unable to draw unique optimum")
        prefix = "Use Resources and Unit profit only. Choose nonnegative REAL amounts of products X and Y, not integers. Each resource row requires X coefficient times X plus Y coefficient times Y <= Capacity. Maximize X profit times X plus Y profit times Y. The optimum is unique. "
        save(
            "continuous_two_product_resource_allocation",
            {
                "Resources": [
                    ["Resource", "X coefficient", "Y coefficient", "Capacity"],
                    *[[f"R{i + 1}", *row] for i, row in enumerate(constraints)],
                ],
                "Unit profit": [
                    ["Product", "Profit"],
                    ["X", profits[0]],
                    ["Y", profits[1]],
                ],
            },
            [
                (
                    prefix
                    + "What is the maximum profit? Answer to two decimal places.",
                    fixed(value),
                ),
                (
                    prefix
                    + "What amount of X attains the maximum profit? Answer to two decimal places.",
                    fixed(x),
                ),
                (
                    prefix
                    + "At the optimum, what unused Capacity remains for R3? Answer to two decimal places.",
                    fixed(Fraction(constraints[2][2]) - x - y),
                ),
            ],
        )

    for _ in range(2):
        base = dt.date(2041, rng.randint(1, 11), 1)
        closed = {base + dt.timedelta(days=rng.randint(2, 22)) for _ in range(3)}
        calendar = [(base + dt.timedelta(days=i)) for i in range(31)]
        open_dates = [
            day for day in calendar if day.weekday() < 5 and day not in closed
        ]
        cases = [
            (
                f"Case-{letter}",
                base + dt.timedelta(days=rng.randint(0, 8)),
                rng.randint(1, 5),
            )
            for letter in "ABC"
        ]
        due = service_dates(open_dates, cases)
        cutoff = base + dt.timedelta(days=12)
        prefix = "Use Cases and Working calendar only. Open=1 is a working date; Open=0 is closed. A case is due on its Lead-th open date STRICTLY AFTER its Opened date; never count the opened date itself. The calendar covers every required date. "
        save(
            "working_calendar_service_deadlines",
            {
                "Cases": [["Case", "Opened", "Lead"], *[list(row) for row in cases]],
                "Working calendar": [
                    ["Date", "Open"],
                    *[[day, int(day in open_dates)] for day in calendar],
                ],
            },
            [
                (
                    prefix + "On what date is Case-A due? Answer YYYY-MM-DD.",
                    due["Case-A"].isoformat(),
                ),
                (
                    prefix
                    + f"How many cases have due dates strictly before {cutoff.isoformat()}? Answer only the integer.",
                    str(sum(day < cutoff for day in due.values())),
                ),
                (
                    prefix
                    + "Sum elapsed calendar days from Opened to due date over all cases. Answer only the integer.",
                    str(sum((due[name] - opened).days for name, opened, _ in cases)),
                ),
            ],
        )
    (root / "queries.json").write_text(
        json.dumps(queries, ensure_ascii=False, indent=2)
    )
    with (root / "references.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["id", "answer"])
        writer.writeheader()
        writer.writerows(references)
    paths = [
        root / "queries.json",
        root / "references.csv",
        *sorted(files.glob("*.xlsx")),
    ]
    manifest = {
        "seed": seed,
        "queries": len(queries),
        "families": sorted({q["family"] for q in queries}),
        "numeric_values_randomized": True,
        "generator_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "independent_oracles": [
            "paired rational unit changes",
            "exact rational polygon vertices",
            "open-date rank/select",
        ],
        "hashes": {
            str(path.relative_to(root)): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in paths
        },
    }
    (root / "manifest.json").write_text(json.dumps(manifest, indent=2))
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--seed", type=int, default=16062026)
    args = parser.parse_args()
    print(json.dumps(generate(args.root, args.seed), indent=2))
