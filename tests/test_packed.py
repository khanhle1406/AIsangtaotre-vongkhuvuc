"""Shared planning must not weaken workbook ownership, arithmetic or accounting."""

import io
import json
from pathlib import Path
from typing import Any
from unittest.mock import patch

import openpyxl
import pytest

from benchmarks.batched import run_batched_stage
from benchmarks.gateway import BatchLimits, GatewayAgent
from benchmarks.packed import (
    PLANNING_PROFILES,
    AllPackedAgent,
    DeferredBooks,
    PackedAgent,
    ParallelPackedAgent,
    SampledPackedAgent,
    TenPackedAgent,
    WidePackedAgent,
    namespace_books,
)
from benchmarks.precision import ACCURACY_CONTRACT
from src.excel_parser import CellData, SheetData, WorkbookData
from src.executor import execute_program
from src.llm_agent import BudgetLedger


def test_column_letter_is_one_based_bounded_and_available_in_worker() -> None:
    from src.executor import column_letter

    indices = [1, 26, 27, 52, 702, 703, 16384]
    expected = ["A", "Z", "AA", "AZ", "ZZ", "AAA", "XFD"]
    assert [column_letter(index) for index in indices] == expected
    result = execute_program(
        "result=[column_letter(i) for i in [1,26,27,52,702,703,16384]]"
    )
    assert result.success and result.result == expected
    for invalid in [0, 16385, True, 1.0]:
        with pytest.raises(ValueError, match="1 to 16384"):
            column_letter(invalid)


def test_namespaces_preserve_formulas_metadata_and_originals() -> None:
    source = WorkbookData(
        {
            "Data": SheetData(
                "Data",
                cells={
                    "A1": CellData(3, "number"),
                    "B1": CellData(None, "uncached_formula", formula="=A1+Inputs!A1"),
                },
                merges=["C3:D3"],
                hidden_rows=[7],
            ),
            "Inputs": SheetData(
                "Inputs", cells={"A1": CellData(7, "number")}, state="hidden"
            ),
        }
    )
    other = WorkbookData(
        {"Inputs": SheetData("Inputs", cells={"A1": CellData(99, "number")})}
    )
    before = source.wire()
    packed = namespace_books({"b.xlsx": other, "a.xlsx": source})
    own = packed.owners["a.xlsx"]
    result = execute_program("result=num('W0::Data','B1')", own.wire())
    assert result.success and result.result == 10
    assert set(result.reads) == {"W0::Data!A1", "W0::Data!B1", "W0::Inputs!A1"}
    assert not execute_program("result=num('W1::Inputs','A1')", own.wire()).success
    assert own.sheets["W0::Data"].merges == ["C3:D3"]
    assert own.sheets["W0::Data"].hidden_rows == [7]
    assert own.sheets["W0::Inputs"].state == "hidden"
    assert own.sheets["W0::Data"].cells["B1"].formula == "=A1+Inputs!A1"
    assert source.wire() == before


def test_probability_complements_do_not_claim_cross_sheet_subtraction_ambiguity() -> (
    None
):
    from src.verification import scope_risks

    book = WorkbookData(
        {
            "Exposure": SheetData(
                "Exposure",
                cells={
                    "A1": CellData("Entity", "string"),
                    "B1": CellData("Weight", "string"),
                    "A2": CellData("Cedar", "string"),
                    "B2": CellData(3, "number"),
                    "A3": CellData("Birch", "string"),
                    "B3": CellData(2, "number"),
                },
            ),
            "Rates": SheetData(
                "Rates",
                cells={
                    "A1": CellData("Entity", "string"),
                    "B1": CellData("Rate", "string"),
                    "A2": CellData("Cedar", "string"),
                    "B2": CellData(0.2, "number"),
                    "A3": CellData("Birch", "string"),
                    "B3": CellData(0.1, "number"),
                },
            ),
        }
    ).wire()
    code = "a=num('Exposure','B2');b=num('Rates','B2');c=num('Exposure','B3');d=num('Rates','B3');names=[cell('Exposure','A2'),cell('Exposure','A3')];scores=[a*(1-b),c*(1-d)];result=names[scores.index(max(scores))]"
    question = "Which entity has the highest derived score?"
    execution = execute_program(code, book)
    assert execution.success and execution.result == "Cedar"
    assert not scope_risks(book, execution, question, code)
    difference = code.replace("a*(1-b),c*(1-d)", "a-b,c-d")
    assert any(
        "Cross-sheet derived rank" in risk
        for risk in scope_risks(
            book, execute_program(difference, book), question, difference
        )
    )
    tolerance = (
        code.split("scores=", 1)[0]
        + "scores=[a*b,c*d];top=max(scores);result=sorted([n for n,v in zip(names,scores) if abs(v-top)<1e-9])[0]"
    )
    tied = execute_program(tolerance, book)
    assert tied.success and tied.result == "Cedar"
    assert not scope_risks(book, tied, question, tolerance)
    signed_predicate = (
        code.split("scores=", 1)[0] + "result=names[0] if a-b>0 else names[1]"
    )
    assert any(
        "Cross-sheet derived rank" in risk
        for risk in scope_risks(
            book, execute_program(signed_predicate, book), question, signed_predicate
        )
    )


@pytest.mark.parametrize(
    "nested",
    [
        {"assumptions": [], "needs_verification": False},
        {"assumptions": ["conflict"], "needs_verification": False},
        {"needs_verification": True},
        {"needs_verification": 0},
    ],
)
def test_format_metadata_only_normalizes_identical_typed_duplicates(
    tmp_path: Path, nested: dict[str, Any]
) -> None:
    from benchmarks.structured_batch import StructuredBatchedAgent

    agent = StructuredBatchedAgent(
        api_key="test-placeholder",
        gateway_url="https://gateway.invalid/v1",
        gateway_model="ag/gemini-3-flash",
        batch_limits=BatchLimits(tmp_path / "batch.json"),
    )
    agent.structured_state.structured_expected = ["q0"]
    original = json.dumps(
        {
            "programs": {
                "q0": {
                    "code": "result=num('Changed','G408')",
                    "format": {"kind": "integer", **nested},
                    "assumptions": [],
                    "needs_verification": False,
                }
            }
        }
    )
    response = io.BytesIO(
        json.dumps(
            {"choices": [{"message": {"content": original}, "finish_reason": "stop"}]}
        ).encode()
    )
    response.headers = {"Content-Type": "application/json"}
    result = agent._read_response(response)["choices"][0]["message"]["content"]
    if nested == {"assumptions": [], "needs_verification": False}:
        program = json.loads(result)["programs"][0]
        assert program["code"] == "result=num('Changed','G408')"
        assert program["assumptions"] == [] and not program["needs_verification"]
        assert agent.structured_state.original_content == original
        assert (
            "validated-identical-format-metadata"
            in agent.structured_state.envelope_normalization
        )
    else:
        assert result == original


@pytest.mark.parametrize(
    "mode", ["normal", "foreign", "region", "independent", "large_view", "memory"]
)
@pytest.mark.parametrize(
    "agent_class",
    [
        PackedAgent,
        ParallelPackedAgent,
        SampledPackedAgent,
        AllPackedAgent,
        TenPackedAgent,
    ],
)
def test_packed_stage_owns_sources_in_every_path_and_resumes(
    tmp_path: Path, mode: str, agent_class: type[PackedAgent]
) -> None:
    files = tmp_path / "new_path"
    files.mkdir()
    for name, value in [("z.xlsx", 17), ("a.xlsx", 29)]:
        book = openpyxl.Workbook()
        book.active.title = "Same"
        book.active["C407"], book.active["D407"] = "Entity", "Units"
        book.active["C408"], book.active["D408"] = "Birch", value
        book.save(files / name)
        book.close()
    text = "What is Units for Birch? Answer only the integer."
    if mode == "independent":
        text = "How many more Units for Birch across all categories? Answer only the integer."
    # Identical query text and sheet names must remain bound to different files.
    queries = [
        {"id": str(i), "file_name": name, "query": text}
        for i, name in enumerate(["z.xlsx", "a.xlsx"])
    ]
    refs = tmp_path / "references.csv"
    refs.write_text("id,answer\n0,17\n1,29\n")
    environment = {
        "DEV_LLM_BASE_URL": "https://gateway.invalid/v1",
        "DEV_LLM_API_KEY": "test-placeholder",
        "DEV_LLM_MODEL": "ag/gemini-3-flash",
    }
    seen_limits = []

    def respond(request: Any, **kwargs: Any) -> io.BytesIO:
        payload = json.loads(request.data)
        seen_limits.append(payload["max_tokens"])
        content = payload["messages"][-1]["content"]
        if "\nQUESTIONS:\n" in content:
            questions = json.loads(
                content.split("\nQUESTIONS:\n", 1)[1].split(
                    "\n" + ACCURACY_CONTRACT, 1
                )[0]
            )
            programs = {}
            for i, question in enumerate(questions):
                sheet = question["allowed_sheets"][0]
                program = {
                    "code": f"result=num({sheet!r},'D408')",
                    "format": {"kind": "integer"},
                }
                if i == 0 and mode in {"foreign", "region"}:
                    foreign = questions[1]["allowed_sheets"][0]
                    if mode == "foreign":
                        program["code"] = f"result=num({foreign!r},'D408')"
                    else:
                        program["request_regions"] = [
                            {"sheet": foreign, "range": "D408"}
                        ]
                programs[f"q{i}"] = program
            body = {"programs": programs}
        else:
            # Only the owned workbook is in a repair/inspection/fallback context.
            data = payload["messages"][1]["content"].split("\n\nQUESTION:\n", 1)[0]
            namespaces = [
                name
                for name in ["W0::Same", "W1::Same"]
                if 'SHEET "' + name + '"' in data
            ]
            assert len(namespaces) == 1
            body = {
                "code": f"result=num({namespaces[0]!r},'D408')",
                "format": {"kind": "integer"},
            }
        result = io.BytesIO(
            json.dumps(
                {
                    "model": "gemini-3-flash",
                    "choices": [
                        {
                            "message": {"content": json.dumps(body)},
                            "finish_reason": "stop",
                        }
                    ],
                    "usage": {"total_tokens": 101},
                }
            ).encode()
        )
        result.headers = {"Content-Type": "application/json"}
        return result

    original_view = agent_class.planning_text

    def view(agent: PackedAgent, book: WorkbookData, wire: dict[str, Any]) -> str:
        from benchmarks.packed import ScopedBooks

        return (
            "x" * 800_001
            if isinstance(book, ScopedBooks)
            else original_view(agent, book, wire)
        )

    from contextlib import ExitStack

    with ExitStack() as stack:
        stack.enter_context(patch.dict("os.environ", environment))
        calls = stack.enter_context(
            patch.object(GatewayAgent, "_open_request", side_effect=respond)
        )
        if mode == "large_view":
            stack.enter_context(patch.object(agent_class, "planning_text", view))
        if mode == "memory":
            stack.enter_context(patch.object(agent_class, "file_cache_bytes", 1))
            assert isinstance(agent_class.load_group(files, queries), DeferredBooks)
        limits = BatchLimits(tmp_path / "batch.json")
        kwargs = {
            "root": tmp_path / "out",
            "files": files,
            "queries": queries,
            "references": refs,
            "limits": limits,
            "ledger": BudgetLedger(0.01),
            "agent_class": agent_class,
            "group_size": 8,
        }
        first = run_batched_stage(**kwargs)
        count = calls.call_count
        resumed = run_batched_stage(**kwargs)
        assert calls.call_count == count
        assert first["score"]["correct"] == resumed["score"]["correct"] == 2
        assert first["provider_tokens_known"] == 101 * count
        assert first["tokens"] == first["tokens_conservative_upper"] == 101 * count
        assert first["http_attempts"] == count and first["usage_missing_attempts"] == 0
        if mode == "normal":
            assert count == 1 and seen_limits == [agent_class.planning_max_tokens]
            from benchmarks.batched_cli import solve
            from benchmarks.budget import DevelopmentBudget

            query_path = tmp_path / "different queries.json"
            query_path.write_text(json.dumps(queries))
            profile = next(
                name for name, cls in PLANNING_PROFILES.items() if cls is agent_class
            )
            with (
                patch(
                    "benchmarks.batched_cli.workspace_budget",
                    return_value=DevelopmentBudget(tmp_path / "cli_ledger.json", 0),
                ),
                patch(
                    "benchmarks.batched.read_answers",
                    side_effect=AssertionError("No oracle in inference"),
                ),
            ):
                metrics = solve(
                    files,
                    query_path,
                    tmp_path / "cli result.csv",
                    planning_profile=profile,
                )
            assert metrics["answered"] == 2 and not metrics["failures"]
            assert metrics["tokens"] == 101 and metrics["planning_profile"] == profile
            with pytest.raises(ValueError, match="keyed"):
                solve(
                    files,
                    query_path,
                    tmp_path / "invalid.csv",
                    structured=False,
                    planning_profile=profile,
                )
            count = calls.call_count
        if mode in {"foreign", "region"}:
            assert count == 2 and seen_limits == [agent_class.planning_max_tokens, 2048]
        # Any input change must invalidate the checkpoint before another call.
        book = openpyxl.load_workbook(files / "z.xlsx")
        book.active["D408"] = 73
        book.save(files / "z.xlsx")
        book.close()
        with pytest.raises(ValueError, match="mismatch"):
            run_batched_stage(**kwargs)
        assert calls.call_count == count


@pytest.mark.parametrize(
    "agent_class,count,cap", [(WidePackedAgent, 16, 8192), (TenPackedAgent, 20, 6144)]
)
def test_sixteen_colliding_sheet_names_keep_ownership_and_exact_token_allocation(
    tmp_path: Path,
    agent_class: type[PackedAgent],
    count: int,
    cap: int,
) -> None:
    files = tmp_path / "books"
    files.mkdir()
    queries = []
    for i in range(count):
        book = openpyxl.Workbook()
        book.active.title = "Same"
        book.active["A1"], book.active["B1"] = "Entity", "Units"
        book.active["A2"], book.active["B2"] = "Birch", 20 + i
        filename = f"changed_name_{i}.xlsx"
        book.save(files / filename)
        book.close()
        queries.append(
            {
                "id": str(i),
                "file_name": filename,
                "query": "What is Units for Birch? Answer only the integer.",
            }
        )
    refs = tmp_path / "gold.csv"
    refs.write_text(
        "id,answer\n" + "\n".join(f"{i},{20 + i}" for i in range(count)) + "\n"
    )

    def respond(request: Any, **kwargs: Any) -> io.BytesIO:
        payload = json.loads(request.data)
        assert payload["max_tokens"] == cap
        questions = json.loads(
            payload["messages"][-1]["content"]
            .split("\nQUESTIONS:\n", 1)[1]
            .split("\n" + ACCURACY_CONTRACT, 1)[0]
        )
        assert len(questions) == agent_class.maximum_group_size
        programs = {
            f"q{i}": {
                "code": f'result=num({q["allowed_sheets"][0]!r},"B2")',
                "format": {"kind": "integer"},
            }
            for i, q in enumerate(questions)
        }
        response = io.BytesIO(
            json.dumps(
                {
                    "model": "gemini-3-flash",
                    "choices": [
                        {
                            "message": {"content": json.dumps({"programs": programs})},
                            "finish_reason": "stop",
                        }
                    ],
                    "usage": {"total_tokens": 101},
                }
            ).encode()
        )
        response.headers = {"Content-Type": "application/json"}
        return response

    with (
        patch.dict(
            "os.environ",
            {
                "DEV_LLM_BASE_URL": "https://gateway.invalid/v1",
                "DEV_LLM_API_KEY": "test-placeholder",
                "DEV_LLM_MODEL": "ag/gemini-3-flash",
            },
        ),
        patch.object(GatewayAgent, "_open_request", side_effect=respond) as calls,
    ):
        report = run_batched_stage(
            tmp_path / "out",
            files,
            queries,
            refs,
            BatchLimits(tmp_path / "batch.json"),
            BudgetLedger(0.01),
            agent_class=agent_class,
            group_size=agent_class.maximum_group_size,
        )
    assert report["score"]["correct"] == count
    expected_calls = count // agent_class.maximum_group_size
    assert report["provider_tokens_known"] == 101 * expected_calls
    assert calls.call_count == expected_calls
