"""Keep transfer/repair integration coverage independent of benchmark fixture helpers."""

import hashlib
import io
import json
from pathlib import Path
from typing import Any
from unittest.mock import patch

import openpyxl
import pytest

from benchmarks.qwen_batched import QwenBatchedAgent
from src.excel_parser import load_workbook_data
from src.llm_agent import DEFAULT_MODEL, QueryState


def client() -> QwenBatchedAgent:
    return QwenBatchedAgent(api_key="test-placeholder", pricing=(1e-7, 1e-7))


def test_exhausted_batch_is_unassessed_and_cannot_vacuously_qualify(
    tmp_path: Path,
) -> None:
    import time

    from benchmarks.transfer_qualification import qualify

    dataset = tmp_path / "sealed"
    dataset.mkdir()
    (dataset / "manifest.json").write_text(json.dumps({"hashes": {}}))
    allowance = {
        "fatal_error": None,
        "deadline_unix": time.time() + 60,
        "attempts": 200,
        "attempt_limit": 200,
        "token_upper": 1000,
        "token_limit": 2000000,
    }
    with (
        patch("benchmarks.transfer_qualification.source_manifest", return_value={}),
        patch("pathlib.Path.is_file", return_value=True),
        patch("benchmarks.transfer_qualification.BatchLimits") as limits,
        patch("benchmarks.transfer_qualification.run_stage") as baseline,
        patch("benchmarks.transfer_qualification.run_batched_stage") as candidate,
    ):
        limits.return_value.update.return_value = allowance
        report = qualify(tmp_path / "output", dataset)
    assert report["stages"] == {} and "public_1" in report["unassessed"]
    assert report["holdout_status"] == "not evaluated; failed qualification"
    assert not (dataset / "evaluated.json").exists()
    baseline.assert_not_called()
    candidate.assert_not_called()


def test_extended_gate_declares_all_254_consumed_cases_before_new_holdout(
    tmp_path: Path,
) -> None:
    import time

    from benchmarks.transfer_qualification import qualify

    heldout = tmp_path / "fresh"
    heldout.mkdir()
    (heldout / "manifest.json").write_text(json.dumps({"hashes": {}, "queries": 18}))
    with (
        patch("benchmarks.transfer_qualification.source_manifest", return_value={}),
        patch("pathlib.Path.is_file", return_value=True),
        patch("benchmarks.transfer_qualification.BatchLimits") as limits,
        patch("benchmarks.transfer_qualification.run_batched_stage") as network,
    ):
        limits.return_value.update.return_value = {
            "fatal_error": None,
            "deadline_unix": time.time() + 60,
            "attempts": 200,
            "attempt_limit": 200,
            "token_upper": 0,
            "token_limit": 2_000_000,
        }
        report = qualify(
            tmp_path / "report",
            heldout,
            extended_regression=True,
            planning_profile="parallel8",
        )
    lock = report["source_lock"]
    counts = lock["expected_case_counts"]
    assert (
        sum(
            value
            for name, value in counts.items()
            if not name.startswith(("public_", "baseline_"))
        )
        == 254
    )
    assert lock["regression_extensions"]["robustness_24"]["count"] == 24
    assert lock["configuration"]["group_size"] == 8
    assert not (heldout / "evaluated.json").exists()
    network.assert_not_called()


def test_assessment_only_evidence_cannot_enter_official_signoff(tmp_path: Path) -> None:
    from benchmarks.qwen_batch_signoff import signoff, source_manifest

    (tmp_path / "source_lock.json").write_text(
        json.dumps({"modules": source_manifest(), "assessment_only": True})
    )
    with patch("urllib.request.urlopen") as network:
        report = signoff(tmp_path)
    assert (
        report["status"] == "blocked_assessment_only" and report["new_generations"] == 0
    )
    network.assert_not_called()


def test_official_signoff_cannot_substitute_same_book_for_packed_profile(
    tmp_path: Path,
) -> None:
    from benchmarks.qwen_batch_signoff import signoff, source_manifest

    (tmp_path / "source_lock.json").write_text(
        json.dumps(
            {
                "modules": source_manifest(),
                "configuration": {"planning_profile": "independent8"},
            }
        )
    )
    with patch("urllib.request.urlopen") as network:
        report = signoff(tmp_path)
    assert report["status"] == "blocked_unsupported_planning_profile"
    assert not report["promoted"] and report["new_generations"] == 0
    network.assert_not_called()


def test_assessment_still_rejects_incomplete_qualification(tmp_path: Path) -> None:
    from benchmarks.qwen_batch_signoff import source_manifest
    from benchmarks.robustness_proof import proof

    parent = tmp_path / "candidate"
    (parent / "public_1").mkdir(parents=True)
    (parent / "source_lock.json").write_text(
        json.dumps({"modules": source_manifest(), "validation_dataset": "unused"})
    )
    (parent / "public_1/report.json").write_text(
        json.dumps({"complete": False, "score": {"correct": 19, "total": 20}})
    )
    dataset = tmp_path / "unseen"
    dataset.mkdir()
    with (
        patch("urllib.request.urlopen") as network,
        pytest.raises(ValueError, match="requires complete"),
    ):
        proof(parent, dataset, tmp_path / "output", assessment_only=True)
    assert not (dataset / "evaluated.json").exists()
    network.assert_not_called()


@pytest.mark.parametrize(
    "flag", ["review_candidate", "deduplicate_headers", "boundary_contract"]
)
def test_official_signoff_rejects_changed_review_policy_before_network(
    tmp_path: Path, flag: str
) -> None:
    from benchmarks.qwen_batch_signoff import signoff, source_manifest

    (tmp_path / "source_lock.json").write_text(
        json.dumps({"modules": source_manifest(), "configuration": {flag: True}})
    )
    (tmp_path / "public_1").mkdir()
    (tmp_path / "public_1/report.json").write_text(
        json.dumps(
            {
                "complete": True,
                "score": {"correct": 20, "total": 20},
                "config": {flag: False},
            }
        )
    )
    with patch("urllib.request.urlopen") as network:
        report = signoff(tmp_path)
    assert report["status"] == "blocked_qualification"
    assert report["dataset"] == "public_1" and report["new_generations"] == 0
    network.assert_not_called()


def test_unseen_proof_cannot_consume_data_after_partial_qualification(
    tmp_path: Path,
) -> None:
    from benchmarks.qwen_batch_signoff import source_manifest
    from benchmarks.robustness_proof import proof

    parent = tmp_path / "qualified"
    (parent / "public_1").mkdir(parents=True)
    (parent / "source_lock.json").write_text(
        json.dumps(
            {
                "modules": source_manifest(),
                "validation_dataset": str(tmp_path / "consumed"),
            }
        )
    )
    (parent / "public_1/report.json").write_text(
        json.dumps({"complete": True, "score": {"correct": 1, "total": 1}})
    )
    dataset = tmp_path / "unseen"
    dataset.mkdir()
    with (
        patch("urllib.request.urlopen") as network,
        pytest.raises(ValueError, match="Qualification failed"),
    ):
        proof(parent, dataset, tmp_path / "proof")
    network.assert_not_called()
    assert not (dataset / "evaluated.json").exists()
    assert not (tmp_path / "proof").exists()


@pytest.mark.parametrize(
    "configuration",
    [
        {"thinking": "medium", "concurrency": 8},
        {"thinking": "low", "concurrency": 4, "review_candidate": True},
        {"thinking": "low", "concurrency": 4, "deduplicate_headers": True},
        {"thinking": "low", "concurrency": 4, "boundary_contract": True},
    ],
)
def test_unseen_proof_rejects_profile_changed_since_qualification(
    tmp_path: Path, configuration: dict[str, Any]
) -> None:
    from benchmarks.qwen_batch_signoff import source_manifest
    from benchmarks.robustness_proof import proof

    parent = tmp_path / "qualified"
    (parent / "public_1").mkdir(parents=True)
    (parent / "source_lock.json").write_text(
        json.dumps(
            {
                "modules": source_manifest(),
                "configuration": configuration,
                "validation_dataset": str(tmp_path / "consumed"),
            }
        )
    )
    (parent / "public_1/report.json").write_text(
        json.dumps(
            {
                "complete": True,
                "score": {"correct": 20, "total": 20},
                "config": {"thinking": "low", "concurrency": 4},
            }
        )
    )
    dataset = tmp_path / "new"
    dataset.mkdir()
    with (
        patch("urllib.request.urlopen") as network,
        pytest.raises(ValueError, match="configuration changed"),
    ):
        proof(parent, dataset, tmp_path / "proof")
    network.assert_not_called()
    assert not (dataset / "evaluated.json").exists()


def test_independent_proof_requires_all_declared_extended_regressions(
    tmp_path: Path,
) -> None:
    from benchmarks.qwen_batch_signoff import source_manifest
    from benchmarks.robustness_proof import proof

    parent = tmp_path / "qualified"
    parent.mkdir()
    configuration = {"thinking": "low", "concurrency": 4}
    (parent / "source_lock.json").write_text(
        json.dumps(
            {
                "modules": source_manifest(),
                "configuration": configuration,
                "validation_dataset": str(tmp_path / "consumed"),
                "regression_extensions": {
                    "extra_known": {"folder": str(tmp_path / "known"), "count": 2}
                },
            }
        )
    )
    for label, count in {
        "public_1": 20,
        "public_2": 20,
        "public_3": 20,
        "regression_60": 60,
        "regression_32": 32,
        "known_18": 18,
        "synthetic_24": 24,
        "consumed_validation": 24,
        "extra_known": 2,
    }.items():
        stage = parent / label
        stage.mkdir()
        (stage / "report.json").write_text(
            json.dumps(
                {
                    "complete": label != "extra_known",
                    "score": {
                        "correct": count - 1 if label == "extra_known" else count,
                        "total": count,
                    },
                    "config": configuration,
                    "fingerprint": "stable-fixture",
                }
            )
        )
    dataset = tmp_path / "new"
    dataset.mkdir()
    with (
        patch("benchmarks.robustness_proof.fingerprint", return_value="stable-fixture"),
        patch("benchmarks.robustness_proof.validate_queries", return_value=[]),
        patch("urllib.request.urlopen") as network,
        pytest.raises(ValueError, match="Qualification failed on extra_known"),
    ):
        proof(parent, dataset, tmp_path / "proof")
    network.assert_not_called()
    assert not (dataset / "evaluated.json").exists()


@pytest.mark.parametrize("present", [False, True])
def test_official_gate_checks_extended_regressions_before_network(
    tmp_path: Path, present: bool
) -> None:
    from benchmarks.qwen_batch_signoff import signoff, source_manifest

    (tmp_path / "source_lock.json").write_text(
        json.dumps(
            {
                "modules": source_manifest(),
                "regression_extensions": {
                    "extra_known": {"folder": str(tmp_path / "data"), "count": 2}
                },
            }
        )
    )
    for label, count in {
        "public_1": 20,
        "public_2": 20,
        "public_3": 20,
        "regression_60": 60,
        "regression_32": 32,
        "known_18": 18,
        "synthetic_24": 24,
    }.items():
        stage = tmp_path / label
        stage.mkdir()
        (stage / "report.json").write_text(
            json.dumps({"complete": True, "score": {"correct": count, "total": count}})
        )
    if present:
        stage = tmp_path / "extra_known"
        stage.mkdir()
        (stage / "report.json").write_text(
            json.dumps({"complete": True, "score": {"correct": 1, "total": 2}})
        )
    with patch("urllib.request.urlopen") as network:
        report = signoff(tmp_path)
    assert (
        report["status"] == "blocked_qualification"
        and report["dataset"] == "extra_known"
    )
    assert report["new_generations"] == 0
    network.assert_not_called()


def test_new_oracles_have_independent_hand_checked_quantiles_and_counts(
    tmp_path: Path,
) -> None:
    import csv

    from benchmarks.robustness_holdout import generate

    dataset = tmp_path / "unseen"
    generate(dataset)
    queries = json.loads((dataset / "queries.json").read_text())
    assert len(queries) == 24 and len({q["id"] for q in queries}) == 24
    assert len({q["family"] for q in queries}) == 3
    with (dataset / "references.csv").open() as stream:
        answers = [row["answer"] for row in csv.DictReader(stream)]
    assert answers[:9] == [
        "27",
        "41",
        "58",
        "89",
        "44.88",
        "52.00",
        "12",
        "31",
        "73.75",
    ]
    manifest = json.loads((dataset / "manifest.json").read_text())
    assert all(
        hashlib.sha256((dataset / name).read_bytes()).hexdigest() == expected
        for name, expected in manifest["hashes"].items()
    )
    with pytest.raises(ValueError, match="Preserve"):
        generate(dataset)


def workbook(path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "New scope"
    for row, values in enumerate([["Label", "Units"], ["One", 139], ["Two", 73]], 306):
        for col, value in enumerate(values, 4):
            sheet.cell(row, col, value)
    book.save(path)
    book.close()


def test_structural_oracle_uses_verified_fifo_areas_and_leaf_expansion(
    tmp_path: Path,
) -> None:
    import csv

    from benchmarks.structural_holdout import generate

    root = tmp_path / "structural"
    generate(root)
    with (root / "references.csv").open() as stream:
        values = [row["answer"] for row in csv.DictReader(stream)]
    assert values == [
        "25",
        "Order-D",
        "72",
        "3",
        "Order-E",
        "3",
        "12.00",
        "14.50",
        "16.14",
        "234.29",
        "291.00",
        "13.86",
        "14, 5, 4",
        "54.10",
        "12",
        "77.90",
        "23.80",
        "Alloy",
    ]
    queries = json.loads((root / "queries.json").read_text())
    assert len(queries) == 18 and len({row["family"] for row in queries}) == 3
    with pytest.raises(ValueError, match="Preserve"):
        generate(root)


def response(content: dict[str, Any]) -> io.BytesIO:
    return io.BytesIO(
        json.dumps(
            {
                "model": DEFAULT_MODEL,
                "choices": [
                    {
                        "message": {"content": json.dumps(content)},
                        "finish_reason": "stop",
                    }
                ],
                "usage": {"total_tokens": 31, "cost": 0.0001},
            }
        ).encode()
    )


@pytest.mark.parametrize(
    "choices", [[], [None], [{"message": None}], [{"message": {"content": None}}]]
)
def test_keyed_decoder_preserves_usage_for_invalid_choices(choices: list[Any]) -> None:
    agent = client()
    agent.structured_state.structured_expected = ["q0"]
    data = {"choices": choices, "usage": {"total_tokens": 23}}
    assert agent._read_response(io.BytesIO(json.dumps(data).encode())) == data


def test_compositional_oracles_are_hand_checked_and_frozen(tmp_path: Path) -> None:
    import csv

    from benchmarks.compositional_holdout import generate

    root = tmp_path / "new_families"
    generate(root)
    with (root / "references.csv").open(newline="") as stream:
        answers = [row["answer"] for row in csv.DictReader(stream)]
    assert answers == [
        "16",
        "Survey, Design, Review, Launch",
        "4",
        "15",
        "1",
        "9",
        "0.764706",
        "2.500000",
        "3.000000",
        "8",
        "2.166667",
        "Nora",
        "3",
        "ACME Inc., Blue Co., Echo Ltd.",
        "96",
        "ACME Inc.",
        "32.00",
        "2",
    ]
    manifest = json.loads((root / "manifest.json").read_text())
    assert manifest["questions"] == 18 and len(manifest["families"]) == 3
    for name, expected in manifest["hashes"].items():
        assert hashlib.sha256((root / name).read_bytes()).hexdigest() == expected
    with pytest.raises(ValueError, match="Preserve"):
        generate(root)


def test_keyed_decoder_normalizes_only_unambiguous_single_or_shared_metadata() -> None:
    from benchmarks.batched import parse_plans

    one = {"code": "result=num('Moved','E309')", "format": {"kind": "integer"}}
    agent = client()
    agent.structured_state.structured_expected = ["q0"]
    decoded = agent._read_response(response(one))
    plans = parse_plans(decoded["choices"][0]["message"]["content"], 1)
    assert json.loads(plans[0])["code"] == one["code"]
    agent = client()
    agent.structured_state.structured_expected = ["q0", "q1"]
    envelope = {
        "programs": {"q0": one, "q1": {"code": "result=num('Moved','F309')"}},
        "format": {"kind": "auto", "separator": "; "},
        "result_variable": "result",
        "assumptions": ["Common scope remains explicit"],
        "needs_verification": True,
    }
    decoded = agent._read_response(response(envelope))
    plans = [
        json.loads(p)
        for p in parse_plans(decoded["choices"][0]["message"]["content"], 2)
    ]
    assert plans[0]["format"]["kind"] == "integer"
    assert all(p["format"]["separator"] == "; " for p in plans)
    assert all(
        p["needs_verification"] and p["assumptions"] == envelope["assumptions"]
        for p in plans
    )
    assert decoded["usage"]["total_tokens"] == 31
    for rejected in [
        one,  # One program cannot be guessed into two question slots.
        {
            **envelope,
            "format": {"kind": "decimal"},
        },  # Conflicting global/per-item format.
        {
            **envelope,
            "result_variable": "different",
            "programs": {"q0": {**one, "result_variable": "result"}, "q1": one},
        },
        {**envelope, "extra": "ignored?"},
        {**envelope, "programs": {"q0": one}},
    ]:
        original = json.dumps(rejected)
        decoded = agent._read_response(response(rejected))
        assert decoded["choices"][0]["message"]["content"] == original


def test_regression_receipt_cannot_be_reused_as_an_independent_test(
    tmp_path: Path,
) -> None:
    from benchmarks.transfer_qualification import qualify

    heldout = tmp_path / "sealed"
    heldout.mkdir()
    queries = heldout / "queries.json"
    queries.write_text("[]")
    (heldout / "manifest.json").write_text(
        json.dumps(
            {
                "hashes": {
                    "queries.json": hashlib.sha256(queries.read_bytes()).hexdigest()
                }
            }
        )
    )
    (heldout / "evaluated.json").write_text("{}")
    with (
        patch("urllib.request.urlopen") as network,
        pytest.raises(ValueError, match="consumed"),
    ):
        qualify(tmp_path / "new", heldout)
    network.assert_not_called()
    queries.write_text("changed")
    with (
        patch("urllib.request.urlopen") as network,
        pytest.raises(ValueError, match="changed"),
    ):
        qualify(tmp_path / "other", heldout)
    network.assert_not_called()


def test_individual_repair_uses_the_lossless_context(tmp_path: Path) -> None:
    path = tmp_path / "renamed.xlsx"
    workbook(path)
    agent = client()
    observed: list[list[dict[str, str]]] = []

    def request(messages: list[dict[str, str]], state: QueryState) -> str:
        observed.append(messages)
        return json.dumps(
            {
                "code": "result=num('New scope','E306')"
                if len(observed) == 1
                else "result=num('New scope','E307')"
            }
        )

    with patch.object(agent, "_request", side_effect=request):
        result = agent.solve_workbook(
            load_workbook_data(path), "What is Units for One? Answer only the integer."
        )
    assert result.answer == "139" and result.error is None
    assert len(observed) == 2
    assert all(
        "LOSSLESS CELL RUNS" in messages[1]["content"]
        and "D308:" in messages[1]["content"]
        for messages in observed
    )


def test_runtime_repair_retains_code_without_oracle_values(tmp_path: Path) -> None:
    path = tmp_path / "new.xlsx"
    workbook(path)
    agent = client()
    rejected = "result=sum(values('New scope','E307:E308')[0])"
    observed: list[list[dict[str, str]]] = []

    def request(messages: list[dict[str, str]], state: QueryState) -> str:
        observed.append(messages)
        return json.dumps(
            {
                "code": rejected
                if len(observed) == 1
                else "result=sum(values('New scope','E307:E308'))"
            }
        )

    with patch.object(agent, "_request", side_effect=request):
        result = agent.solve_workbook(
            load_workbook_data(path),
            "What is the sum of Units for One and Two? Answer only the integer.",
        )
    assert result.answer == "212" and result.error is None
    assert json.dumps(rejected) in observed[1][-1]["content"]
    assert "212" not in observed[1][-1]["content"]


def test_multi_axis_queries_are_independent_and_shared_costs_count_once(
    tmp_path: Path,
) -> None:
    path = tmp_path / "new.xlsx"
    workbook(path)
    agent = client()
    questions = [
        {"id": str(i), "file_name": path.name, "query": query}
        for i, query in enumerate(
            [
                "How many more Units are One than Two across all locations? Answer only the integer.",
                "What is Units for One? Answer only the integer.",
                "What is Units for Two? Answer only the integer.",
            ]
        )
    ]
    shared = {
        "programs": {
            "q0": {"code": "result=num('New scope','E307')"},
            "q1": {"code": "result=num('New scope','E308')"},
        }
    }
    single = {"code": "result=num('New scope','E307')-num('New scope','E308')"}
    with patch.object(
        agent, "_open_request", side_effect=[response(shared), response(single)]
    ) as calls:
        records, event = agent.solve_group(load_workbook_data(path), questions)
    assert calls.call_count == 2
    assert [record["answer"] for record in records] == ["66", "139", "73"]
    assert agent.total_cost == pytest.approx(0.0002)
    assert event["independent_question_ids"] == ["0"] and len(event["calls"]) == 1


def test_signoff_requires_a_passed_frozen_validation_before_network(
    tmp_path: Path,
) -> None:
    from benchmarks.qwen_batch_signoff import signoff, source_manifest

    for label in (
        "public_1",
        "public_2",
        "public_3",
        "regression_60",
        "regression_32",
        "known_18",
        "synthetic_24",
    ):
        stage = tmp_path / label
        stage.mkdir()
        (stage / "report.json").write_text(
            json.dumps(
                {
                    "complete": True,
                    "score": {
                        "correct": 20
                        if label.startswith("public_")
                        else int(label.rsplit("_", 1)[1]),
                        "total": 20
                        if label.startswith("public_")
                        else int(label.rsplit("_", 1)[1]),
                    },
                }
            )
        )
    heldout = tmp_path / "sealed"
    heldout.mkdir()
    manifest = heldout / "manifest.json"
    manifest.write_text(json.dumps({"questions": 24, "hashes": {}}))
    (tmp_path / "source_lock.json").write_text(
        json.dumps(
            {
                "modules": source_manifest(),
                "validation_dataset": str(heldout),
                "validation_manifest_sha256": hashlib.sha256(
                    manifest.read_bytes()
                ).hexdigest(),
            }
        )
    )
    with patch("urllib.request.urlopen") as network:
        result = signoff(tmp_path)
    assert result["dataset"] == "heldout_proof" and result["new_generations"] == 0
    network.assert_not_called()


def test_generalization_oracles_match_hand_calculated_examples(tmp_path: Path) -> None:
    import csv
    import hashlib

    from benchmarks.generalization_holdout import generate

    root = tmp_path / "independent"
    generate(root)
    with (root / "references.csv").open(newline="") as stream:
        answers = [row["answer"] for row in csv.DictReader(stream)]
    assert answers == [
        "3",
        "6",
        "Cedar",
        "Aster",
        "12",
        "4",
        "2.393443",
        "1.672131",
        "8.852459",
        "3.016393",
        "7.142857",
        "0.970674",
        "0.125000",
        "0.437500",
        "0.265625",
        "1.437500",
        "0.359375",
        "2.625000",
    ]
    manifest = json.loads((root / "manifest.json").read_text())
    assert manifest["questions"] == 18 and len(manifest["families"]) == 3
    for name, expected in manifest["hashes"].items():
        assert hashlib.sha256((root / name).read_bytes()).hexdigest() == expected
    with pytest.raises(ValueError, match="Preserve"):
        generate(root)
