import datetime as dt
import json
import time
import urllib.error
from pathlib import Path
from unittest.mock import Mock, patch

import openpyxl
import pytest

from evaluate import evaluate, read_answers
from run import run_queries, validate_queries
from src.excel_parser import grid_to_clean_text, load_workbook_data, unmerge_sheet
from src.executor import execute_code, execute_program
from src.llm_agent import (
    BudgetExceeded,
    BudgetLedger,
    LLMAgent,
    QueryState,
    SolveResult,
)
from src.postprocessor import AnswerFormat, format_answer


@pytest.fixture
def book(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Value", "Date", "Percent"])
    ws.append(["ACME Inc.", 10, dt.datetime(2024, 3, 4), 0.25])
    ws["D2"].number_format = "0%"
    ws["A201"] = "Tail"
    ws["B201"] = 42
    ws["A202"] = "Numeric merge"
    ws["B202"] = 7
    ws.merge_cells("B202:B203")
    ws["A205"] = "Header"
    ws.merge_cells("A205:C205")
    ws["E2"] = "=B2*2"
    ws["E3"] = "=SUM(B2:B2)"
    ws["E4"] = "=E4"
    ws["E5"] = "=VLOOKUP(A2,A1:B2,2)"
    ws["XFD1048576"].number_format = "0.00"
    ws.row_dimensions[201].hidden = True
    second = wb.create_sheet("Other")
    second["A1"] = "Other table"
    second["B2"] = 5
    path = tmp_path / "book.xlsx"
    wb.save(path)
    wb.close()
    return path


def agent() -> LLMAgent:
    return LLMAgent("unit-test-placeholder", pricing=(0.0000002, 0.0000008))


def test_sparse_parser_preserves_tail_types_and_merges(book: Path) -> None:
    workbook = load_workbook_data(book)
    data = workbook.sheets["Data"]
    assert "B201<number>=42" in workbook.to_text()
    assert len(data.cells) < 30
    assert data.cells["C2"].value == dt.datetime(2024, 3, 4)
    assert data.cells["D2"].value == 0.25
    assert data.cells["D2"].number_format == "0%"
    assert "B203" not in data.cells
    assert "A205:C205" in data.merges
    assert data.hidden_rows == [201]
    assert data.cells["E2"].kind == "uncached_formula"
    assert "Other" in workbook.sheets


def test_explicit_parser_limits_do_not_silently_truncate(book: Path) -> None:
    assert "Tail" in grid_to_clean_text([[None]] * 200 + [["Tail"]])
    with pytest.raises(ValueError, match="truncate"):
        grid_to_clean_text([[1]] * 201, 200)
    with pytest.raises(MemoryError):
        load_workbook_data(book, max_bytes=100)


def test_legacy_unmerge_only_inherits_text() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "Header", 10
    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:B2")
    assert unmerge_sheet(ws) == [["Header", 10], ["Header", None]]
    wb.close()


@pytest.mark.parametrize(
    ("code", "expected"),
    [
        ("values=[1,2,3]\nresult=sum(x for x in values if x>values[0])", 5),
        ("values=[1,2,3]\ndef calc():\n return sum(values)\nresult=calc()", 6),
        ("values=[1,2,3]\nresult=sorted(values,key=lambda x: -x)", [3, 2, 1]),
        (
            "from datetime import datetime\nresult=datetime.fromisoformat('2024-03-04').year",
            2024,
        ),
        ("result=0\nfor i in range(5):\n result+=i", 10),
    ],
)
def test_interpreter_python_scopes(code: str, expected: object) -> None:
    assert execute_code(code) == (True, expected, None)


@pytest.mark.parametrize(
    "code",
    [
        "scratch=[1,2,3]",
        "import os\nresult=1",
        "result=open('/etc/passwd').read()",
        "result=(1).__class__",
        "result=float('nan')",
        "result=[0]*100000000",
        "result=2**10000000",
        "class X: pass\nresult=1",
    ],
)
def test_worker_rejects_unsafe_or_invalid_programs(code: str) -> None:
    assert not execute_program(code).success


def test_worker_timeout_and_no_credential_access() -> None:
    result = execute_program("while True: pass", timeout=0.1)
    assert not result.success
    assert "timed out" in result.error
    with patch(
        "src.executor.subprocess.Popen", wraps=__import__("subprocess").Popen
    ) as popen:
        execute_code("result=1")
        assert "OPENROUTER_API_KEY" not in popen.call_args.kwargs["env"]


def test_runtime_reads_and_formula_cache_policy(book: Path) -> None:
    wire = load_workbook_data(book).wire()
    result = execute_program(
        "result=[num('Data','B201'), cell('Data','B203'), label('Data','C205'), num('Other','B2')]",
        wire,
    )
    assert result.success and result.result == [42, None, "Header", 5]
    assert "Data!B201" in result.reads
    assert execute_program("result=num('Data','E2')", wire).result == 20
    assert execute_program("result=num('Data','E3')", wire).result == 10
    assert "Circular" in execute_program("result=num('Data','E4')", wire).error
    assert not execute_program("result=num('Data','E5')", wire).success


@pytest.mark.parametrize(
    ("value", "query", "expected"),
    [
        ("ACME Inc.", "Answer with only the company name.", "ACME Inc."),
        ([1, 2], "Answer with exactly two decimal places.", "1.00, 2.00"),
        ([True, False], "Answer as two Yes or No values.", "Yes, No"),
        (
            [1234, 5678],
            "Answer as two integers separated by a comma and one space.",
            "1234, 5678",
        ),
        ("2024-03-04T00:00:00", "Answer with the date in ISO format.", "2024-03-04"),
        (3.1375, "Answer rounded to two decimal places.", "3.14"),
    ],
)
def test_answer_contracts(value: object, query: str, expected: str) -> None:
    assert format_answer(value, query) == expected


@pytest.mark.parametrize(
    ("value", "query"),
    [
        ("ID 42, unavailable", "Answer with only the integer."),
        ("1,234,5,678", "Answer as two integers."),
        ("03/04/2024", "Answer in ISO date format."),
        (1.5, "Answer with the integer."),
        (float("nan"), "Answer with only the number."),
        ([], "Answer with only the list."),
    ],
)
def test_formatter_does_not_invent_answers(value: object, query: str) -> None:
    with pytest.raises((ValueError, TypeError)):
        format_answer(value, query)


def test_explicit_query_format_overrides_model() -> None:
    assert (
        format_answer(1, "Use two decimal places", AnswerFormat(kind="integer"))
        == "1.00"
    )


def response(content: str, usage: dict[str, object] | None = None) -> Mock:
    mock = Mock()
    mock.__enter__ = Mock(return_value=mock)
    mock.__exit__ = Mock(return_value=False)
    mock.read.return_value = json.dumps(
        {
            "choices": [{"message": {"content": content}, "finish_reason": "stop"}],
            "usage": usage or {},
        }
    ).encode()
    return mock


def test_api_429_retries_and_accounts_uncertain_attempt() -> None:
    client = agent()
    state = QueryState(time.monotonic() + 30, 0.09)
    rate_limit = urllib.error.HTTPError(
        "https://example.invalid", 429, "rate limited", {}, None
    )
    with (
        patch(
            "urllib.request.urlopen",
            side_effect=[
                rate_limit,
                response("{}", {"cost": 0.0001, "total_tokens": 10}),
            ],
        ) as call,
        patch("time.sleep"),
    ):
        assert client._request([{"role": "user", "content": "test"}], state) == "{}"
    assert call.call_count == 2 and state.requests == 2
    assert state.cost > 0.0001
    assert state.calls[0]["cost_estimated"] is True
    assert client.ledger.spent == pytest.approx(state.cost)


def test_missing_usage_is_conservatively_charged() -> None:
    client = agent()
    state = QueryState(time.monotonic() + 30, 0.09)
    with patch("urllib.request.urlopen", return_value=response("{}")):
        client._request([{"role": "user", "content": "test"}], state)
    assert state.cost > 0 and state.calls[0]["cost_estimated"]


def test_budget_reservation_prevents_request() -> None:
    client = agent()
    client.ledger = BudgetLedger(0.000001)
    with patch("urllib.request.urlopen") as call, pytest.raises(BudgetExceeded):
        client._request(
            [{"role": "user", "content": "test"}],
            QueryState(time.monotonic() + 30, 0.09),
        )
    call.assert_not_called()


def test_transport_request_limit() -> None:
    client = agent()
    state = QueryState(time.monotonic() + 30, 0.09)
    with (
        patch(
            "urllib.request.urlopen", side_effect=urllib.error.URLError("offline")
        ) as call,
        patch("time.sleep"),
        pytest.raises(TimeoutError),
    ):
        client._request([{"role": "user", "content": "test"}], state)
    assert call.call_count == 5


def test_no_workbook_reads_triggers_repair(book: Path) -> None:
    client = agent()
    outputs = [
        json.dumps({"code": "result=999"}),
        json.dumps({"code": "result=num('Data','B2')"}),
    ]
    with patch.object(client, "_request", side_effect=outputs) as call:
        result = client.solve_workbook(
            load_workbook_data(book), "Return the integer value for ACME Inc."
        )
    assert result.answer == "10" and call.call_count == 2


def test_ambiguity_triggers_independent_verification(book: Path) -> None:
    client = agent()
    initial = {
        "code": "result=num('Data','B2')",
        "needs_verification": True,
        "assumptions": ["unclear scope"],
    }
    with patch.object(
        client,
        "_request",
        side_effect=[json.dumps(initial), json.dumps({"code": initial["code"]})],
    ) as call:
        result = client.solve_workbook(
            load_workbook_data(book), "Return the integer value for ACME Inc."
        )
    assert result.answer == "10" and call.call_count == 2


def test_strict_evaluator_and_invalid_ids(tmp_path: Path) -> None:
    gt, pred = tmp_path / "gt.csv", tmp_path / "pred.csv"
    gt.write_text("id,answer\nx,abc\n", encoding="utf8")
    pred.write_text("id,answer\nx,ABC\n", encoding="utf8")
    assert evaluate(pred, gt) == (0, 1)
    pred.write_text("id,answer\nx, abc \n", encoding="utf8")
    assert evaluate(pred, gt) == (0, 1)
    pred.write_text("id,answer\nx,abc\nx,ABC\n", encoding="utf8")
    with pytest.raises(ValueError, match="duplicate"):
        evaluate(pred, gt)
    pred.write_text("id,answer\nx,abc\ny,42\n", encoding="utf8")
    with pytest.raises(ValueError, match="ID mismatch"):
        evaluate(pred, gt)


def test_input_validation_and_checkpoints(book: Path, tmp_path: Path) -> None:
    queries = [
        {"id": 1, "file_name": book.name, "query": "value"},
        {"id": 2, "file_name": "missing.xlsx", "query": "value"},
    ]
    path = tmp_path / "queries.json"
    path.write_text(json.dumps(queries))
    assert validate_queries(path, tmp_path) == queries
    fake = agent()
    fake.solve_workbook = Mock(return_value=SolveResult(answer="10"))
    output = tmp_path / "output.csv"
    metrics = run_queries(tmp_path, queries, output, fake)
    assert read_answers(output) == {"1": "10", "2": ""}
    assert "2" in metrics["failures"]
    queries.append({"id": "1", "file_name": book.name, "query": "value"})
    path.write_text(json.dumps(queries))
    with pytest.raises(ValueError, match="Duplicate"):
        validate_queries(path, tmp_path)
    path.write_text(
        json.dumps([{"id": 1, "file_name": "../outside.xlsx", "query": "value"}])
    )
    with pytest.raises(ValueError, match="inside"):
        validate_queries(path, tmp_path)


def test_unused_read_cannot_ground_a_literal_answer(book: Path) -> None:
    result = execute_program(
        "unused=cell('Data','B2')\nresult=999", load_workbook_data(book).wire()
    )
    assert not result.success and "does not depend" in result.error


def test_numeric_list_with_unambiguous_separator() -> None:
    assert (
        format_answer(
            "1,234, 5,678", "Answer as two integers separated by a comma and one space."
        )
        == "1234, 5678"
    )
    assert (
        format_answer(
            "40.60, 37.60", "Answer as `men, women`, using exactly two decimal places."
        )
        == "40.60, 37.60"
    )


def test_generic_header_and_age_selection_checks(tmp_path: Path) -> None:
    from src.llm_agent import selection_errors

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Enrollment"
    ws.append(["Surname", "Birth", "Admission", "Score"])
    ws.append(["A", dt.datetime(1990, 1, 1), dt.datetime(2010, 1, 1), 50])
    ws.append(["B", dt.datetime(1992, 1, 1), dt.datetime(2025, 1, 1), 60])
    path = tmp_path / "age.xlsx"
    wb.save(path)
    wb.close()
    wire = load_workbook_data(path).wire()
    reads = {"Enrollment!B2": wire["Enrollment"]["cells"]["B2"]}
    assert selection_errors(wire, reads, "Who was the youngest on admission?")
    reads["Enrollment!C2"] = wire["Enrollment"]["cells"]["C2"]
    assert not selection_errors(wire, reads, "Who was the youngest on admission?")
    reads = {"Enrollment!D2": wire["Enrollment"]["cells"]["D2"]}
    assert selection_errors(wire, reads, 'Return the value of "Birth".')


@pytest.mark.parametrize(
    ("offset", "name_col", "sold_col", "noise"),
    [(0, 1, 2, False), (310, 4, 2, True), (7, 2, 5, True)],
)
def test_metamorphic_layout_invariance(
    tmp_path: Path, offset: int, name_col: int, sold_col: int, noise: bool
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Changed sheet name"
    ws.cell(offset + 1, name_col, "Company")
    ws.cell(offset + 1, sold_col, "Units sold")
    ws.cell(offset + 2, name_col, "Delta Co.")
    ws.cell(offset + 2, sold_col, 123)
    ws.cell(offset + 3, name_col, "Other Inc.")
    ws.cell(offset + 3, sold_col, 456)
    if noise:
        wb.create_sheet("Distractor").append(["Delta Co.", 999])
    path = tmp_path / "changed-file.xlsx"
    wb.save(path)
    wb.close()
    # Resolve coordinates using the independently generated layout; calculation
    # still reads workbook data, and never receives the oracle answer.
    from openpyxl.utils import get_column_letter

    code = f"sheet=sheets[0]\nresult=sum(num(sheet,f'{get_column_letter(sold_col)}{{r}}') for r in rows(sheet) if cell(sheet,f'{get_column_letter(name_col)}{{r}}')=='Delta Co.')"
    result = execute_program(code, load_workbook_data(path).wire())
    assert result.success and result.result == 123


def test_published_total_does_not_assume_children_add_exactly(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Counts"
    ws.append(["Group", "Workers"])
    ws.append(["All", 101])
    ws.append(["East", 40])
    ws.append(["West", 60])
    path = tmp_path / "totals.xlsx"
    wb.save(path)
    wb.close()
    result = execute_program(
        "result=num('Counts','B2')", load_workbook_data(path).wire()
    )
    assert result.result == 101


def test_matrix_and_flat_helpers_have_distinct_shapes(book: Path) -> None:
    wire = load_workbook_data(book).wire()
    assert execute_program("result=cells('Data','B2:B2')", wire).result == [[10]]
    assert execute_program("result=values('Data','B2:B2')", wire).result == [10]


def test_age_rank_direction_is_observable(tmp_path: Path) -> None:
    from src.verification import age_sort_errors

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ages"
    ws.append(["Name", "Birth", "Admission", "Score"])
    for name, birth, admission, score in [
        ("A", 1990, 2010, 10),
        ("B", 1992, 2025, 20),
        ("C", 1994, 2024, 30),
    ]:
        ws.append([name, dt.datetime(birth, 1, 1), dt.datetime(admission, 1, 1), score])
    path = tmp_path / "ages.xlsx"
    wb.save(path)
    wb.close()
    wire = load_workbook_data(path).wire()
    base = "people=[(age_days('Ages',f'B{r}',f'C{r}'),num('Ages',f'D{r}')) for r in range(2,5)]\n"
    wrong = execute_program(
        base + "ordered=sorted(people,key=lambda p:p[0])\nresult=ordered[1][1]", wire
    )
    assert age_sort_errors(wire, wrong, "What score did the oldest person receive?")
    right = execute_program(
        base
        + "ordered=sorted(people,key=lambda p:p[0],reverse=True)\nresult=ordered[0][1]",
        wire,
    )
    assert not age_sort_errors(wire, right, "What score did the oldest person receive?")
    assert right.result == 20


def test_rank_labels_use_requested_options() -> None:
    query = "Rank the following sectors: East, West, and North. Answer with three sector names in descending order."
    assert (
        format_answer(["West division", "North", "East district"], query)
        == "West, North, East"
    )
    with pytest.raises(ValueError, match="exactly once"):
        format_answer(["West", "West", "East"], query)


def test_column_age_ranges_are_not_mistaken_for_row_constraints(tmp_path: Path) -> None:
    from src.verification import row_scope_errors

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hours"
    ws.append(["Group", "1 to 34 hours", "35 hours or more"])
    ws.append(["All", 10, 20])
    path = tmp_path / "hours.xlsx"
    wb.save(path)
    wb.close()
    wire = load_workbook_data(path).wire()
    reads = {
        "Hours!B2": wire["Hours"]["cells"]["B2"],
        "Hours!C2": wire["Hours"]["cells"]["C2"],
    }
    assert not row_scope_errors(
        wire, reads, "Return totals for 1 to 34 hours and 35 hours or more."
    )
    assert (
        format_answer(
            "Leisure\nand\nhospitality", "Answer with only the industry name."
        )
        == "Leisure and hospitality"
    )


def test_in_place_sort_is_audited() -> None:
    execution = execute_program(
        "items=[(10,1),(30,2),(20,3)]\nitems.sort(key=lambda p:p[0])\nresult=items[1][1]"
    )
    assert execution.success and execution.result == 3
    assert execution.sorts[0]["keys"] == [10, 30, 20]
    assert execution.sorts[0]["selected"] == [1]


def test_mean_coverage_and_named_option_alignment(tmp_path: Path) -> None:
    from src.verification import row_scope_errors

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Metrics"
    ws.append(["Metric", "Level A", "Total", "Level B"])
    ws.append(["Rate", 1, 2, 3])
    ws.append(["Alpha", 10, 20, 30])
    ws.append(["Decoy", 100, 200, 300])
    ws.append(["Beta", 40, 50, 60])
    path = tmp_path / "metrics.xlsx"
    wb.save(path)
    wb.close()
    wire = load_workbook_data(path).wire()
    reads = {"Metrics!B2": wire["Metrics"]["cells"]["B2"]}
    assert row_scope_errors(
        wire, reads, "Calculate the average Rate across all levels."
    )
    reads = {f"Metrics!{c}2": wire["Metrics"]["cells"][f"{c}2"] for c in "BCD"}
    assert not row_scope_errors(
        wire, reads, "Calculate the average Rate across all levels."
    )
    reads = {
        "Metrics!B3": wire["Metrics"]["cells"]["B3"],
        "Metrics!B4": wire["Metrics"]["cells"]["B4"],
    }
    assert row_scope_errors(
        wire, reads, "Rank the following regions: Alpha, Beta. Answer with the names."
    )


def test_dynamic_row_lookup_does_not_assume_adjacency(book: Path) -> None:
    wire = load_workbook_data(book).wire()
    result = execute_program(
        "r=find_rows('Data','A','Tail')[0]\nresult=num('Data','B'+str(r))", wire
    )
    assert result.success and result.result == 42
    assert "Data!A201" in result.reads


def test_typed_primitives_exclude_headers_without_swallowing_errors(book: Path) -> None:
    wire = load_workbook_data(book).wire()
    result = execute_program("result=data_rows('Data',['B'],'number')", wire)
    assert result.success and result.result == [2, 201, 202]
    assert execute_program("result=num('Data','B2:B2')", wire).result == [10]
    assert execute_program("result=argmax('Data','B2:B202')", wire).result == "B201"
    assert not execute_program("result=data_rows('Data',['E'],'number')", wire).success


def test_parent_child_aggregate_overlap_is_rejected(tmp_path: Path) -> None:
    from src.verification import aggregation_overlap_errors

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Counts"
    ws.append(["Group", "Workers"])
    ws.append(["Total", 101])
    ws.append(["East", 40])
    ws.append(["West", 60])
    path = tmp_path / "counts.xlsx"
    wb.save(path)
    wb.close()
    wire = load_workbook_data(path).wire()
    reads = {f"Counts!B{r}": wire["Counts"]["cells"][f"B{r}"] for r in (2, 3, 4)}
    assert aggregation_overlap_errors(
        wire, reads, "Calculate the total workers across all groups."
    )
    assert not aggregation_overlap_errors(
        wire,
        {"Counts!B2": reads["Counts!B2"]},
        "Calculate the total workers across all groups.",
    )
    assert not aggregation_overlap_errors(
        wire, reads, "Calculate the difference between total workers and East."
    )


def test_aggregation_contract_completion_is_data_derived(tmp_path: Path) -> None:
    from src.verification import complete_mean_contract, complete_total_contract

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric", "One", "Total", "Two"])
    ws.append(["Rate", 1, 2, 3])
    ws.append(["Total workers", 101, 0, 0])
    ws.append(["East", 40, 0, 0])
    path = tmp_path / "summary.xlsx"
    wb.save(path)
    wb.close()
    wire = load_workbook_data(path).wire()
    partial = execute_program("result=num('Summary','B2')", wire)
    code = complete_mean_contract(
        wire, partial, "Calculate the average Rate across all levels."
    )
    assert code and execute_program(code, wire).result == 2
    assert (
        complete_mean_contract(
            wire, partial, "Calculate a weighted average Rate across all levels."
        )
        is None
    )
    duplicate = execute_program("result=num('Summary','B3')+num('Summary','B4')", wire)
    code = complete_total_contract(
        wire, duplicate, "Calculate total workers across all regions."
    )
    assert code and execute_program(code, wire).result == 101


def test_ordinary_underscore_variables_and_boolean_contract() -> None:
    assert execute_code("items=[(1,2),(3,4)]\nresult=sum(x for x,_ in items)") == (
        True,
        4,
        None,
    )
    assert not execute_code("result=__builtins__")[0]
    with pytest.raises(TypeError, match="Boolean"):
        format_answer(True, "Answer with only the integer.")


def test_explicit_final_prefix_fences_and_column_headers(book: Path) -> None:
    wire = load_workbook_data(book).wire()
    result = execute_program("```python\nresult_value=num('Data','B2')\n```", wire)
    assert result.success and result.result == 10
    assert execute_program("result=label('Data','B')", wire).result == "Value"
    assert not execute_program("scratch=num('Data','B2')", wire).success
    assert execute_code("import string\nresult=string.ascii_uppercase[0]") == (
        True,
        "A",
        None,
    )


def test_soft_word_wrap_does_not_delete_date_range_hyphens() -> None:
    from src.verification import canonical_label

    assert canonical_label("Transpor-\ntation") == "Transportation"
    assert canonical_label("Sep. 2024-\nOct. 2024") == "Sep. 2024- Oct. 2024"


def test_declared_output_name_and_syntax_repair(book: Path) -> None:
    wire = load_workbook_data(book).wire()
    assert (
        execute_program(
            "chosen=num('Data','B2')", wire, result_variable="chosen"
        ).result
        == 10
    )
    client = agent()
    with patch.object(
        client,
        "_request",
        side_effect=[
            json.dumps({"code": "result = [1:"}),
            json.dumps(
                {"code": "chosen=num('Data','B2')", "result_variable": "chosen"}
            ),
        ],
    ):
        result = client.solve_workbook(
            load_workbook_data(book), "Return the integer value for ACME Inc."
        )
    assert result.answer == "10" and result.error is None


def test_composite_axis_caption_does_not_hide_section_entities(tmp_path: Path) -> None:
    from src.verification import entity_label_errors

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Groups"
    ws.append(["Metric, sex, race and age", "Count"])
    ws.append(["Group A", None])
    ws.append(["Employed", 10])
    path = tmp_path / "groups.xlsx"
    wb.save(path)
    wb.close()
    wire = load_workbook_data(path).wire()
    result = execute_program("result=cell('Groups','A2')", wire)
    assert not entity_label_errors(wire, result, "Answer with only the race name.")


def test_explicit_event_reference_overrides_admission_default(tmp_path: Path) -> None:
    from src.llm_agent import selection_errors
    from src.verification import age_reference_columns

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Events"
    ws.append(["Name", "Birth", "Admission", "Marriage", "Score"])
    ws.append(
        [
            "Vela",
            dt.datetime(1990, 1, 1),
            dt.datetime(2010, 1, 1),
            dt.datetime(2025, 1, 1),
            10,
        ]
    )
    ws.append(
        [
            "Nova",
            dt.datetime(1992, 1, 1),
            dt.datetime(2025, 1, 1),
            dt.datetime(2010, 1, 1),
            20,
        ]
    )
    path = tmp_path / "events.xlsx"
    wb.save(path)
    wb.close()
    wire = load_workbook_data(path).wire()
    question = "Who was the oldest on their marriage date? Answer with only the name."
    assert age_reference_columns(wire["Events"], question) == ("B", "D")
    assert age_reference_columns(wire["Events"], "Who was the oldest?") == ("B", "C")
    assert (
        age_reference_columns(wire["Events"], "Who was the oldest by date of birth?")
        is None
    )
    result = execute_program(
        "r=age_order('Events','B','D',descending=True)[0]\nresult=cell('Events','A'+str(r))",
        wire,
    )
    assert result.success and result.result == "Vela"
    assert not selection_errors(wire, result.reads, question)


def test_account_failure_stops_subsequent_network_requests() -> None:
    client = agent()
    error = urllib.error.HTTPError(
        "https://example.invalid", 402, "payment required", {}, None
    )
    with patch("urllib.request.urlopen", side_effect=error) as call:
        with pytest.raises(ValueError, match="HTTP 402"):
            client._request(
                [{"role": "user", "content": "test"}],
                QueryState(time.monotonic() + 30, 0.09),
            )
        with pytest.raises(ValueError, match="HTTP 402"):
            client._request(
                [{"role": "user", "content": "test"}],
                QueryState(time.monotonic() + 30, 0.09),
            )
    assert call.call_count == 1
