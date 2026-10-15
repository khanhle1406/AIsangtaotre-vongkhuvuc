"""A smaller planning view must never masquerade as a complete workbook."""

import json
from pathlib import Path
from unittest.mock import patch

import openpyxl

from benchmarks.backend_view import backend_view
from benchmarks.efficient import dense_text
from benchmarks.gateway import BatchLimits
from benchmarks.packed import SampledPackedAgent
from src.context import inspect_regions
from src.excel_parser import load_workbook_data
from src.executor import execute_program


def test_label_cleaning_is_typed_and_keeps_meaningful_punctuation() -> None:
    import pytest

    from src.excel_parser import CellData, SheetData, WorkbookData
    from src.executor import clean_label

    assert (
        clean_label("Trans-\nportation\nand utilities")
        == "Transportation and utilities"
    )
    assert clean_label("  ACME-Group Inc.  ") == "ACME-Group Inc."
    for invalid in [None, 17, False]:
        with pytest.raises(TypeError, match="expects text"):
            clean_label(invalid)
    wire = WorkbookData(
        {
            "Moved": SheetData(
                "Moved", cells={"G408": CellData("Con-\nstruction", "string")}
            )
        }
    ).wire()
    actual = execute_program("result=clean_label(cell('Moved','G408'))", wire)
    assert actual.success and actual.result == "Construction"
    literal = execute_program(
        "unused=cell('Moved','G408');result=clean_label('Copied answer')", wire
    )
    assert not literal.success and "does not depend" in literal.error


def test_sheet_names_support_list_and_snapshot_without_foreign_access() -> None:
    from src.excel_parser import CellData, SheetData, WorkbookData

    book = WorkbookData(
        {"Renamed": SheetData("Renamed", cells={"C408": CellData(23, "number")})}
    ).wire()
    for expression in ["sheets[0]", "sheets()[0]", "list(sheets)[0]"]:
        result = execute_program(f"result=cell({expression},'C408')", book)
        assert result.success and result.result == 23
    result = execute_program(
        "snapshot=sheets();snapshot.append('Foreign');result=cell(sheets[0],'C408')",
        book,
    )
    assert result.success and result.result == 23
    assert not execute_program("result=cell('Foreign','C408')", book).success
    assert not execute_program("result=cell(sheets('Foreign')[0],'C408')", book).success
    shadowed = execute_program(
        "sheets=['Renamed'];result=cell(sheets()[0],'C408')", book
    )
    assert not shadowed.success and "not callable: sheets" in shadowed.error


def test_numeric_sample_discloses_original_ranges_and_preserves_full_backend(
    tmp_path: Path,
) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Moved"
    sheet["F301"], sheet["G301"], sheet["H301"] = "Entity", "Units (thousands)", "Rate"
    sheet.merge_cells("F299:H299")
    sheet["F299"] = "Complete measured population; missing is not zero"
    for index in range(100):
        row = 302 + index
        sheet[f"F{row}"] = f"Entity-{index}"
        sheet[f"G{row}"] = index * 7 - 25
        sheet[f"H{row}"] = index / 97
        sheet[f"H{row}"].number_format = "0.00%"
    sheet["G315"] = -900  # Interior extreme must remain available, despite omission.
    sheet["I301"], sheet["I315"] = "Formula", "=G315*2"
    sheet.row_dimensions[315].hidden = True
    path = tmp_path / "renamed.xlsx"
    book.save(path)
    book.close()
    data = load_workbook_data(path)
    original = data.wire()
    text = backend_view(data, original)
    catalog = json.loads(text.rsplit("metadata): ", 1)[1])["Moved"]
    assert catalog["omitted_numeric_values"] == 194
    assert catalog["visible_cells"] + catalog["omitted_numeric_values"] == len(
        data.sheets["Moved"].cells
    )
    assert "ALL stored cells follow" not in text and "no rows omitted" not in text
    assert "NOT blanks or zero" in text and "Entity-13" in text
    assert "-900" not in text and "=G315*2" in text
    assert "0.00%" in text and "F299:H299" in text and '"hidden_rows":[315]' in text
    assert "Units (thousands)" in text and '"headers"' in text
    assert len(text.encode()) < len(
        dense_text(data, original, column_catalog=True).encode()
    )
    assert (
        inspect_regions(original, [{"sheet": "Moved", "range": "G315"}]).find("-900")
        >= 0
    )
    result = execute_program("result=min(values('Moved','G302:G401'))", original)
    assert result.success and result.result == -900
    assert data.wire() == original
    agent = SampledPackedAgent(
        api_key="test-placeholder",
        gateway_url="https://gateway.invalid/v1",
        gateway_model="ag/gemini-3-flash",
        batch_limits=BatchLimits(tmp_path / "batch.json"),
    )
    seen = []
    responses = [
        {"code": "result=0", "request_regions": [{"sheet": "Moved", "range": "G315"}]},
        {"code": "result=num('Moved','G315')", "format": {"kind": "integer"}},
    ]

    def request(messages: list[dict[str, str]], _state: object) -> str:
        seen.append(messages[-1]["content"])
        return json.dumps(responses[len(seen) - 1])

    with patch.object(agent, "_request", side_effect=request):
        answer = agent.solve_workbook(
            data, "What is Units for Entity-13? Answer as an integer."
        )
    assert answer.answer == "-900" and answer.error is None
    assert "BACKEND PROGRAM VIEW" in seen[0] and "-900" not in seen[0]
    assert "Requested original workbook cells" in seen[1] and "-900" in seen[1]


def test_small_and_string_tables_keep_the_existing_full_encoding(
    tmp_path: Path,
) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Entity", "Units"])
    sheet.append(["Birch", 17])
    sheet.append(["Cedar", 23])
    sheet.append(["Footnote", "Published category boundaries apply."])
    path = tmp_path / "small.xlsx"
    book.save(path)
    book.close()
    data = load_workbook_data(path)
    assert backend_view(data, data.wire()) == dense_text(
        data, data.wire(), column_catalog=True
    )


def test_without_derived_hints_keeps_identical_original_cells_and_metadata(
    tmp_path: Path,
) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.merge_cells("G307:H307")
    sheet["G307"] = "Measured population (thousands)"
    sheet["F308"], sheet["G308"], sheet["H308"] = "Entity", "Overall", "Restricted"
    sheet["F309"], sheet["G309"], sheet["H309"] = "Birch", 17, 11
    sheet["H309"].number_format = "0.0%"
    path = tmp_path / "headers.xlsx"
    book.save(path)
    book.close()
    data = load_workbook_data(path)
    full = dense_text(data, data.wire(), column_catalog=True)
    raw = dense_text(data, data.wire(), column_catalog=True, header_hints=False)
    assert [s for s in full.splitlines() if not s.startswith("HINTS ")] == [
        s for s in raw.splitlines() if not s.startswith("HINTS ")
    ]
    hints = json.loads(next(s[6:] for s in raw.splitlines() if s.startswith("HINTS ")))
    assert hints["merges"] == ["G307:H307"] and "headers" not in hints
    assert "Overall" in raw and "Restricted" in raw and "0.0%" in raw
