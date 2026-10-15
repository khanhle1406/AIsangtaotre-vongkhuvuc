"""Selective agreement uses real sources and fresh prompts, never reference answers."""

import json
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

from benchmarks.gateway import BatchLimits
from benchmarks.packed import IndependentPackedAgent
from src.excel_parser import CellData, SheetData, WorkbookData, load_workbook_data
from src.executor import execute_program
from src.verification import named_row_scope_risks


def test_header_domain_review_is_fresh_even_when_candidate_audit_is_enabled(
    tmp_path: Path,
) -> None:
    from src.llm_agent import LLMAgent

    book = openpyxl.Workbook()
    sheet = book.active
    sheet["F407"], sheet["F408"], sheet["F409"] = "Category", "Birch", "Cedar"
    path = tmp_path / "moved.xlsx"
    book.save(path)
    book.close()
    data = load_workbook_data(path)
    first = "r=data_rows('Sheet',{'F':'string'},407,409);result=len([cell('Sheet','F'+str(i)) for i in r])"
    second = first.replace("407,409", "408,409")
    agent = LLMAgent(api_key="test-placeholder", pricing=(0, 0), review_candidate=True)
    with patch.object(
        agent,
        "_request",
        side_effect=[
            json.dumps({"code": c, "format": {"kind": "integer"}})
            for c in [first, second]
        ],
    ) as request:
        result = agent.solve_workbook(
            data, "How many required categories are listed? Answer as an integer."
        )
    assert result.answer == "2" and result.error is None and request.call_count == 2
    assert result.attempts[0]["verification_mode"] == "fresh-selection"
    assert all(first not in m["content"] for m in request.call_args_list[1].args[0])


def test_repeated_runtime_failure_gets_fresh_third_prompt() -> None:
    from src.llm_agent import LLMAgent

    book = WorkbookData(
        {"Events": SheetData("Events", cells={"C407": CellData(23, "number")})}
    )
    invalid = "v=num('Events','C407');result=v+missing"
    valid = "result=num('Events','C407')*2"
    agent = LLMAgent(api_key="test-placeholder", pricing=(0, 0))
    responses = [
        json.dumps({"code": code, "format": {"kind": "integer"}})
        for code in [invalid, invalid, valid]
    ]
    with patch.object(agent, "_request", side_effect=responses) as request:
        result = agent.solve_workbook(
            book, "Return twice the stored Units as an integer."
        )
    assert result.answer == "46" and result.error is None and request.call_count == 3
    assert invalid in request.call_args_list[1].args[0][-1]["content"]
    assert (
        "same failure occurred twice"
        in request.call_args_list[2].args[0][-1]["content"]
    )
    assert all(invalid not in m["content"] for m in request.call_args_list[2].args[0])
    assert any(
        a.get("repair_mode") == "fresh-after-repeated-failure" for a in result.attempts
    )


@pytest.mark.parametrize("column", ["A", "G", "AA"])
def test_named_record_binding_uses_actual_rows_and_preserves_joins(
    tmp_path: Path, column: str
) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Moved"
    measure = openpyxl.utils.get_column_letter(
        openpyxl.utils.column_index_from_string(column) + 1
    )
    for row, label, amount in [
        (406, "Entity", "Units"),
        (407, "Birch", 17),
        (408, "Cedar", 29),
        (409, "Aster", 11),
    ]:
        sheet[f"{column}{row}"] = label
        sheet[f"{measure}{row}"] = amount
    path = tmp_path / "renamed.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    bad = execute_program(f"result=num('Moved','{measure}407')", wire)
    question = "What is Units for Cedar? Answer only the integer."
    assert bad.success and named_row_scope_risks(wire, bad, question)
    good = execute_program(f"result=num('Moved','{measure}408')", wire)
    assert not named_row_scope_risks(wire, good, question)
    joined = execute_program(
        f"key=cell('Moved','{column}408'); result=num('Moved','{measure}407') if key else 0",
        wire,
    )
    # Reading the actual identity may feed a join: this guard alone cannot
    # establish the correct numeric row after that join.
    assert not named_row_scope_risks(wire, joined, question)
    assert not named_row_scope_risks(wire, bad, "What is Units across all entities?")
    wire["Moved"]["cells"][f"{column}409"]["value"] = "Cedar"
    assert not named_row_scope_risks(wire, bad, question)


@pytest.mark.parametrize("mode", ["agree", "disagree", "single", "no_turn"])
def test_independent_solution_is_unanchored_and_bounded(
    tmp_path: Path, mode: str
) -> None:
    book = WorkbookData(
        {
            name: SheetData(name, cells={"B8": CellData(value, "number")})
            for name, value in [("Source", 7), ("Weights", 3)]
        }
    )
    agent = IndependentPackedAgent(
        api_key="test-placeholder",
        gateway_url="https://gateway.invalid/v1",
        gateway_model="ag/gemini-3-flash",
        batch_limits=BatchLimits(tmp_path / "batch.json"),
    )
    first = "a=num('Source','B8');b=num('Weights','B8');result=a*b"
    second = first if mode == "agree" else first.replace("a*b", "a+b")
    if mode == "single":
        first = second = "result=num('Source','B8')"
    codes = [first, second, second]
    if mode == "no_turn":
        codes = ["result=num('Source','B8')/0", first, second]
    responses = [
        json.dumps({"code": code, "format": {"kind": "integer"}}) for code in codes
    ]
    with patch.object(agent, "_request", side_effect=responses) as request:
        result = agent.solve_workbook(
            book, "Calculate the specified score as an integer."
        )
    if mode == "single":
        assert request.call_count == 1 and result.answer == "7"
        return
    if mode == "no_turn":
        assert request.call_count == 2 and result.answer == "21"
        assert result.error is None
        assert "Two inference turns" in result.attempts[-1]["independent_check_skipped"]
        return
    assert request.call_count == (2 if mode == "agree" else 3)
    independent_call = request.call_args_list[1]
    independent_prompt = independent_call.args[0]
    assert (
        "No earlier program or answer is supplied" in independent_prompt[-1]["content"]
    )
    assert all(first not in message["content"] for message in independent_prompt)
    assert result.error is None and result.answer == ("21" if mode == "agree" else "10")
    assert result.attempts[0]["verification_mode"] == "independent-solution"
    assert result.attempts[1]["independent_agreement"] == (mode == "agree")
    if mode == "disagree":
        adjudication = request.call_args_list[2].args[0][-1]["content"]
        assert "unverified candidates" in adjudication
        assert first in json.loads(adjudication.split("Candidates:\n", 1)[1])[0]["code"]
