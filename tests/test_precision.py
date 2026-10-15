"""Structural fixes must hold after relocation and with independent numeric values."""

from pathlib import Path
from typing import Any

import openpyxl
import pytest

from src.excel_parser import load_workbook_data
from src.executor import execute_program
from src.verification import (
    aggregation_overlap_errors,
    complete_total_contract,
    evidence_errors,
)


@pytest.mark.parametrize(
    "offset,column,total", [(0, 2, 117), (310, 4, 263), (9, 6, 49)]
)
def test_omitted_total_and_nested_age_bands(
    tmp_path: Path, offset: int, column: int, total: int
) -> None:
    from openpyxl.utils import get_column_letter

    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "Unseen ledger"
    ws.cell(offset + 1, 1, "Age")
    ws.cell(offset + 1, column, "Participants")
    for r, label, value in [
        (2, "Total, 16 years and over", total),
        (3, "16 to 19 years", 30),
        (4, "16 to 17 years", 10),
        (5, "18 to 19 years", 20),
        (6, "20 years and over", total - 30),
    ]:
        ws.cell(offset + r, 1, label)
        ws.cell(offset + r, column, value)
    path = tmp_path / "renamed.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    col = get_column_letter(column)
    code = f"result=sum(values('Unseen ledger','{col}{offset + 3}:{col}{offset + 6}'))"
    before = execute_program(code, wire)
    query = "Calculate the total participants across all age groups. Answer only the integer."
    assert aggregation_overlap_errors(wire, before.reads, query)
    repair = complete_total_contract(wire, before, query)
    assert repair is not None
    after = execute_program(repair, wire)
    assert after.success and after.result == total
    assert (
        complete_total_contract(
            wire, before, "Calculate total participants aged 16 to 19 years."
        )
        is None
    )


def test_age_overlap_does_not_merge_separate_tables(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "S"
    for row, label, value in [
        (1, "Age", "Workers"),
        (2, "16 to 19 years", 30),
        (3, "20 years and over", 70),
    ]:
        ws.cell(row, 1, label)
        ws.cell(row, 2, value)
        ws.cell(row, 5, label)
        ws.cell(row, 6, value)
    path = tmp_path / "tables.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    result = execute_program("result=num('S','B2')+num('S','F2')", wire)
    assert not aggregation_overlap_errors(
        wire,
        result.reads,
        "Combined total across all age groups for the two independent regions.",
    )


def test_category_helper_excludes_only_grand_total(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "S"
    ws.append(["Category", "Rate"])
    for row in [
        ("Total, all staff", 100),
        ("1 to 39 hours", 30),
        ("40 hours and over", 70),
        ("40 to 49 hours", 45),
        ("Allied contractors", 12),
    ]:
        ws.append(row)
    path = tmp_path / "categories.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    result = execute_program(
        "r=max(category_rows('S','A','B'), key=lambda r:num('S',f'B{r}'))\nresult=cell('S',f'A{r}')",
        wire,
    )
    assert result.success and result.result == "40 hours and over"
    assert any(e.get("value") == "Allied contractors" for e in result.reads.values())


def test_total_repair_uses_its_own_evidence(tmp_path: Path) -> None:
    from unittest.mock import patch

    from src.llm_agent import LLMAgent

    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "S"
    ws.append(["Age", "Workers"])
    for row in [
        ("Total, 16 years and over", 101),
        ("16 to 19 years", 20),
        ("16 to 17 years", 8),
        ("18 to 19 years", 12),
        ("20 years and over", 81),
    ]:
        ws.append(row)
    path = tmp_path / "one.xlsx"
    book.save(path)
    book.close()
    import json

    program = {
        "code": "result=sum(values('S','B3:B6'))",
        "format": {"kind": "integer"},
        "selection": [{"sheet": "S", "range": "B3:B6"}],
    }
    client = LLMAgent("test-placeholder", pricing=(0, 0))
    with patch.object(client, "_request", return_value=json.dumps(program)) as request:
        answer = client.solve_workbook(
            load_workbook_data(path),
            "Calculate total workers across all age groups. Answer with only the integer.",
        )
    assert answer.answer == "101" and request.call_count == 1
    assert not evidence_errors(
        load_workbook_data(path).wire(),
        answer.attempts[-1]["reads"],
        answer.attempts[-1]["selection"],
    )


def test_sealed_dataset_has_disjoint_files_and_independent_identity_oracle(
    tmp_path: Path,
) -> None:
    import csv
    import hashlib
    import json

    from benchmarks.generalization import generate

    root = tmp_path / "blind"
    queries, refs = generate(root, seed=413779)
    questions = json.loads(queries.read_text())
    with refs.open(newline="") as stream:
        gold = {r["id"]: r["answer"] for r in csv.DictReader(stream)}
    assert len(questions) == 24 and len({q["family"] for q in questions}) == 4
    manifest = json.loads((root / "manifest.json").read_text())
    assert all(
        hashlib.sha256((root / path).read_bytes()).hexdigest() == sha
        for path, sha in manifest["sha256"].items()
    )
    book = openpyxl.load_workbook(
        root / "xlsx_files/cohort_duplicate.xlsx", read_only=True, data_only=True
    )
    completed = [
        row for row in list(book["Membership"].values)[1:] if row[2] == "Completed"
    ]
    assert gold["blind-001"] == str(len({row[0] for row in completed}))
    assert gold["blind-002"] == str(len(completed))
    assert gold["blind-001"] != gold["blind-002"]
    book.close()
    # Independently read the published H1 value, rather than reusing generator sums.
    book = openpyxl.load_workbook(
        root / "xlsx_files/fiscal_partition.xlsx", read_only=True, data_only=True
    )
    value = book["Fiscal sales"]["B2"].value
    assert gold["blind-019"] == str(value) and gold["blind-020"] == str(value * 1000)
    book.close()
    with pytest.raises(ValueError, match="new directory"):
        generate(root)


def test_null_optional_format_uses_defaults_but_bad_types_still_fail() -> None:
    from pydantic import ValidationError

    from src.postprocessor import AnswerFormat, format_answer

    supplied = {
        "kind": "integer",
        "separator": None,
        "rounding": None,
        "date_format": None,
    }
    spec = AnswerFormat.model_validate(supplied)
    assert format_answer(39, "Answer only the integer.", spec) == "39"
    assert supplied["separator"] is None
    assert (
        format_answer([2, 3], "Answer two integers, comma and one space.", spec)
        == "2, 3"
    )
    with pytest.raises(ValidationError):
        AnswerFormat.model_validate({"separator": 15})
    with pytest.raises(ValidationError):
        AnswerFormat.model_validate({"kind": "guessed"})


def test_gateway_finishes_without_waiting_for_connection_close(tmp_path: Path) -> None:
    import json
    from collections.abc import Iterator

    from benchmarks.gateway import BatchLimits, GatewayAgent

    class Response:
        def __init__(self) -> None:
            self.headers = {"Content-Type": "text/event-stream"}

        def __iter__(self) -> Iterator[bytes]:
            yield (
                b"data: "
                + json.dumps(
                    {
                        "model": "gemini-3-flash",
                        "choices": [
                            {"delta": {"content": "{}"}, "finish_reason": "stop"}
                        ],
                    }
                ).encode()
                + b"\n"
            )
            yield (
                b"data: "
                + json.dumps({"choices": [], "usage": {"total_tokens": 4}}).encode()
                + b"\n"
            )
            raise TimeoutError("Proxy kept connection alive after complete response")

    agent = GatewayAgent(
        api_key="test-placeholder",
        gateway_url="https://gateway.invalid/v1",
        gateway_model="ag/gemini-3-flash",
        batch_limits=BatchLimits(tmp_path / "batch.json"),
    )
    data = agent._read_response(Response())
    assert (
        data["usage"]["total_tokens"] == 4
        and data["choices"][0]["message"]["content"] == "{}"
    )


def test_category_contract_keeps_intermediate_categories_and_numeric_evidence(
    tmp_path: Path,
) -> None:
    from src.verification import complete_category_contract

    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "S"
    ws.append(["Category", "Percent"])
    for row in [
        ("Total", 100),
        ("Under 40 hours", 25),
        ("40 hours and over", 75),
        ("40 to 49 hours", 45),
        ("50 hours and over", 30),
    ]:
        ws.append(row)
    path = tmp_path / "ranking.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    wrong = execute_program("data=values('S','B2:B7')\nresult=cell('S','A5')", wire)
    query = "Which hour category had the highest percentage? Answer with only the hour category."
    repair = complete_category_contract(wire, wrong, query)
    assert repair is not None
    fixed = execute_program(repair, wire)
    assert fixed.success and fixed.result == "40 hours and over"
    assert any(e["kind"] == "number" for e in fixed.reads.values())
    assert (
        complete_category_contract(
            wire, wrong, "Which leaf category had the highest percentage?"
        )
        is None
    )
    assert (
        complete_category_contract(
            wire,
            wrong,
            "Which category among those under 40 hours had the highest percentage?",
        )
        is None
    )


def test_precision_qwen_keeps_judge_model_and_openrouter_transport() -> None:
    import json
    import time
    from unittest.mock import patch

    from benchmarks.precision import ACCURACY_CONTRACT
    from benchmarks.qwen_signoff import CandidateAgent
    from src.llm_agent import DEFAULT_MODEL, OPENROUTER_URL, QueryState

    agent = CandidateAgent(
        api_key="test-placeholder", pricing=(0, 0), precision_profile="robust"
    )
    assert agent.model == DEFAULT_MODEL and agent._request_url() == OPENROUTER_URL
    assert agent.iterator_helpers and agent.compact_prompt

    from io import BytesIO

    response = BytesIO(
        json.dumps(
            {
                "model": DEFAULT_MODEL,
                "choices": [{"message": {"content": "{}"}, "finish_reason": "stop"}],
                "usage": {"total_tokens": 4, "cost": 0},
            }
        ).encode()
    )
    # Inspect the exact outgoing payload without making any network request.
    with patch.object(agent, "_open_request", return_value=response) as request:
        agent._request(
            [{"role": "user", "content": "task"}],
            QueryState(time.monotonic() + 60, 0.09),
        )
    payload = json.loads(request.call_args.args[0].data)
    assert payload["model"] == DEFAULT_MODEL and payload["temperature"] == 0
    assert (
        payload["reasoning"]["max_tokens"] == 256 and "reasoning_effort" not in payload
    )
    assert ACCURACY_CONTRACT in payload["messages"][-1]["content"]


def test_lookup_candidates_are_not_claimed_as_selected_data(tmp_path: Path) -> None:
    from src.verification import evidence_errors

    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "S"
    ws.append(["Group", "Units"])
    ws.append(["Delta", 7])
    ws.append(["Other", 11])
    ws.append(["Delta", 19])
    path = tmp_path / "lookup.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    result = execute_program(
        "r=find_rows('S','A','Delta')[0]\nresult=num('S',f'B{r}')", wire
    )
    assert result.success and result.result == 7
    assert result.reads["S!A4"]["kind"] == "lookup_probe"
    assert not evidence_errors(wire, result.reads, [{"sheet": "S", "range": "A2:B2"}])
    # Actual selected numeric reads still cannot escape the declaration.
    assert evidence_errors(wire, result.reads, [{"sheet": "S", "range": "A2"}])
    explicit = execute_program(
        "r=find_rows('S','A','Delta')[0]\nresult=cell('S',f'A{r}')", wire
    )
    assert explicit.reads["S!A2"]["kind"] == "string"


def test_precision_cli_uses_actual_workbook_and_official_qwen(tmp_path: Path) -> None:
    import io
    import json
    from unittest.mock import patch

    from benchmarks.precision_submission import main
    from src.llm_agent import DEFAULT_MODEL, LLMAgent

    files = tmp_path / "new inputs"
    files.mkdir()
    book = openpyxl.Workbook()
    book.active.title = "Changed"
    book.active.append(["Name", "Units"])
    book.active.append(["Item", 47])
    book.save(files / "opaque.xlsx")
    book.close()
    queries = tmp_path / "new queries.json"
    queries.write_text(
        json.dumps(
            [
                {
                    "id": "new-id",
                    "file_name": "opaque.xlsx",
                    "query": "How many units? Answer only the integer.",
                }
            ]
        )
    )
    output = tmp_path / "outputs/submission.csv"
    program = json.dumps(
        {"code": "result=num(sheets[0],'B2')", "format": {"kind": "integer"}}
    )
    response = io.BytesIO(
        json.dumps(
            {
                "model": DEFAULT_MODEL,
                "choices": [{"message": {"content": program}, "finish_reason": "stop"}],
                "usage": {"total_tokens": 50, "cost": 0},
            }
        ).encode()
    )
    with (
        patch(
            "sys.argv",
            [
                "precision_submission",
                "--input-dir",
                str(files),
                "--queries",
                str(queries),
                "--output",
                str(output),
                "--api-key",
                "test-placeholder",
            ],
        ),
        patch("src.llm_agent.get_pricing", return_value=(2e-7, 8e-7)),
        patch.object(LLMAgent, "_open_request", return_value=response) as request,
    ):
        main()
    from evaluate import read_answers

    assert read_answers(output) == {"new-id": "47"}
    assert (
        request.call_args.args[0].full_url
        == "https://openrouter.ai/api/v1/chat/completions"
    )
    assert json.loads(request.call_args.args[0].data)["model"] == DEFAULT_MODEL


def test_acceptance_uses_accuracy_before_time_then_tokens() -> None:
    from benchmarks.qwen_signoff import acceptance_verdict

    def result(correct: int, seconds: float, tokens: int) -> dict[str, Any]:
        return {
            "score": {"correct": correct, "total": 20},
            "metrics": {
                "failures": {},
                "answered": 20,
                "elapsed_seconds": seconds,
                "tokens": tokens,
            },
            "fatal_error": None,
        }

    baseline = result(19, 20, 100)
    improved = result(20, 200, 1000)
    assert acceptance_verdict(
        {"baseline_public": baseline, "candidate_public": improved}, ["public"]
    )["promotion_eligible"]
    baseline = result(20, 20, 100)
    assert not acceptance_verdict(
        {"baseline_public": baseline, "candidate_public": improved}, ["public"]
    )["promotion_eligible"]
    tied = result(20, 20, 90)
    assert acceptance_verdict(
        {"baseline_public": baseline, "candidate_public": tied}, ["public"]
    )["promotion_eligible"]
    tied["metrics"]["answered"] = 19
    assert not acceptance_verdict(
        {"baseline_public": baseline, "candidate_public": tied}, ["public"]
    )["promotion_eligible"]


@pytest.mark.parametrize("value", [-273, -8, 0, 11, 719])
def test_explicit_magnitude_is_general_and_signed_outputs_remain_signed(
    value: int,
) -> None:
    from src.postprocessor import format_answer

    assert format_answer(
        value,
        "Report the positive magnitude of the adjustment. Answer only the integer.",
    ) == str(abs(value))
    assert format_answer(
        value, "Report the net adjustment. Keep its sign. Answer only the integer."
    ) == str(value)
    assert format_answer(
        str(value), "Return the absolute value. Answer only the integer."
    ) == str(abs(value))
    assert format_answer(
        value,
        "Explain the absolute value elsewhere, but report the adjustment and keep its sign. Answer only the integer.",
    ) == str(value)


def test_total_repair_respects_explicit_leaf_contract(tmp_path: Path) -> None:
    from src.verification import complete_total_contract

    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "S"
    ws.append(["Age", "Count"])
    for row in [("Total", 101), ("16 to 19 years", 20), ("20 years and over", 80)]:
        ws.append(row)
    path = tmp_path / "rounding.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    observed = execute_program("result=sum(values('S','B2:B4'))", wire)
    assert (
        complete_total_contract(
            wire,
            observed,
            "Sum leaf rows across all age groups; exclude the published total.",
        )
        is None
    )


def test_category_minimum_and_tie_deferral(tmp_path: Path) -> None:
    from src.verification import complete_category_contract

    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "S"
    ws.append(["Category", "Percent"])
    for row in [("Total", 100), ("Small", 20), ("Large", 80)]:
        ws.append(row)
    path = tmp_path / "minimum.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    wrong = execute_program("data=values('S','B2:B4')\nresult=cell('S','A4')", wire)
    repair = complete_category_contract(
        wire, wrong, "Which category had the lowest percentage?"
    )
    assert repair and execute_program(repair, wire).result == "Small"
    wire["S"]["cells"]["B3"]["value"] = 80
    assert (
        complete_category_contract(
            wire, wrong, "Which category had the highest percentage?"
        )
        is None
    )
