"""Projection evidence is conservative and cannot change source values or audit reads."""

import copy
from typing import Any

import pytest
from openpyxl.utils import get_column_letter

from src.excel_parser import CellData, SheetData, WorkbookData
from src.executor import execute_program
from src.source_projection import column_projection
from src.verification import column_scope_errors


def fixture(
    offset: int = 0, left: int = 1, flipped: bool = False
) -> tuple[dict[str, Any], str, str]:
    sheet = "Moved original"
    letters = [get_column_letter(left + i) for i in range(5)]
    label, first, second, third, fourth = letters
    raw, adjusted = (third, first) if flipped else (first, third)
    raw_last = get_column_letter(left + (4 if flipped else 2))
    adjusted_last = get_column_letter(left + (2 if flipped else 4))
    cells = {
        f"{label}{offset + 1}": CellData("Entity", "string"),
        f"{raw}{offset + 1}": CellData("Raw movement", "string"),
        f"{adjusted}{offset + 1}": CellData("Adjusted movement", "string"),
    }
    for column in [first, third]:
        cells[f"{column}{offset + 2}"] = CellData("May to June", "string")
    for column in [second, fourth]:
        cells[f"{column}{offset + 2}"] = CellData("June to July", "string")
    for row, (name, raw_value, adjusted_value) in enumerate(
        [("ACME Inc.", -37, 2), ("Birch", 5, 900), ("Cedar", 2, 17)], offset + 3
    ):
        for column in letters[1:]:
            cells[f"{column}{row}"] = CellData(0, "number")
        cells[f"{label}{row}"] = CellData(name, "string")
        cells[f"{raw_last}{row}"] = CellData(raw_value, "number")
        cells[f"{adjusted_last}{row}"] = CellData(adjusted_value, "number")
    book = WorkbookData(
        {
            sheet: SheetData(
                sheet,
                cells=cells,
                merges=[
                    f"{raw}{offset + 1}:{raw_last}{offset + 1}",
                    f"{adjusted}{offset + 1}:{adjusted_last}{offset + 1}",
                ],
            )
        }
    )
    wrong_index = 2 if flipped else 4
    code = f"data=cells({sheet!r},'{label}{offset + 3}:{fourth}{offset + 5}')\nbest=None\nmaximum=-1\nfor item in data:\n    name=item[0]\n    amount=item[{wrong_index}]\n    magnitude=abs(float(amount))\n    if magnitude>maximum:\n        maximum=magnitude\n        best=name\nresult=best"
    correct = code.replace(
        f"amount=item[{wrong_index}]", f"amount=item[{4 if flipped else 2}]"
    )
    return book.wire(), code, correct


@pytest.mark.parametrize("quoted", [True, False])
@pytest.mark.parametrize(
    "offset,left,flipped", [(0, 1, False), (300, 6, False), (507, 9, True)]
)
def test_loaded_requested_branch_does_not_mask_wrong_projected_metric(
    offset: int, left: int, flipped: bool, quoted: bool
) -> None:
    book, code, correct = fixture(offset, left, flipped)
    bad = execute_program(code, book)
    assert bad.success and bad.result == "Birch"
    original = copy.deepcopy(bad.reads)
    projected = column_projection(bad, code)
    assert projected is not None and len(projected.reads) == 6 and len(bad.reads) == 15
    query = 'Which entity had the highest absolute "Raw movement" from June to July?'
    if not quoted:
        query = query.replace('"', "")
    assert column_scope_errors(book, bad, query, code)
    assert not column_scope_errors(
        book, bad, query
    )  # Full rectangle alone conceals the error.
    assert bad.reads == original and projected.result == bad.result
    good = execute_program(correct, book)
    assert good.success and good.result == "ACME Inc."
    assert not column_scope_errors(book, good, query, correct)


@pytest.mark.parametrize(
    "mutation",
    [
        "alias=data\n",
        "data=list(data)\n",
        "data.append([])\n",
        "for other in data:\n    pass\n",
        "def unused(item):\n    return item[0]\n",
    ],
)
def test_escaping_rebinding_or_unknown_scopes_leave_full_evidence(
    mutation: str,
) -> None:
    book, code, _ = fixture()
    changed = code.replace("best=None\n", mutation + "best=None\n")
    execution = execute_program(changed, book)
    assert column_projection(execution, changed) is None


@pytest.mark.parametrize(
    "replacement",
    [
        "amount=item[index]",
        "amount=item[-1]",
        "amount=item[1:3]",
        "amount=tuple(item)[4]",
        "item[1]=0\n    amount=item[4]",
        "alias=item\n    amount=alias[4]",
    ],
)
def test_dynamic_indices_row_aliases_and_mutation_are_not_projected(
    replacement: str,
) -> None:
    book, code, _ = fixture()
    changed = code.replace("amount=item[4]", replacement)
    execution = execute_program(changed, book)
    assert column_projection(execution, changed) is None


def test_additional_reads_and_formula_dependencies_prevent_narrowing() -> None:
    book, code, _ = fixture()
    extra = code.replace("best=None", "outside=cell('Moved original','A1')\nbest=None")
    assert column_projection(execute_program(extra, book), extra) is None
    book["Moved original"]["cells"]["E3"] = {
        "value": None,
        "kind": "uncached_formula",
        "formula": "=G9",
    }
    book["Moved original"]["cells"]["G9"] = {"value": 17, "kind": "number"}
    assert column_projection(execute_program(code, book), code) is None


def test_hint_without_physical_caption_cannot_establish_a_contradiction() -> None:
    book, code, _ = fixture()
    bad = execute_program(code, book)
    book["Moved original"]["cells"]["B1"]["value"] = "Something different"
    query = 'Which entity had the highest absolute "Raw movement" from June to July?'
    assert not column_scope_errors(book, bad, query, code)


def test_derived_metric_may_differ_from_a_reported_quoted_column() -> None:
    book, code, _ = fixture()
    derived = code.replace(
        "magnitude=abs(float(amount))", "magnitude=abs(float(amount))*100"
    )
    execution = execute_program(derived, book)
    assert column_projection(execution, derived) is not None
    assert not column_scope_errors(
        book,
        execution,
        'Which entity has the highest "Raw movement" calculated in another scale?',
        derived,
    )


def test_blank_row_bands_with_same_effective_headers_preserve_column_identity() -> None:
    from src.excel_parser import CellData, SheetData, WorkbookData

    original, code, _ = fixture()
    data = {
        address: CellData(entry["value"], entry["kind"])
        for address, entry in original["Moved original"]["cells"].items()
    }
    for row in range(9, 12):
        data[f"A{row}"] = CellData(f"Noise-{row}", "string")
        for column in ["B", "C", "D", "E"]:
            data[f"{column}{row}"] = CellData(row, "number")
    book = WorkbookData(
        {
            "Moved original": SheetData(
                "Moved original", cells=data, merges=["B1:C1", "D1:E1"]
            )
        }
    ).wire()
    code = code.replace("A3:E5", "A3:E11").replace(
        "    magnitude=abs(float(amount))",
        "    if amount is None:\n        continue\n    magnitude=abs(float(amount))",
    )
    result = execute_program(code, book)
    assert result.success and len(book["Moved original"]["tables"]) == 2
    assert column_scope_errors(
        book, result, "Which entity has the highest absolute Raw movement?", code
    )
