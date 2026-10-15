"""Official routing, shared-call charging and failures must never borrow dev settings."""

import io
import json
import time
import urllib.error
from pathlib import Path
from typing import Any
from unittest.mock import patch

import openpyxl
import pytest

from benchmarks.batched_submission import run_batched_queries
from benchmarks.qwen_batched import QwenBatchedAgent
from src.llm_agent import DEFAULT_MODEL, BudgetExceeded, BudgetLedger, QueryState


def qualification(root: Path) -> None:
    from benchmarks.qwen_batch_signoff import source_manifest

    (root / "source_lock.json").write_text(json.dumps({"modules": source_manifest()}))
    for label, count in [
        ("public_1", 20),
        ("public_2", 20),
        ("public_3", 20),
        ("regression_60", 60),
        ("regression_32", 32),
        ("known_18", 18),
        ("synthetic_24", 24),
    ]:
        stage = root / label
        stage.mkdir()
        (stage / "report.json").write_text(
            json.dumps({"complete": True, "score": {"correct": count, "total": count}})
        )


def test_signoff_zero_credit_never_sends_inference(tmp_path: Path) -> None:
    from benchmarks.qwen_batch_signoff import signoff

    qualification(tmp_path)
    with (
        patch.dict("os.environ", {"OPENROUTER_API_KEY": "judge-placeholder"}),
        patch(
            "urllib.request.urlopen",
            return_value=io.BytesIO(b'{"data":{"total_credits":0,"total_usage":0.2}}'),
        ) as credit,
        patch.object(QwenBatchedAgent, "_open_request") as inference,
        patch("benchmarks.qwen_batch_signoff.get_pricing") as pricing,
    ):
        result = signoff(tmp_path)
    assert result["status"] == "blocked_credit" and result["new_generations"] == 0
    assert credit.call_args.args[0].full_url == "https://openrouter.ai/api/v1/credits"
    inference.assert_not_called()
    pricing.assert_not_called()
    with pytest.raises(ValueError, match="Preserve"):
        signoff(tmp_path)


def test_credit_resume_preserves_evidence_and_cannot_reset_paid_limit(
    tmp_path: Path,
) -> None:
    from benchmarks.qwen_batch_signoff import signoff

    qualification(tmp_path)
    original = tmp_path / "qwen_batch_signoff.json"
    original.write_text(json.dumps({"status": "blocked_credit", "new_generations": 0}))
    prior = original.read_bytes()
    with (
        patch.dict("os.environ", {"OPENROUTER_API_KEY": "test-placeholder"}),
        patch(
            "urllib.request.urlopen",
            return_value=io.BytesIO(b'{"data":{"total_credits":0,"total_usage":0.2}}'),
        ),
    ):
        report = signoff(tmp_path, attempt="restored-1")
    assert report["status"] == "blocked_credit" and original.read_bytes() == prior
    assert (tmp_path / "qwen_batch_signoff_restored-1.json").exists()
    original.write_text(
        json.dumps(
            {
                "status": "incomplete_official_check",
                "new_generations": 0,
                "http_attempts": 1,
            }
        )
    )
    with (
        patch("urllib.request.urlopen") as network,
        pytest.raises(ValueError, match="paid acceptance cap"),
    ):
        signoff(tmp_path, attempt="restored-2")
    network.assert_not_called()


def test_signoff_failed_or_stale_qualification_stops_before_network(
    tmp_path: Path,
) -> None:
    from benchmarks.qwen_batch_signoff import signoff

    qualification(tmp_path)
    (tmp_path / "public_2/report.json").write_text(
        json.dumps({"complete": True, "score": {"correct": 19, "total": 20}})
    )
    with patch("urllib.request.urlopen") as network:
        result = signoff(tmp_path)
    assert result["status"] == "blocked_qualification"
    network.assert_not_called()
    other = tmp_path / "changed"
    other.mkdir()
    qualification(other)
    source = json.loads((other / "source_lock.json").read_text())
    source["modules"]["src/executor.py"] = "changed"
    (other / "source_lock.json").write_text(json.dumps(source))
    with (
        patch("urllib.request.urlopen") as network,
        pytest.raises(ValueError, match="changed after lock"),
    ):
        signoff(other)
    network.assert_not_called()


def test_signoff_counts_shared_responses_once(tmp_path: Path) -> None:
    from benchmarks.qwen_batch_signoff import inference_counts

    call = {"request": 1, "shared_http_id": "live-group", "finish_reason": "stop"}
    rows = [
        {"planning_event": {"calls": [call]}},
        {"id": "one", "calls": [call]},
        {"id": "two", "calls": [call, {"request": 2, "error": "TimeoutError"}]},
    ]
    path = tmp_path / "trace.jsonl"
    path.write_text("\n".join(json.dumps(row) for row in rows))
    assert inference_counts(path) == {"http_attempts": 2, "model_responses": 1}


def response(content: dict[str, Any], usage: dict[str, Any] | None) -> io.BytesIO:
    return io.BytesIO(
        json.dumps(
            {
                "model": DEFAULT_MODEL,
                "choices": [
                    {
                        "message": {"content": json.dumps(content)},
                        "finish_reason": "stop",
                    }
                ],
                "usage": usage,
            }
        ).encode()
    )


def agent(**kwargs: Any) -> QwenBatchedAgent:
    return QwenBatchedAgent(api_key="judge-placeholder", pricing=(1e-7, 1e-7), **kwargs)


def workbook(path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "New scope"
    sheet.cell(306, 4, "Label")
    sheet.cell(306, 5, "Units")
    sheet.cell(307, 4, "One")
    sheet.cell(307, 5, 139)
    sheet.cell(308, 4, "Two")
    sheet.cell(308, 5, 73)
    book.save(path)
    book.close()


def test_qwen_group_uses_official_endpoint_and_charges_once(tmp_path: Path) -> None:
    files = tmp_path / "new inputs"
    files.mkdir()
    workbook(files / "renamed.xlsx")
    queries = [
        {
            "id": f"changed-{i}",
            "file_name": "renamed.xlsx",
            "query": f"What is Units for {label}? Answer only the integer.",
        }
        for i, label in enumerate(("One", "Two"))
    ]
    body = {
        "programs": {
            f"q{i}": {
                "code": f"result=num('New scope','E{i + 307}')",
                "format": {"kind": "integer"},
            }
            for i in range(2)
        }
    }
    ledger = BudgetLedger(0.05)
    client = agent(ledger=ledger)
    output = tmp_path / "new output" / "submission.csv"
    with patch.object(
        client,
        "_open_request",
        return_value=response(body, {"total_tokens": 101, "cost": "0.0001"}),
    ) as request:
        metrics = run_batched_queries(files, queries, output, client)
    assert not metrics["failures"] and metrics["answered"] == 2
    assert request.call_count == 1 and metrics["tokens"] == 101
    assert metrics["cost_usd"] == ledger.spent == pytest.approx(0.0001)
    req = request.call_args.args[0]
    payload = json.loads(req.data)
    assert req.full_url == "https://openrouter.ai/api/v1/chat/completions"
    assert req.get_header("Authorization") == "Bearer judge-placeholder"
    assert payload["model"] == DEFAULT_MODEL
    assert payload["temperature"] == 0 and payload["reasoning"] == {"max_tokens": 256}
    assert payload["response_format"] == {"type": "json_object"}
    assert "reasoning_effort" not in payload
    assert output.read_text().splitlines() == [
        "id,answer",
        "changed-0,139",
        "changed-1,73",
    ]
    rows = [
        json.loads(line)
        for line in output.with_suffix(".trace.jsonl").read_text().splitlines()
        if "planning_event" not in json.loads(line)
    ]
    assert sum(r["tokens"] for r in rows) == 101
    assert sum(r["cost"] for r in rows) == pytest.approx(0.0001)


def test_qwen_missing_usage_is_not_zero() -> None:
    client = agent()
    state = QueryState(time.monotonic() + 60, 0.09)
    with patch.object(
        client,
        "_open_request",
        return_value=response({"code": "result=num('S','B2')"}, None),
    ):
        client._request(
            [{"role": "user", "content": "Compute from the workbook"}], state
        )
    assert state.tokens > 0 and state.calls[0]["usage_missing"]
    assert (
        state.cost == pytest.approx(state.calls[0]["reserved_usd"]) and state.cost > 0
    )
    assert client.total_tokens_used == state.tokens


def test_qwen_timeouts_keep_cost_and_token_estimates() -> None:
    client = agent()
    state = QueryState(time.monotonic() + 60, 0.09)
    with (
        patch.object(client, "_open_request", side_effect=TimeoutError),
        patch("time.sleep"),
        pytest.raises(TimeoutError),
    ):
        client._request([{"role": "user", "content": "Compute from workbook"}], state)
    assert state.requests == 5 and len(state.calls) == 5
    assert state.tokens == sum(c["reserved_token_bound"] for c in state.calls) > 0
    assert state.cost == pytest.approx(sum(c["reserved_usd"] for c in state.calls))
    assert client.total_tokens_used == state.tokens


def test_qwen_budget_blocks_before_request() -> None:
    client = agent(ledger=BudgetLedger(1e-12))
    with (
        patch.object(client, "_open_request") as request,
        pytest.raises(BudgetExceeded),
    ):
        client._request(
            [{"role": "user", "content": "Read the workbook"}],
            QueryState(time.monotonic() + 60, 0.09),
        )
    request.assert_not_called()


def test_qwen_shared_retry_keeps_unknown_attempt_in_every_allocated_total(
    tmp_path: Path,
) -> None:
    from src.excel_parser import load_workbook_data

    path = tmp_path / "book.xlsx"
    workbook(path)
    queries = [
        {
            "id": str(i),
            "file_name": path.name,
            "query": f"How many Units for {label}? Answer only the integer.",
        }
        for i, label in enumerate(("One", "Two"))
    ]
    body = {
        "programs": {
            f"q{i}": {
                "code": f"result=num('New scope','E{i + 307}')",
                "format": {"kind": "integer"},
            }
            for i in range(2)
        }
    }
    client = agent(ledger=BudgetLedger(0.05))
    temporary = urllib.error.HTTPError(
        "https://openrouter.ai/api/v1/chat/completions",
        429,
        "busy",
        {},
        io.BytesIO(b"{}"),
    )
    with (
        patch.object(
            client,
            "_open_request",
            side_effect=[
                temporary,
                response(body, {"total_tokens": 101, "cost": 0.0001}),
            ],
        ) as request,
        patch("time.sleep"),
    ):
        rows, event = client.solve_group(load_workbook_data(path), queries)
    assert request.call_count == 2 and len(event["calls"]) == 2
    assert [row["answer"] for row in rows] == ["139", "73"]
    assert all(row["requests"] == 2 for row in rows)
    assert sum(row["tokens"] for row in rows) == client.total_tokens_used > 101
    assert sum(row["cost"] for row in rows) == pytest.approx(client.total_cost)
    assert client.total_cost > 0.0001 and event["calls"][0]["usage_missing"]


@pytest.mark.parametrize("code", [401, 402, 403])
def test_qwen_auth_failure_has_no_paid_fallback(code: int) -> None:
    client = agent()
    error = urllib.error.HTTPError(
        "https://openrouter.ai/api/v1/chat/completions",
        code,
        "denied",
        {},
        io.BytesIO(b"{}"),
    )
    with (
        patch.object(client, "_open_request", side_effect=error) as request,
        pytest.raises(ValueError),
    ):
        client._request(
            [{"role": "user", "content": "Read workbook"}],
            QueryState(time.monotonic() + 60, 0.09),
        )
    assert request.call_count == 1 and client.fatal_error


def test_qwen_overcharge_stops_further_calls() -> None:
    client = agent()
    with patch.object(
        client,
        "_open_request",
        return_value=response({"code": "result=1"}, {"total_tokens": 10, "cost": 0.02}),
    ) as request:
        with pytest.raises(BudgetExceeded):
            client._request(
                [{"role": "user", "content": "Read workbook"}],
                QueryState(time.monotonic() + 60, 0.09),
            )
        with pytest.raises(ValueError, match="Billed cost"):
            client._request(
                [{"role": "user", "content": "Read workbook"}],
                QueryState(time.monotonic() + 60, 0.09),
            )
    assert request.call_count == 1 and client.total_cost == pytest.approx(0.02)


def test_judge_never_uses_dev_key_or_model() -> None:
    with (
        patch.dict(
            "os.environ",
            {"DEV_LLM_API_KEY": "development-placeholder", "OPENROUTER_API_KEY": ""},
        ),
        pytest.raises(ValueError, match="OPENROUTER_API_KEY"),
    ):
        QwenBatchedAgent(pricing=(1e-7, 1e-7))
    with pytest.raises(ValueError, match="competition Qwen"):
        QwenBatchedAgent(
            api_key="judge-placeholder", model="different-model", pricing=(1e-7, 1e-7)
        )
