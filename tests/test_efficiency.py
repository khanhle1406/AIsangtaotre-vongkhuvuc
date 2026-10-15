"""Compact context must preserve values, coordinates, metadata and safety limits."""

import datetime as dt
import json
import re
from pathlib import Path
from typing import Any
from unittest.mock import patch

import openpyxl
import pytest
from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries

from benchmarks.efficient import EfficientAgent, dense_text, metadata_ranges
from benchmarks.gateway import BatchLimits
from src.excel_parser import CellData, SheetData, WorkbookData, load_workbook_data
from src.llm_agent import SolveResult


def test_batched_envelope_rejects_missing_duplicate_and_extra_indices() -> None:
    from benchmarks.batched import parse_plans

    good = {"index": 0, "code": "result=num('Data','B2')"}
    assert len(parse_plans(json.dumps({"programs": [good]}), 1)) == 1
    echo = json.dumps({"programs": [{**good, "query": "Read units"}]})
    assert len(parse_plans(echo, 1, ["Read units"])) == 1
    with pytest.raises(ValueError, match="disagrees"):
        parse_plans(echo, 1, ["Read costs"])
    for entries in (
        [good, good],
        [good],
        [{**good, "index": True}, {**good, "index": 1}],
    ):
        with pytest.raises(ValueError):
            parse_plans(json.dumps({"programs": entries}), 2)


def test_shared_tokens_are_allocated_exactly_once() -> None:
    from benchmarks.batched import apportioned_calls

    original = [
        {
            "tokens": 101,
            "provider_total_tokens": 101,
            "reserved_token_bound": 200,
            "shared_http_id": "live-event",
            "cost_usd": 0.006,
        }
    ]
    divided = [apportioned_calls(original, index, 3)[0] for index in range(3)]
    assert sum(c["tokens"] for c in divided) == 101
    assert sum(c["provider_total_tokens"] for c in divided) == 101
    assert sum(c["reserved_token_bound"] for c in divided) == 200
    assert sum(c["cost_usd"] for c in divided) == pytest.approx(0.006)
    assert {c["shared_http_id"] for c in divided} == {"live-event"}


def test_structured_keys_cannot_be_silently_overwritten() -> None:
    from benchmarks.structured_batch import unique_object

    with pytest.raises(ValueError, match="Duplicate"):
        json.loads(
            '{"programs":{"q0":{},"q0":{},"q1":{}}}', object_pairs_hook=unique_object
        )
    assert (
        json.loads('{"q0":{"code":"result=1"}}', object_pairs_hook=unique_object)["q0"][
            "code"
        ]
        == "result=1"
    )


@pytest.mark.parametrize("offset", [0, 310])
def test_extremum_cannot_skip_intermediate_category(
    tmp_path: Path, offset: int
) -> None:
    from src.executor import execute_program
    from src.verification import complete_category_contract

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "New scope"
    for i, row in enumerate(
        [
            ["Category", "Share"],
            ["All categories", 100],
            ["Small group", 17],
            ["Broad group", 72],
            ["Narrow group", 41],
        ],
        offset + 1,
    ):
        for c, value in enumerate(row, 1):
            sheet.cell(i, c, value)
    path = tmp_path / "changed.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    before = execute_program(
        f"v=max([(num('New scope','B{offset + 3}'),'A{offset + 3}'),(num('New scope','B{offset + 5}'),'A{offset + 5}')]); result=cell('New scope',v[1])",
        wire,
    )
    query = "Which category has the highest share? Answer only the category."
    repair = complete_category_contract(wire, before, query)
    assert repair is not None
    assert execute_program(repair, wire).result == "Broad group"
    assert (
        complete_category_contract(
            wire,
            before,
            "Which category among Small group and Narrow group has the highest share?",
        )
        is None
    )


def test_difference_must_read_both_named_measures(tmp_path: Path) -> None:
    from src.executor import execute_program
    from src.verification import row_scope_errors

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Relocated"
    for i, row in enumerate(
        [
            ["Measure", "Amount"],
            ["Available capacity", 179],
            ["Used capacity", 61],
            ["Reserved capacity", 54],
        ],
        301,
    ):
        for c, value in enumerate(row, 4):
            sheet.cell(i, c, value)
    path = tmp_path / "unseen.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    question = (
        "What is the absolute difference between Available capacity and Used capacity?"
    )
    wrong = execute_program(
        "result=abs(num('Relocated','E303')-num('Relocated','E304'))", wire
    )
    assert any(
        "available capacity" in error
        for error in row_scope_errors(wire, wrong.reads, question)
    )
    correct = execute_program(
        "result=abs(num('Relocated','E302')-num('Relocated','E303'))", wire
    )
    assert not row_scope_errors(wire, correct.reads, question)


def test_empty_type_filter_does_not_silently_erase_typed_dates(tmp_path: Path) -> None:
    from src.executor import execute_program

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Typed events"
    sheet.append(["Date", "Team", "Amount"])
    sheet.append([dt.datetime(2028, 4, 9), "North", 47])
    path = tmp_path / "dates.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    wrong = execute_program(
        "r=data_rows('Typed events',{'A':'string','B':'string','C':'number'}); result=sum(num('Typed events','C'+str(i)) for i in r)",
        wire,
    )
    assert not wrong.success and "datetime" in wrong.error
    correct = execute_program(
        "r=data_rows('Typed events',{'A':'datetime','B':'string','C':'number'}); result=sum(num('Typed events','C'+str(i)) for i in r)",
        wire,
    )
    assert correct.success and correct.result == 47
    empty = execute_program(
        "result=sum(num('Typed events','C'+str(i)) for i in data_rows('Typed events',{'A':'datetime','C':'number'}) if cell('Typed events','B'+str(i))=='Absent')",
        wire,
    )
    assert empty.success and empty.result == 0


def test_age_helper_sort_is_audited_like_explicit_sort(tmp_path: Path) -> None:
    from src.executor import execute_program
    from src.verification import age_sort_errors

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Admissions"
    sheet.append(["Name", "Birth date", "Admission date", "Score"])
    for name, birth, event, score in [
        ("A", dt.datetime(1988, 2, 1), dt.datetime(2010, 2, 1), 41),
        ("B", dt.datetime(1991, 2, 1), dt.datetime(2026, 2, 1), 73),
        ("C", dt.datetime(1998, 2, 1), dt.datetime(2025, 2, 1), 58),
    ]:
        sheet.append([name, birth, event, score])
    path = tmp_path / "people.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    query = "Who was the youngest at their admission date? Answer only the name."
    wrong = execute_program(
        "r=age_order('Admissions','B','C',descending=True)[0]; result=cell('Admissions','A'+str(r))",
        wire,
    )
    assert wrong.success and wrong.sorts
    assert age_sort_errors(wire, wrong, query)
    right = execute_program(
        "r=age_order('Admissions','B','C',descending=False)[0]; result=cell('Admissions','A'+str(r))",
        wire,
    )
    assert right.success and right.result == "A"
    assert not age_sort_errors(wire, right, query)


def test_live_shared_plan_executes_each_question_and_resumes(tmp_path: Path) -> None:
    import io

    from benchmarks.batched import run_batched_stage
    from benchmarks.gateway import GatewayAgent
    from src.llm_agent import BudgetLedger

    files = tmp_path / "xlsx_files"
    files.mkdir()
    book = openpyxl.Workbook()
    book.active.title = "Renamed"
    book.active.append(["Label", "Count"])
    book.active.append(["Alpha", 139])
    book.active.append(["Beta", 73])
    book.save(files / "unfamiliar.xlsx")
    book.close()
    queries = [
        {
            "id": str(i),
            "file_name": "unfamiliar.xlsx",
            "query": f"How many for {label}? Answer only the integer.",
        }
        for i, label in enumerate(("Alpha", "Beta"))
    ]
    refs = tmp_path / "gold.csv"
    refs.write_text("id,answer\n0,139\n1,73\n")
    body = {
        "programs": [
            {
                "index": i,
                "code": f"result=num('Renamed','B{i + 2}')",
                "format": {"kind": "integer"},
            }
            for i in range(2)
        ]
    }

    def response(*_args: Any, **_kwargs: Any) -> io.BytesIO:
        result = io.BytesIO(
            json.dumps(
                {
                    "model": "gemini-3-flash",
                    "choices": [
                        {
                            "message": {"content": json.dumps(body)},
                            "finish_reason": "stop",
                        }
                    ],
                    "usage": {"total_tokens": 101},
                }
            ).encode()
        )
        result.headers = {"Content-Type": "application/json"}
        return result

    environment = {
        "DEV_LLM_BASE_URL": "https://gateway.invalid/v1",
        "DEV_LLM_API_KEY": "test-placeholder",
        "DEV_LLM_MODEL": "ag/gemini-3-flash",
    }
    limits = BatchLimits(tmp_path / "batch.json")
    with (
        patch.dict("os.environ", environment),
        patch.object(GatewayAgent, "_open_request", side_effect=response) as request,
    ):
        first = run_batched_stage(
            tmp_path / "out", files, queries, refs, limits, BudgetLedger(0.01)
        )
        resumed = run_batched_stage(
            tmp_path / "out", files, queries, refs, limits, BudgetLedger(0.01)
        )
    assert first["score"]["correct"] == resumed["score"]["correct"] == 2
    assert first["http_attempts"] == request.call_count == 1
    assert first["tokens"] == first["provider_tokens_known"] == 101
    assert first["usage_missing_attempts"] == 0
    assert first["peak_total_rss_bytes"] == resumed["peak_total_rss_bytes"]
    from benchmarks.structured_batch import StructuredBatchedAgent

    body = {
        "programs": {
            f"q{i}": {
                "code": f"result=num('Renamed','B{i + 2}')",
                "format": {"kind": "integer"},
            }
            for i in range(2)
        }
    }
    with (
        patch.dict("os.environ", environment),
        patch.object(GatewayAgent, "_open_request", side_effect=response) as request,
    ):
        keyed = run_batched_stage(
            tmp_path / "keyed",
            files,
            queries,
            refs,
            limits,
            BudgetLedger(0.01),
            agent_class=StructuredBatchedAgent,
        )
    assert keyed["score"]["correct"] == 2 and keyed["tokens"] == 101
    assert (
        json.loads(request.call_args.args[0].data)["response_format"]["type"]
        == "json_schema"
    )
    # Restore the classic mock response before testing the default CLI.
    body = {
        "programs": [
            {
                "index": i,
                "code": f"result=num('Renamed','B{i + 2}')",
                "format": {"kind": "integer"},
            }
            for i in range(2)
        ]
    }
    from benchmarks.batched_cli import solve
    from benchmarks.budget import DevelopmentBudget

    queries_path = tmp_path / "renamed queries.json"
    queries_path.write_text(json.dumps(queries))
    with (
        patch.dict("os.environ", environment),
        patch.object(GatewayAgent, "_open_request", side_effect=response) as request,
        patch(
            "benchmarks.batched_cli.workspace_budget",
            return_value=DevelopmentBudget(tmp_path / "ledger.json", 0),
        ),
        patch(
            "benchmarks.batched.read_answers",
            side_effect=AssertionError("CLI must not read an oracle"),
        ),
    ):
        metrics = solve(
            files, queries_path, tmp_path / "output with spaces" / "answers.csv"
        )
    assert not metrics["failures"] and metrics["answered"] == 2
    assert metrics["tokens"] == 101 and request.call_count == 1
    with pytest.raises(ValueError, match="overwrite"):
        solve(files, queries_path, queries_path)


def decode_runs(text: str) -> dict[str, dict[str, Any]]:
    """Independent test decoder verifies a complete round trip, not string shape."""
    cells: dict[str, dict[str, Any]] = {}
    for line in text.splitlines():
        if line.startswith("CELL METADATA "):
            for group in json.loads(line.split(") ", 1)[1]):
                for area in group["ranges"]:
                    left, top, right, bottom = range_boundaries(area)
                    for row in range(top, bottom + 1):
                        for col in range(left, right + 1):
                            cells[f"{get_column_letter(col)}{row}"].update(
                                group["attributes"]
                            )
        elif re.match(r"^[A-Z]+[1-9]\d*:\[", line):
            anchor, raw = line.split(":", 1)
            row, col = coordinate_to_tuple(anchor)
            for offset, value in enumerate(json.loads(raw)):
                kind = {
                    str: "string",
                    int: "number",
                    float: "number",
                    bool: "boolean",
                }.get(type(value))
                cells[f"{get_column_letter(col + offset)}{row}"] = {
                    "value": value,
                    "kind": kind,
                    "number_format": "General",
                    "formula": None,
                    "indent": 0,
                    "bold": False,
                }
    return cells


def test_dense_context_round_trip_preserves_typed_cells(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Unfamiliar "section"'
    sheet.append(["Group", "Rate", "Event", "Flag", "Formula"])
    sheet.append(
        ["ACME Inc.\ncontinued", 0.325, dt.datetime(2027, 2, 11), True, "=B2*2"]
    )
    sheet["B2"].number_format = "0.0%"
    sheet["A2"].alignment = openpyxl.styles.Alignment(indent=2)
    sheet["A1"].font = openpyxl.styles.Font(bold=True)
    sheet.merge_cells("A4:C4")
    sheet["A4"] = -15
    sheet.row_dimensions[5].hidden = True
    sheet.column_dimensions["D"].hidden = True
    path = tmp_path / "table.xlsx"
    book.save(path)
    book.close()
    parsed = load_workbook_data(path)
    wire = parsed.wire()
    text = dense_text(parsed, wire)
    assert decode_runs(text) == wire[sheet.title]["cells"]
    hints = json.loads(
        next(line[6:] for line in text.splitlines() if line.startswith("HINTS "))
    )
    assert hints["merges"] == ["A4:C4"]
    assert hints["hidden_rows"] == [5]
    assert hints["hidden_columns"] == wire[sheet.title]["hidden_columns"]
    assert hints["headers"] == wire[sheet.title]["headers"]
    assert hints["tables"] == wire[sheet.title]["tables"]


def test_dense_context_sparse_extent_and_grouped_styles() -> None:
    sheet = SheetData("Sparse")
    sheet.cells = {
        "A1": CellData("Title", "string"),
        "XFD1048576": CellData(-7, "number"),
    }
    book = WorkbookData({"Sparse": sheet})
    wire = book.wire()
    text = dense_text(book, wire)
    assert len(text) < 2500
    assert decode_runs(text) == wire["Sparse"]["cells"]
    assert metadata_ranges(["B8", "B9", "B11", "D6"]) == ["B8:B9", "B11", "D6"]


def test_efficiency_agent_keeps_full_backend_and_existing_solver(
    tmp_path: Path,
) -> None:
    sheet = SheetData("Shifted", cells={"D306": CellData(139, "number")})
    workbook = WorkbookData({"Shifted": sheet})
    agent = EfficientAgent(
        api_key="test-placeholder",
        gateway_url="https://gateway.invalid/v1",
        gateway_model="ag/gemini-3-flash",
        batch_limits=BatchLimits(tmp_path / "batch.json"),
    )
    with patch.object(agent, "_solve", return_value=SolveResult()) as solve:
        agent.solve_workbook(workbook, "Read the audited value")
    text, wire, _, _ = solve.call_args.args
    assert "D306:[139]" in text
    assert wire["Shifted"]["cells"]["D306"]["value"] == 139
    assert agent.iterator_helpers and agent.model == "ag/gemini-3-flash"
