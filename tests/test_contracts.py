"""Task contracts use source structure and arithmetic intent, never sample answers."""

import datetime as dt
import json
from pathlib import Path

import openpyxl
import pytest

from src.contracts import (
    extremum_evidence_errors,
    header_count_errors,
    scale_conversion_errors,
    signed_comparison_errors,
)
from src.excel_parser import load_workbook_data
from src.executor import execute_program
from src.verification import scope_risks


def test_ranked_labels_stay_bound_to_source_rows(tmp_path: Path) -> None:
    from src.verification import ranked_label_binding_errors

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Moved"
    for address, value in {
        "F507": "Category",
        "G507": "Units",
        "F508": "Birch",
        "G508": 17,
        "F509": "Cedar",
        "G509": 41,
        "F510": "Aster",
        "G510": 9,
    }.items():
        sheet[address] = value
    path = tmp_path / "renamed.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    code = (
        "names=['Cedar','Birch','Aster']; "
        "v=values('Moved','G508:G510'); "
        "pairs=sorted(zip(names,v),key=lambda p:p[1],reverse=True); "
        "result=[p[0] for p in pairs]"
    )
    query = "Rank the categories by Units in descending order."
    bad = execute_program(code, wire)
    assert bad.success and ranked_label_binding_errors(wire, bad, query)
    good = execute_program(
        code.replace("['Cedar','Birch','Aster']", "['Birch','Cedar','Aster']"), wire
    )
    assert good.success and not ranked_label_binding_errors(wire, good, query)
    ascending = execute_program(code.replace("reverse=True", "reverse=False"), wire)
    assert ranked_label_binding_errors(
        wire, ascending, query.replace("descending", "ascending")
    )
    # A transformed objective is not proven by matching raw source measures.
    transformed = execute_program(code.replace("p[1],reverse", "-p[1],reverse"), wire)
    assert not ranked_label_binding_errors(wire, transformed, query)
    assert not ranked_label_binding_errors(wire, bad, "List the category names.")
    duplicate = dict(wire["Moved"]["cells"]["F509"])
    wire["Moved"]["cells"]["F510"] = duplicate
    assert not ranked_label_binding_errors(wire, bad, query)


def test_failed_row_unpack_can_use_width_matched_cells_without_another_model_call() -> (
    None
):
    from unittest.mock import patch

    from src.contracts import row_unpack_contract
    from src.excel_parser import CellData, SheetData, WorkbookData
    from src.llm_agent import LLMAgent

    book = WorkbookData(
        {
            "Moved": SheetData(
                "Moved",
                cells={
                    "F508": CellData(17, "number"),
                    "G508": CellData("Birch", "string"),
                    "F509": CellData(29, "number"),
                    "G509": CellData("Cedar", "string"),
                },
            )
        }
    )
    code = "pairs=values('Moved','F508:G509'); amounts={name:n for n,name in pairs}; answer=amounts['Birch']"
    agent = LLMAgent(api_key="test-placeholder", pricing=(0, 0))
    with patch.object(
        agent,
        "_request",
        return_value=json.dumps(
            {"code": code, "result_variable": "answer", "format": {"kind": "integer"}}
        ),
    ) as request:
        result = agent.solve_workbook(
            book, "What is Units for Birch? Answer only the integer."
        )
    assert result.answer == "17" and result.error is None and request.call_count == 1
    assert result.attempts[0]["error"] and result.attempts[0]["code"] == code
    assert result.attempts[1]["contract_kind"] == "row-unpack-shape"
    assert "cells(" in result.attempts[1]["code"]
    assert set(result.attempts[0]["reads"]) == set(result.attempts[1]["reads"])
    for rejected in [
        code + "; unused=sum(pairs)",
        code.replace("for n,name in pairs", "for n,name,other in pairs"),
        code.replace("F508:G509", "F508:F509"),
        "cells=17; " + code,
    ]:
        assert (
            row_unpack_contract(rejected, book.wire(), "Units", failure_line=1) is None
        )
    assert (
        row_unpack_contract(code, book.wire(), "Extract the letters", failure_line=1)
        is None
    )
    assert row_unpack_contract(code, book.wire(), "Units", failure_line=99) is None
    text_wire = book.wire()
    for address in text_wire["Moved"]["cells"]:
        text_wire["Moved"]["cells"][address].update(value="AB", kind="string")
    assert row_unpack_contract(code, text_wire, "Units", failure_line=1) is None


def test_unbounded_string_header_domain_requests_review_but_preserves_data_rows(
    tmp_path: Path,
) -> None:
    import ast

    from src.verification import header_domain_risks

    book = openpyxl.Workbook()
    sheet = book.active
    sheet["G407"], sheet["G408"], sheet["G409"] = "Category", "Birch", "Cedar"
    path = tmp_path / "categories.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    code = "r=data_rows('Sheet',{'G':'string'}); names=[cell('Sheet','G'+str(i)) for i in r]; result=len(names)"
    execution = execute_program(code, wire)
    assert execution.result == 3  # Existing type-filter interface is unchanged.
    query = "How many required categories are listed?"
    assert header_domain_risks(wire, execution, query, ast.parse(code))
    bounded = code.replace(
        "{'G':'string'})", "{'G':'string'},start_row=408,end_row=409)"
    )
    good = execute_program(bounded, wire)
    assert good.result == 2 and not header_domain_risks(
        wire, good, query, ast.parse(bounded)
    )
    filtered = code.replace("for i in r]", "for i in r if i>407]")
    checked = execute_program(filtered, wire)
    assert checked.result == 2 and not header_domain_risks(
        wire, checked, query, ast.parse(filtered)
    )
    assert not header_domain_risks(
        wire, execution, "Count categories including headers.", ast.parse(code)
    )
    # Even matching an inferred caption cannot prove that a measure-bearing
    # first row is a heading: keep it as a genuine record.
    wire["Sheet"]["cells"]["H407"] = {
        **wire["Sheet"]["cells"]["G407"],
        "value": 17,
        "kind": "number",
    }
    wire["Sheet"]["headers"] = {"G": ["Category"]}
    wire["Sheet"]["tables"] = [{"range": "G407:H409", "headers": {"G": ["Category"]}}]
    assert not header_domain_risks(wire, execution, query, ast.parse(code))


def test_scalar_shape_errors_identify_operation_and_fix_source_helper() -> None:
    from src.excel_parser import CellData, SheetData, WorkbookData

    wire = WorkbookData(
        {
            "Shifted": SheetData(
                "Shifted",
                cells={
                    "H503": CellData(17, "number"),
                    "I503": CellData("Birch", "string"),
                    "H504": CellData(29, "number"),
                    "I504": CellData("Cedar", "string"),
                },
            )
        }
    ).wire()
    for code, fragment in [
        (
            "result={name:amount for amount,name in values('Shifted','H503:I504')}",
            "Cannot unpack scalar int",
        ),
        (
            "v=values('Shifted','H503:I504'); result=v[0][1]",
            "Cannot index scalar int in v[0][1]",
        ),
    ]:
        result = execute_program(code, wire)
        assert not result.success and fragment in result.error
        assert "line 1" in result.error and "cells(sheet,range)" in result.error
    good = execute_program(
        "d={name:amount for amount,name in cells('Shifted','H503:I504')}; result=d['Birch']",
        wire,
    )
    assert good.success and good.result == 17
    assert execute_program("result=values('Shifted','H503:H504')[1]", wire).result == 29


def test_flat_string_projection_requests_review_without_banning_initials(
    tmp_path: Path,
) -> None:
    from unittest.mock import patch

    from src.llm_agent import LLMAgent

    book = openpyxl.Workbook()
    sheet = book.active
    for row, cells in enumerate(
        [["Category", "Units"], ["Birch", 17], ["Cedar", 29]], 407
    ):
        for column, value in enumerate(cells, 6):
            sheet.cell(row, column, value)
    path = tmp_path / "flat.xlsx"
    book.save(path)
    book.close()
    workbook = load_workbook_data(path)
    wrong = "names=[c[0] for c in values('Sheet','F408:F409')]; result=sum(num('Sheet','G'+str(r)) for r in [408,409] if cell('Sheet','F'+str(r)) in names)"
    correct = wrong.replace("[c[0] for c", "[c for c")
    query = "What is combined Units for the listed categories? Answer only the integer."
    execution = execute_program(wrong, workbook.wire())
    assert execution.success and execution.result == 0
    assert any(
        "Flat values()" in risk
        for risk in scope_risks(workbook.wire(), execution, query, wrong)
    )
    assert not scope_risks(
        workbook.wire(), execution, "Count names whose first character is B.", wrong
    )
    agent = LLMAgent(api_key="test-placeholder", pricing=(0, 0), review_candidate=True)
    with patch.object(
        agent,
        "_request",
        side_effect=[json.dumps({"code": wrong}), json.dumps({"code": correct})],
    ) as request:
        result = agent.solve_workbook(workbook, query)
    assert request.call_count == 2 and result.answer == "46" and result.error is None


def test_list_insert_keeps_python_order_and_worker_item_limit() -> None:
    from unittest.mock import patch

    from src.executor import Interpreter

    result = execute_program(
        "p=['End']; p.insert(0,'Start'); p.insert(-1,'Middle'); result=p"
    )
    assert result.success and result.result == ["Start", "Middle", "End"]
    interpreter = Interpreter({})
    items = [1, 2]
    with (
        patch("src.executor.MAX_ITEMS", 2),
        pytest.raises(MemoryError, match="item limit"),
    ):
        interpreter.insert_value(items, 0, 3)
    assert items == [1, 2]


def test_set_update_and_relations_keep_iterable_semantics_and_caps() -> None:
    from unittest.mock import patch

    from src.executor import Interpreter

    result = execute_program(
        "s={1}; s.update([2,2], (3,)); result=[sorted(s), {1,2}.issubset([1,2,3]), s.issuperset([1,3]), s.isdisjoint([4]), s.isdisjoint([3])] "
    )
    assert result.success and result.result == [[1, 2, 3], True, True, True, False]
    interpreter = Interpreter({})
    with patch("src.executor.MAX_ITEMS", 2):
        target = {1, 2}
        interpreter.update_set(target, [2, 1])
        with pytest.raises(MemoryError, match="item limit"):
            interpreter.update_set(target, [3])
        with pytest.raises(MemoryError, match="item limit"):
            interpreter.set_relation({1}, "issubset", iter([1, 2, 3]))
        assert target == {1, 2}


def test_typed_date_iso_equality_cannot_silently_filter_every_record() -> None:
    from src.excel_parser import CellData, SheetData, WorkbookData

    wire = WorkbookData(
        {
            "Events": SheetData(
                "Events", cells={"G408": CellData(dt.datetime(2031, 4, 19), "datetime")}
            )
        }
    ).wire()
    for expression in ["d=='2031-04-19T00:00:00'", "'2031-04-19'!=d"]:
        checked = execute_program("d=cell('Events','G408'); result=" + expression, wire)
        assert (
            not checked.success and "Typed date compared with ISO text" in checked.error
        )
    good = execute_program(
        "d=cell('Events','G408'); result=[d==datetime.datetime(2031,4,19),d.isoformat()=='2031-04-19T00:00:00',d.date().isoformat()=='2031-04-19',d is None]",
        wire,
    )
    assert good.success and good.result == [True, True, True, False]


def test_bit_masks_and_augmented_shifts_obey_integer_limits() -> None:
    from src.executor import bounded_left_shift

    result = execute_program(
        "n=1; n<<=4; result=[n, n>>2, [i for i in range(1<<3) if i & (1<<1)]]"
    )
    assert result.success and result.result == [16, 4, [2, 3, 6, 7]]
    assert bounded_left_shift(1, 4095).bit_length() == 4096
    for count in [4096, 1000000000]:
        with pytest.raises(MemoryError, match="bit limit"):
            bounded_left_shift(1, count)
    with pytest.raises(ValueError, match="negative shift"):
        bounded_left_shift(1, -1)
    too_large = execute_program("n=1; n<<=1000000000; result=n")
    assert not too_large.success and "bit limit" in too_large.error


def test_quoted_parent_metric_reviews_unrequested_restricted_child(
    tmp_path: Path,
) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Shifted metrics"
    sheet["E307"], sheet["F307"], sheet["G307"] = "Entity", "Average duration", None
    sheet.merge_cells("F307:G307")
    sheet["F308"], sheet["G308"] = "Overall visits", "Overnight visits"
    sheet["E309"], sheet["F309"], sheet["G309"] = "Birch", 17.4, 32.9
    sheet["E310"], sheet["F310"], sheet["G310"] = "Cedar", 19.2, 34.7
    path = tmp_path / "scopes.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    query = 'What is "Average duration" for Birch? Answer to one decimal place.'
    child_code = "result=num('Shifted metrics','G309')"
    child = execute_program(child_code, wire)
    overall_code = "result=num('Shifted metrics','F309')"
    overall = execute_program(overall_code, wire)
    assert any(
        "restricted children" in risk
        for risk in scope_risks(wire, child, query, child_code)
    )
    assert not scope_risks(wire, overall, query, overall_code)
    mixed = "result=[num('Shifted metrics','F309'),num('Shifted metrics','G310')]"
    both = query.replace("for Birch", "for Birch and Cedar")
    assert scope_risks(wire, execute_program(mixed, wire), both, mixed)
    correct = mixed.replace("G310", "F310")
    assert not scope_risks(wire, execute_program(correct, wire), both, correct)
    # Reading both siblings in each row does not establish a restricted-only record.
    full = "result=values('Shifted metrics','F309:G310')"
    assert not scope_risks(wire, execute_program(full, wire), both, full)
    from unittest.mock import patch

    from src.llm_agent import LLMAgent

    agent = LLMAgent(api_key="test-placeholder", pricing=(0, 0), review_candidate=True)
    with patch.object(
        agent,
        "_request",
        side_effect=[
            json.dumps(
                {"code": child_code, "format": {"kind": "decimal", "decimals": 1}}
            ),
            json.dumps(
                {"code": overall_code, "format": {"kind": "decimal", "decimals": 1}}
            ),
        ],
    ) as requests:
        solved = agent.solve_workbook(load_workbook_data(path), query)
    assert requests.call_count == 2 and solved.answer == "17.4" and solved.error is None
    assert solved.attempts[0]["verification_mode"] == "candidate-audit"
    restricted = query.replace("for Birch", "for Birch's overnight visits")
    assert not scope_risks(wire, child, restricted, child_code)
    # Without explicit named-parent evidence this check makes no scope assumption.
    assert not scope_risks(
        wire, child, "What is the overnight duration for Birch?", child_code
    )
    # Vertical subgroup headings may define the identity column only; the metric
    # columns still inherit the sheet's original parent/child paths.
    for table in wire["Shifted metrics"]["tables"]:
        table["headers"] = {"E": ["Entity"]}
    assert any(
        "restricted children" in risk
        for risk in scope_risks(wire, child, query, child_code)
    )
    for table in wire["Shifted metrics"]["tables"]:
        table["headers"].update({"F": ["Other metric"], "G": ["Unrelated metric"]})
    assert not scope_risks(wire, child, query, child_code)


def test_named_comparison_requires_both_actual_numeric_header_columns(
    tmp_path: Path,
) -> None:
    from src.verification import column_scope_errors

    book = openpyxl.Workbook()
    sheet = book.active
    for row, values in enumerate(
        [["Group", "Manufacturing", "Construction", "Other"], ["Young", 37, 61, 89]],
        407,
    ):
        for column, value in enumerate(values, 4):
            sheet.cell(row, column, value)
    path = tmp_path / "permuted.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    question = "For Young, what is the sum of Construction and Manufacturing?"
    wrong = execute_program("result=num('Sheet','F408')+num('Sheet','G408')", wire)
    right = execute_program("result=num('Sheet','E408')+num('Sheet','F408')", wire)
    assert column_scope_errors(wire, wrong, question)
    assert not column_scope_errors(wire, right, question)


def test_cross_sheet_difference_rank_requests_independent_semantic_review(
    tmp_path: Path,
) -> None:
    book = openpyxl.Workbook()
    charges = book.active
    charges.title = "Charges"
    charges.append(["Item", "Charge"])
    charges.append(["Rotor", 57])
    settled = book.create_sheet("Settlements")
    settled.append(["Item", "Paid"])
    settled.append(["Rotor", 13])
    path = tmp_path / "balances.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    code = (
        "owed=num('Charges','B2')-num('Settlements','B2'); result=cell('Charges','A2')"
    )
    execution = execute_program(code, wire)
    risks = scope_risks(
        wire, execution, "Which item has largest remaining balance?", code
    )
    assert any("subtraction direction" in risk for risk in risks)
    assert not scope_risks(wire, execution, "What is the Rotor name?", code)
    assert not scope_risks(
        wire,
        execution,
        "Each step adds Charge and subtracts Paid. Which item has largest remaining balance?",
        code,
    )


def test_flat_matrix_aggregate_can_finish_after_two_runtime_repairs(
    tmp_path: Path,
) -> None:
    from unittest.mock import patch

    from src.llm_agent import LLMAgent

    book = openpyxl.Workbook()
    sheet = book.active
    for row, values in enumerate(
        [
            ["Actual", "Predicted", "Count"],
            ["North", "North", 17],
            ["North", "East", 3],
            ["East", "North", 5],
            ["East", "East", 13],
        ],
        402,
    ):
        for col, value in enumerate(values, 6):
            sheet.cell(row, col, value)
    path = tmp_path / "counts.xlsx"
    book.save(path)
    book.close()
    body = "tp=sum(x[2] for x in data if x[0]=='North' and x[1]=='North')\nfp=sum(x[2] for x in data if x[0]!='North' and x[1]=='North')\nfn=sum(x[2] for x in data if x[0]=='North' and x[1]!='North')\nresult=100*2*tp/(2*tp+fp+fn)"
    programs = [
        {"code": code, "format": {"kind": "decimal", "decimals": 2}}
        for code in [
            "result =",
            "data=values('Sheet','F403:H406')\n" + body,
            "data=cells('Sheet','F403:H406')\n" + body,
        ]
    ]
    client = LLMAgent(api_key="test-placeholder", pricing=(0, 0))
    with patch.object(
        client, "_request", side_effect=[json.dumps(p) for p in programs]
    ) as call:
        result = client.solve_workbook(
            load_workbook_data(path),
            "Calculate North F1 = 2*TP/(2*TP+FP+FN), expressed as percent to two decimal places.",
        )
    assert call.call_count == 3 and result.error is None
    assert result.answer == f"{100 * 34 / (34 + 3 + 5):.2f}"


def test_row_count_is_grounded_in_structure_but_unused_probe_cannot_ground_literal(
    tmp_path: Path,
) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    for row, values in enumerate([["Key", "Value"], ["A", 37], ["B", 61]], 307):
        for col, value in enumerate(values, 4):
            sheet.cell(row, col, value)
    path = tmp_path / "shifted_rows.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    count = execute_program("result=len([r for r in rows('Sheet') if r>307])", wire)
    assert count.success and count.result == 2
    assert count.reads and all(
        entry["kind"] == "type_probe" for entry in count.reads.values()
    )
    guessed = execute_program("unused=rows('Sheet'); result=2", wire)
    assert not guessed.success and "depend" in guessed.error


def test_typed_date_compared_as_space_string_to_iso_t_cutoff_is_rejected(
    tmp_path: Path,
) -> None:
    from src.contracts import temporal_comparison_errors

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Event"])
    sheet.append([dt.datetime(2029, 5, 1, 8)])
    sheet.append([dt.datetime(2029, 4, 30, 23)])
    path = tmp_path / "times.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    wrong = "r=data_rows('Sheet',{'A':'datetime'}); result=sum(1 for i in r if str(cell('Sheet','A'+str(i)))<'2029-05-01T00:00:00')"
    execution = execute_program(wrong, wire)
    query = "How many events occur strictly before 2029-05-01 00:00?"
    assert execution.result == 2 and temporal_comparison_errors(execution, query, wrong)
    assert not temporal_comparison_errors(
        execution, "Compare the timestamp strings lexicographically.", wrong
    )
    correct = "r=data_rows('Sheet',{'A':'datetime'}); result=sum(1 for i in r if cell('Sheet','A'+str(i))<datetime.datetime(2029,5,1))"
    checked = execute_program(correct, wire)
    assert checked.result == 1 and not temporal_comparison_errors(
        checked, query, correct
    )


def test_unused_source_elimination_preserves_result_and_complete_audit(
    tmp_path: Path,
) -> None:
    from unittest.mock import patch

    from src.llm_agent import LLMAgent

    book = openpyxl.Workbook()
    sheet = book.active
    for row, values in enumerate([["Name", "Units"], ["One", 37], ["Two", 61]], 307):
        for column, value in enumerate(values, 4):
            sheet.cell(row, column, value)
    path = tmp_path / "outside_evidence.xlsx"
    book.save(path)
    book.close()
    workbook = load_workbook_data(path)
    client = LLMAgent(api_key="test-placeholder", pricing=(0, 0))
    code = "unused=values('Sheet','E309:E309')[0]; answer=num('Sheet','E308')"
    program = {
        "code": code,
        "result_variable": "answer",
        "selection": [{"sheet": "Sheet", "range": "E308"}],
        "format": {"kind": "integer"},
    }
    with patch.object(client, "_request", return_value=json.dumps(program)) as call:
        result = client.solve_workbook(
            workbook, "What is Units for One? Answer only the integer."
        )
    assert call.call_count == 1 and result.answer == "37" and result.error is None
    assert "Sheet!E309" in result.attempts[0]["reads"]
    assert result.attempts[1]["contract_kind"] == "unused-source-elimination"
    assert "Sheet!E309" not in result.attempts[1]["reads"]
    bad = {"code": "unused=num('Sheet','D308'); result=num('Sheet','E308')"}
    with patch.object(
        client,
        "_request",
        side_effect=[
            json.dumps(bad),
            json.dumps({"code": "result=num('Sheet','E308')"}),
        ],
    ) as call:
        corrected = client.solve_workbook(
            workbook, "What is Units for One? Answer only the integer."
        )
    assert call.call_count == 2 and corrected.answer == "37"
    assert corrected.attempts[0]["error"]


def test_source_simplification_keeps_mutations_shadowed_helpers_and_closures() -> None:
    from src.executor import without_unused_source_reads

    for code in (
        "items=[1,2]; unused=items.pop(); result=len(items)",
        "x=cell('S','A1')\ndef f():\n    return x\nresult=f()",
        "def cell(s,a):\n    items.append(a)\nitems=[]\nunused=cell('S','A1')\nresult=len(items)",
    ):
        assert without_unused_source_reads(code) == code


def test_typed_row_filter_accepts_validated_inclusive_bounds(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    for row in (4, 310, 311, 420):
        sheet.cell(row, 5, "Group")
        sheet.cell(row, 6, row * 3)
    path = tmp_path / "bounded.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    for bounds in ("310,311", "start_row=310,end_row=311"):
        execution = execute_program(
            f"result=data_rows('Sheet',{{'E':'string','F':'number'}},{bounds})", wire
        )
        assert execution.result == [310, 311] and execution.error is None
        assert "Sheet!F420" not in execution.reads
    for bounds in ("311,310", "0,311", "310,True", "310,1048577"):
        execution = execute_program(
            f"result=data_rows('Sheet',{{'F':'number'}},{bounds})", wire
        )
        assert "bounds" in execution.error


def test_iso_age_ranking_uses_each_reference_and_intraday_precision(
    tmp_path: Path,
) -> None:
    from src.verification import age_reference_columns, age_sort_errors

    book = openpyxl.Workbook()
    sheet = book.active
    rows = [
        ["Name", "Birth date", "Admission date", "Audit date"],
        ["A", "2000-01-01T12:00:00", "2020-01-01T18:00:00", "2021-01-01"],
        ["B", "2000-01-01T00:00:00", "2020-01-01T12:00:00", "2020-01-01"],
    ]
    for row, values in enumerate(rows, 309):
        for col, value in enumerate(values, 4):
            sheet.cell(row, col, value)
    path = tmp_path / "ISO.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    question = "Who is oldest at admission?"
    assert age_reference_columns(wire["Sheet"], question) == ("E", "F")
    correct = execute_program(
        "r=age_order('Sheet','E','F',descending=True); result=cell('Sheet','D'+str(r[0]))",
        wire,
    )
    assert correct.result == "B" and not age_sort_errors(wire, correct, question)
    wrong = execute_program(
        "r=age_order('Sheet','E','F'); result=cell('Sheet','D'+str(r[0]))", wire
    )
    assert age_sort_errors(wire, wrong, question)
    assert age_reference_columns(wire["Sheet"], "Who is oldest at audit?") == ("E", "G")
    from src.verification import date_field_errors

    wrong_date = execute_program("result=cell('Sheet','F310')", wire)
    right_date = execute_program("result=cell('Sheet','G310')", wire)
    assert date_field_errors(
        wire, wrong_date, "What is A audit date? Answer only the ISO date."
    )
    assert not date_field_errors(
        wire, right_date, "What is A audit date? Answer only the ISO date."
    )
    assert not date_field_errors(
        wire, wrong_date, "What is A admission date? Answer only the ISO date."
    )


def test_ambiguous_date_text_requires_explicit_format(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Birth", "Reference"])
    sheet.append(["04/03/2000", "05/03/2020"])
    path = tmp_path / "locale.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    ambiguous = execute_program("result=age_order('Sheet','A','B')", wire)
    assert "date_format" in ambiguous.error
    explicit = execute_program(
        "result=age_order('Sheet','A','B',date_format='%d/%m/%Y')", wire
    )
    assert explicit.result == [2] and explicit.error is None
    days = execute_program(
        "result=age_days('Sheet','A2','B2',date_format='%d/%m/%Y')", wire
    )
    assert days.result == (dt.datetime(2020, 3, 5) - dt.datetime(2000, 3, 4)).days


@pytest.mark.parametrize(
    "offset,unit,factor", [(0, "thousands", 1000), (309, "millions", 1000000)]
)
def test_unrequested_scaling_is_detected_but_explicit_conversion_is_allowed(
    tmp_path: Path, offset: int, unit: str, factor: int
) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Unfamiliar scope"
    sheet.cell(offset + 1, 4, f"Amounts in {unit} of USD")
    for i, row in enumerate(
        [["Group", "Amount"], ["North", 47], ["South", 19]], offset + 2
    ):
        for col, value in enumerate(row, 4):
            sheet.cell(i, col, value)
    path = tmp_path / "new.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    code = f"scale={factor}; amount=num('Unfamiliar scope','E{offset + 3}'); result=amount*scale"
    result = execute_program(code, wire)
    question = "What is the Amount for North? Answer only the integer without units."
    assert scale_conversion_errors(wire, result, question, code, "result")
    assert scale_conversion_errors(
        wire, result, f"What is North Amount in {unit}?", code, "result"
    )
    assert scale_conversion_errors(
        wire,
        result,
        "What is North Amount? Format without thousands separators or units.",
        code,
        "result",
    )
    assert not scale_conversion_errors(
        wire, result, "Report the North Amount in actual USD.", code, "result"
    )
    assert not scale_conversion_errors(
        wire,
        result,
        f"Multiply the displayed North Amount by {factor}.",
        code,
        "result",
    )
    untouched = f"unused=2*{factor}; result=num('Unfamiliar scope','E{offset + 3}')"
    assert not scale_conversion_errors(
        wire, execute_program(untouched, wire), question, untouched, "result"
    )
    dynamic = f"scale={factor}\nfor r in [1]:\n    scale=num('Unfamiliar scope','E{offset + 4}')\nresult=num('Unfamiliar scope','E{offset + 3}')*scale"
    assert not scale_conversion_errors(
        wire, execute_program(dynamic, wire), question, dynamic, "result"
    )
    for output, output_question in [
        ("str(amount*scale)", question),
        (
            "f'{int(amount*scale)}, {int(amount*scale)}'",
            "Return two integers, separated by a comma and one space, without thousands separators or units.",
        ),
        ("[str(amount*scale), amount*scale]", "Return two numbers as integers."),
    ]:
        formatted_code = f"scale={factor}; amount=num('Unfamiliar scope','E{offset + 3}'); result={output}"
        formatted = execute_program(formatted_code, wire)
        assert formatted.success
        assert scale_conversion_errors(
            wire, formatted, output_question, formatted_code, "result"
        )
        assert not scale_conversion_errors(
            wire,
            formatted,
            "Report the amount in actual USD as integers.",
            formatted_code,
            "result",
        )
    label_code = f"amount=num('Unfamiliar scope','E{offset + 3}'); result='Large' if amount*{factor}>10 else 'Small'"
    assert not scale_conversion_errors(
        wire,
        execute_program(label_code, wire),
        "Which category label applies?",
        label_code,
        "result",
    )


def test_reporting_year_selection_uses_full_header_path_even_for_scalar_lookup(
    tmp_path: Path,
) -> None:
    from src.verification import column_scope_errors

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Relocated periods"
    sheet["D401"] = "Period table"
    sheet["E402"] = "2027"
    sheet.merge_cells("E402:F402")
    sheet["G402"] = "2028"
    sheet.merge_cells("G402:H402")
    for address, value in {
        "D403": "Region",
        "E403": "Units",
        "F403": "Completion rate",
        "G403": "Units",
        "H403": "Completion rate",
        "D404": "West",
        "E404": 43,
        "F404": 0.3,
        "G404": 71,
        "H404": 0.8,
    }.items():
        sheet[address] = value
    path = tmp_path / "periods.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    question = "For 2028, what is the West completion rate? Answer as percent."
    wrong = execute_program("result=100*num('Relocated periods','F404')", wire)
    right = execute_program("result=100*num('Relocated periods','H404')", wire)
    assert wrong.success and column_scope_errors(wire, wrong, question)
    assert column_scope_errors(wire, right, question) == []
    for historical in [
        "For 2028, what was the previous year's completion rate?",
        "For 2028, forecast completion rate from prior data.",
        "What is the change between 2027 and 2028?",
    ]:
        assert column_scope_errors(wire, wrong, historical) == []


def test_string_filter_header_count_is_reviewed_without_dropping_blank_records(
    tmp_path: Path,
) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Events"
    rows = [
        ["Identifier", "Date", "Value"],
        ["A", dt.datetime(2029, 2, 1), None],
        ["B", dt.datetime(2029, 2, 2), 83],
    ]
    for i, row in enumerate(rows, 311):
        for col, value in enumerate(row, 4):
            sheet.cell(i, col, value)
    path = tmp_path / "shifted.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    wrong = "r=data_rows('Events',{'D':'string'}); result=len(r)"
    result = execute_program(wrong, wire)
    query = "How many event records are in the table? Answer only the integer."
    assert result.result == 3 and header_count_errors(
        wire, result, query, wrong, "result"
    )
    assert not header_count_errors(
        wire, result, "How many rows, including headers?", wrong, "result"
    )
    right = "r=data_rows('Events',{'D':'string','E':'datetime'}); result=len(r)"
    checked = execute_program(right, wire)
    assert checked.result == 2 and not header_count_errors(
        wire, checked, query, right, "result"
    )


def test_signed_minimum_and_floating_tie_tolerance_are_distinct() -> None:
    query = "On which date was the smallest signed change? Break ties by earliest date."
    assert signed_comparison_errors(
        query, "if abs(change) < abs(best):\n    result=change"
    )
    assert not signed_comparison_errors(
        query, "if change < best or abs(change-best) < 1e-9:\n    result=change"
    )
    assert not signed_comparison_errors(
        "Find the smallest absolute change.", "if abs(change)<best:\n    result=change"
    )
    assert signed_comparison_errors(
        query,
        "candidates.append((abs(change), change, date))\ncandidates.sort(key=lambda x:x[0])",
    )


def test_explicit_row_dimensions_resolve_only_the_correct_year_and_region(
    tmp_path: Path,
) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Audited decisions"
    for i, row in enumerate(
        [
            ["Year", "Region", "Metric", "Value"],
            [2028, "East", "Approved", 71],
            [2028, "West", "Approved", 29],
            [2029, "East", "Approved", 113],
            [2029, "West", "Approved", 46],
        ],
        301,
    ):
        for col, value in enumerate(row, 4):
            sheet.cell(i, col, value)
    path = tmp_path / "unseen.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    query = "What is Approved Value for East in 2029?"
    correct = execute_program("result=num('Audited decisions','G304')", wire)
    incorrect = execute_program("result=num('Audited decisions','G302')", wire)
    assert not scope_risks(wire, correct, query)
    assert scope_risks(wire, incorrect, query)
    assert scope_risks(
        wire, incorrect, query, "result=sum([num('Audited decisions','G302')])"
    )


def test_flat_aggregate_qualifiers_are_or_within_column_and_and_across_columns(
    tmp_path: Path,
) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    for row, values in enumerate(
        [
            ["Year", "Region", "Metric", "Value"],
            [2028, "East", "Approved", 71],
            [2028, "West", "Approved", 29],
            [2029, "East", "Approved", 113],
            [2029, "West", "Approved", 46],
            [2029, "North", "Approved", 17],
        ],
        301,
    ):
        for column, value in enumerate(values, 4):
            sheet.cell(row, column, value)
    path = tmp_path / "dimensions.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    query = "What is total Approved Value for East and West in 2029?"
    correct_code = "result=sum(values('Sheet','G304:G305'))"
    correct = execute_program(correct_code, wire)
    assert correct.success and correct.result == 159
    assert scope_risks(wire, correct, query, correct_code) == []
    for area in ["G302:G305", "G304:G306"]:
        wrong_code = f"result=sum(values('Sheet',{area!r}))"
        wrong = execute_program(wrong_code, wire)
        assert scope_risks(wire, wrong, query, wrong_code)


def test_complete_relation_roles_do_not_fabricate_separate_sections(
    tmp_path: Path,
) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    rows = [
        ["Weight", "Origin", "Destination"],
        [0.6, "Birch", "Birch"],
        [0.4, "Birch", "Cedar"],
        [0.2, "Cedar", "Birch"],
        [0.8, "Cedar", "Cedar"],
    ]
    for row, values in enumerate(rows, 402):
        for column, value in enumerate(values, 10):
            sheet.cell(row, column, value)
    path = tmp_path / "links.xlsx"
    book.save(path)
    wire = load_workbook_data(path).wire()
    query = "Starting at Birch, what is the expected number of visits to Cedar?"
    code = "data=cells('Sheet','J403:L406'); result=data[1][0]+data[0][0]*data[1][0]+data[1][0]*data[3][0]"
    execution = execute_program(code, wire)
    assert execution.success
    assert not scope_risks(wire, execution, query, code)
    partial = code.replace("J403:L406", "J403:L404").replace("data[3][0]", "data[1][0]")
    assert scope_risks(wire, execute_program(partial, wire), query, partial)
    # A second physical table remains a genuine source-scope ambiguity.
    sheet["J450"] = "Other period"
    for row, values in enumerate(rows, 451):
        for column, value in enumerate(values, 10):
            sheet.cell(row, column, value)
    book.save(path)
    book.close()
    multiple = load_workbook_data(path).wire()
    assert scope_risks(multiple, execute_program(code, multiple), query, code)


def test_flat_ledger_aggregation_is_not_mistaken_for_repeated_sections(
    tmp_path: Path,
) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    rows = [
        ["Date", "Region", "SKU", "Units"],
        [dt.datetime(2029, 2, 1), "South", "K1", 13],
        [dt.datetime(2029, 2, 3), "South", "K2", 27],
        [dt.datetime(2029, 3, 2), "West", "K3", 52],
    ]
    for row, values in enumerate(rows, 303):
        for col, value in enumerate(values, 4):
            sheet.cell(row, col, value)
    path = tmp_path / "ledger.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    execution = execute_program(
        "r=data_rows('Sheet',{'D':'datetime','E':'string','G':'number'}); result=sum(num('Sheet','G'+str(i)) for i in r if cell('Sheet','E'+str(i)) in ['South','West'] and cell('Sheet','D'+str(i))<datetime.datetime(2029,3,1))",
        wire,
    )
    assert execution.result == 40
    assert not scope_risks(
        wire, execution, "What is the total Units for South and West before 2029-03-01?"
    )


def test_root_batch_metadata_cannot_drop_a_review_request(tmp_path: Path) -> None:
    import io

    from benchmarks.gateway import BatchLimits
    from benchmarks.structured_batch import StructuredBatchedAgent

    class Response(io.BytesIO):
        def __init__(self, data: bytes) -> None:
            super().__init__(data)
            self.headers = {"Content-Type": "application/json"}

    agent = StructuredBatchedAgent(
        api_key="test-placeholder",
        gateway_url="https://gateway.invalid/v1",
        gateway_model="ag/gemini-3-flash",
        batch_limits=BatchLimits(tmp_path / "batch.json"),
    )
    agent.structured_state.structured_expected = ["q0"]
    item = {"code": "result=num('S','B2')", "assumptions": ["local"]}
    body = {
        "programs": {"q0": item},
        "assumptions": ["shared"],
        "needs_verification": True,
        "selection": [],
    }
    raw = {
        "choices": [{"message": {"content": json.dumps(body)}, "finish_reason": "stop"}]
    }
    parsed = agent._read_response(Response(json.dumps(raw).encode()))
    program = json.loads(parsed["choices"][0]["message"]["content"])["programs"][0]
    assert program["needs_verification"] and program["assumptions"] == [
        "local",
        "shared",
    ]

    body["unsupported_answer"] = "42"
    raw["choices"][0]["message"]["content"] = json.dumps(body)
    parsed = agent._read_response(Response(json.dumps(raw).encode()))
    assert "unsupported_answer" in json.loads(
        parsed["choices"][0]["message"]["content"]
    )

    agent.structured_state.structured_expected = ["q0", "q1"]
    body = {
        "programs": {"q0": item, "q1": {"code": "result=num('S','C2')"}},
        "assumptions": [["q1", "second-specific"], ["q0", "first-specific"]],
        "needs_verification": True,
    }
    raw["choices"][0]["message"]["content"] = json.dumps(body)
    parsed = agent._read_response(Response(json.dumps(raw).encode()))
    programs = json.loads(parsed["choices"][0]["message"]["content"])["programs"]
    assert programs[0]["assumptions"] == ["local", "first-specific"]
    assert programs[1]["assumptions"] == ["second-specific"]
    assert programs[0]["needs_verification"] and programs[1]["needs_verification"]
    assert programs[0]["code"] == item["code"]
    for invalid in [
        [["q0", "first"], ["unknown", "other"]],
        [["q0", "first"], ["q0", "duplicate"]],
        [["q0", "first"]],
        [["q0", ["nested"]], ["q1", "second"]],
    ]:
        body["assumptions"] = invalid
        raw["choices"][0]["message"]["content"] = json.dumps(body)
        parsed = agent._read_response(Response(json.dumps(raw).encode()))
        assert json.loads(parsed["choices"][0]["message"]["content"]) == body


@pytest.mark.parametrize(
    "shared_selection", [[], None, {}, [{"sheet": "Other", "range": "A1"}]]
)
def test_shared_selection_only_accepts_empty_neutral_metadata(
    tmp_path: Path, shared_selection: object
) -> None:
    import io

    from benchmarks.gateway import BatchLimits
    from benchmarks.structured_batch import StructuredBatchedAgent

    class Response(io.BytesIO):
        def __init__(self, data: bytes) -> None:
            super().__init__(data)
            self.headers = {"Content-Type": "application/json"}

    agent = StructuredBatchedAgent(
        api_key="test-placeholder",
        gateway_url="https://gateway.invalid/v1",
        gateway_model="ag/gemini-3-flash",
        batch_limits=BatchLimits(tmp_path / "batch.json"),
    )
    agent.structured_state.structured_expected = ["q0", "q1"]
    child_selection = [{"sheet": "Moved", "range": "G508:G509"}]
    envelope = {
        "programs": {
            "q0": {"code": "result=num('Moved','G508')", "selection": child_selection},
            "q1": {"code": "result=num('Moved','G509')"},
        },
        "selection": shared_selection,
        "needs_verification": True,
    }
    original = json.dumps(envelope)
    response = {
        "choices": [{"message": {"content": original}, "finish_reason": "stop"}]
    }
    parsed = agent._read_response(Response(json.dumps(response).encode()))
    content = parsed["choices"][0]["message"]["content"]
    if shared_selection != []:
        assert content == original  # No nonempty/invalid declaration is discarded.
        return
    programs = json.loads(content)["programs"]
    assert len(programs) == 2 and all(p["needs_verification"] for p in programs)
    assert programs[0]["selection"][0]["range"] == "G508:G509"
    assert programs[1]["selection"] == []
    assert (
        "validated-empty-shared-selection"
        in agent.structured_state.envelope_normalization
    )
    assert agent.structured_state.original_content == original
    # Neutral metadata cannot repair missing keys or mask a competing root program.
    for invalid in [
        {**envelope, "code": "result=num('Other','A1')"},
        {**envelope, "programs": {"q0": envelope["programs"]["q0"]}},
    ]:
        text = json.dumps(invalid)
        response["choices"][0]["message"]["content"] = text
        rejected = agent._read_response(Response(json.dumps(response).encode()))
        assert rejected["choices"][0]["message"]["content"] == text


def test_nonexistent_columns_cannot_turn_into_a_zero_count(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Shifted source"
    sheet.cell(307, 4, "Identifier")
    sheet.cell(307, 5, "Score")
    sheet.cell(308, 4, "One")
    path = tmp_path / "empty scores.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    invalid = execute_program(
        "result=len(data_rows('Shifted source',{'A':'string','B':'number'}))", wire
    )
    assert not invalid.success and "no stored cells" in invalid.error
    genuinely_empty = execute_program(
        "result=len(data_rows('Shifted source',{'E':'number'}))", wire
    )
    assert genuinely_empty.success and genuinely_empty.result == 0


@pytest.mark.parametrize("offset_hours", [None, 7])
def test_timestamp_is_portable_for_naive_and_aware_input(
    offset_hours: int | None,
) -> None:
    from src.excel_parser import CellData, SheetData, WorkbookData

    timezone = (
        None if offset_hours is None else dt.timezone(dt.timedelta(hours=offset_hours))
    )
    value = dt.datetime(2029, 8, 13, 15, 30, tzinfo=timezone)
    wire = WorkbookData(
        {"Times": SheetData("Times", cells={"D308": CellData(value, "datetime")})}
    ).wire()
    result = execute_program("result=cell('Times','D308').timestamp()", wire)
    assert result.success
    assert result.result == pytest.approx(
        value.replace(tzinfo=timezone or dt.UTC).timestamp()
    )


def test_only_entity_scope_requests_review_of_unfiltered_reads(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "New ledger"
    sheet.append(["Entity", "Amount"])
    sheet.append(["Birch", 61])
    sheet.append(["Cedar", 9])
    path = tmp_path / "scope.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    question = "For Birch only, what is Amount? Answer only the integer."
    broad = execute_program("result=sum(values('New ledger','B2:B3'))", wire)
    precise = execute_program("result=num('New ledger','B2')", wire)
    assert any(
        "explicitly restricted" in risk for risk in scope_risks(wire, broad, question)
    )
    assert not scope_risks(wire, precise, question)


def test_bounded_combinations_supports_calculation_without_unbounded_pools() -> None:
    from src.excel_parser import CellData, SheetData, WorkbookData

    wire = WorkbookData(
        {
            "Values": SheetData(
                "Values",
                cells={
                    "D301": CellData(17, "number"),
                    "D302": CellData(29, "number"),
                    "D303": CellData(8, "number"),
                },
            )
        }
    ).wire()
    code = "import itertools\nv=values('Values','D301:D303')\nresult=max(sum(pair) for pair in itertools.combinations(v,2))"
    result = execute_program(code, wire)
    assert result.success and result.result == 46
    too_large = execute_program(
        "import itertools\nresult=list(itertools.combinations(range(1000),5))"
    )
    assert not too_large.success and "item budget" in too_large.error
    assert not execute_program("import itertools\nresult=itertools.__dict__").success


def test_across_all_compares_complete_sibling_columns_only_when_unambiguous(
    tmp_path: Path,
) -> None:
    import copy

    from src.verification import all_category_column_errors

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Relocated matrix"
    sheet["D301"] = "Metric"
    sheet["E301"] = "2028"
    sheet.merge_cells("E301:G301")
    for address, value in {
        "E302": "Amber",
        "F302": "Birch",
        "G302": "Copper",
        "D303": "Earlier",
        "E303": 2,
        "F303": 3,
        "G303": 4,
        "D304": "Later",
        "E304": 5,
        "F304": 7,
        "G304": 8,
    }.items():
        sheet[address] = value
    path = tmp_path / "changed.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    question = "How much greater is Later than Earlier across all categories?"
    partial = execute_program(
        "result=sum(values('Relocated matrix','F304:G304'))-sum(values('Relocated matrix','F303:G303'))",
        wire,
    )
    errors = all_category_column_errors(wire, partial.reads, question)
    assert len(errors) == 2 and any("E303" in error for error in errors)
    complete = execute_program(
        "result=sum(values('Relocated matrix','E304:G304'))-sum(values('Relocated matrix','E303:G303'))",
        wire,
    )
    assert complete.success and complete.result == 11
    assert all_category_column_errors(wire, complete.reads, question) == []
    for narrowed in [question + " Excluding Amber.", question + " For Birch only."]:
        assert all_category_column_errors(wire, partial.reads, narrowed) == []
    # A published parent or competing ancestry prevents inferring a partition.
    for replacement in [["2028", "Total"], ["Other measure", "Amber"]]:
        ambiguous = copy.deepcopy(wire)
        for table in ambiguous["Relocated matrix"]["tables"]:
            if "E" in table["headers"]:
                table["headers"]["E"] = replacement
        assert all_category_column_errors(ambiguous, partial.reads, question) == []
    measures = copy.deepcopy(wire)
    for table in measures["Relocated matrix"]["tables"]:
        if "E" in table["headers"]:
            table["headers"]["E"] = ["2028", "Published workers"]
            table["headers"]["F"] = ["2028", "Completion rate"]
            table["headers"]["G"] = ["2028", "Balance"]
    assert (
        all_category_column_errors(
            measures, partial.reads, "How many workers across all departments?"
        )
        == []
    )
    assert (
        all_category_column_errors(
            wire, partial.reads, question + " Use the published total."
        )
        == []
    )
    mixed = copy.deepcopy(wire)
    mixed["Relocated matrix"]["cells"]["F303"]["number_format"] = "0.0%"
    assert all(
        "E303" not in error
        for error in all_category_column_errors(mixed, partial.reads, question)
    )


def test_named_header_parent_requires_evidence_inside_that_branch(
    tmp_path: Path,
) -> None:
    import copy

    from src.verification import column_scope_errors

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Relocated branches"
    sheet["D401"] = "Metric"
    sheet["E401"] = "2031"
    sheet.merge_cells("E401:H401")
    for address, value in {
        "E402": "Overall",
        "F402": "Channel North",
        "H402": "Channel South",
        "F403": "Total",
        "G403": "Active only",
        "D404": "Units",
        "E404": 999,
        "F404": 30,
        "G404": 20,
        "H404": 50,
    }.items():
        sheet[address] = value
    sheet.merge_cells("F402:G402")
    path = tmp_path / "branches.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    question = "Compare Channel North and Channel South; answer as two integers."
    wrong = execute_program(
        "result=[num('Relocated branches','E404'),num('Relocated branches','H404')]",
        wire,
    )
    errors = column_scope_errors(wire, wrong, question)
    assert errors and "Channel North" in errors[0] and "F" in errors[0]
    right = execute_program(
        "result=[num('Relocated branches','F404'),num('Relocated branches','H404')]",
        wire,
    )
    assert column_scope_errors(wire, right, question) == []
    ambiguous = copy.deepcopy(wire)
    for table in ambiguous["Relocated branches"]["tables"]:
        if "F" in table["headers"]:
            table["headers"]["F"] = ["2031", "Older", "Channel North"]
            table["headers"]["G"] = ["2031", "Newer", "Channel North"]
    assert column_scope_errors(ambiguous, wrong, question) == []


def test_standard_filter_match_groups_and_datetime_feedback_are_safe() -> None:
    from src.excel_parser import CellData, SheetData, WorkbookData

    wire = WorkbookData(
        {"Moved": SheetData("Moved", cells={"K407": CellData("Code-318", "string")})}
    ).wire()
    filtered = execute_program(
        "s=cell('Moved','K407'); result=''.join(filter(str.isdigit,s))", wire
    )
    assert filtered.success and filtered.result == "318"
    matched = execute_program(
        "import re\ns=cell('Moved','K407'); result=re.search(r'([A-Za-z]+)-(\\d+)',s).groups()",
        wire,
    )
    assert matched.success and matched.result == ["Code", "318"]
    mapped = execute_program(
        "s=cell('Moved','K407'); result=sum(map(int,filter(str.isdigit,s)))", wire
    )
    assert mapped.success and mapped.result == 12
    wrong = execute_program("result=datetime(2031,5,7)")
    assert not wrong.success and "datetime.datetime" in wrong.error
    assert execute_program("result=datetime.datetime(2031,5,7).day").result == 7
    for unsafe in [
        "result=str.__dict__",
        "import re\nresult=re.search('x','x').__class__",
        "import re\nresult=re.search('x','x').re",
    ]:
        assert not execute_program(unsafe).success


def test_import_diagnostics_identify_rejected_module_and_actual_exports() -> None:
    for code, rejected, available in [
        (
            "from scipy.stats import rankdata\nresult=1",
            "scipy.stats",
            "Available modules",
        ),
        ("import pandas\nresult=1", "pandas", "Available modules"),
        (
            "from statistics import rankdata\nresult=1",
            "statistics.rankdata",
            "mean, median",
        ),
        ("from . import math\nresult=1", "Import not allowed: .", "Available modules"),
    ]:
        result = execute_program(code)
        assert (
            not result.success
            and rejected in result.error
            and available in result.error
        )
    assert (
        execute_program("from statistics import mean\nresult=mean([2,4])").result == 3
    )


def test_flat_loop_accumulation_is_not_misread_as_duplicate_sections(
    tmp_path: Path,
) -> None:
    from src.verification import scope_risks

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "State movements"
    for row, values in enumerate(
        [
            ["Origin", "Destination", "Rate"],
            ["Up", "Up", 0.75],
            ["Up", "Down", 0.25],
            ["Down", "Up", 0.5],
            ["Down", "Down", 0.5],
        ],
        309,
    ):
        for col, value in enumerate(values, 7):
            sheet.cell(row, col, value)
    path = tmp_path / "transitions.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    code = """data=cells('State movements','G310:I313')
transitions={(r[0],r[1]):r[2] for r in data}
prob={'Up':1.0,'Down':0.0}
for step in range(2):
    updated={'Up':0.0,'Down':0.0}
    for source in prob:
        for target in prob:
            updated[target]+=prob[source]*transitions[(source,target)]
    prob=updated
result=prob['Up']"""
    execution = execute_program(code, wire)
    assert execution.success and execution.result == 0.6875
    question = "Starting in Up, what is the probability of Up after two steps?"
    assert scope_risks(wire, execution, question, code) == []
    reward_code = code.rsplit("result=", 1)[0] + "result=2*prob['Up']"
    reward_question = "Starting in Up, what is expected reward after two steps, with 2 for Up and 0 for Down?"
    reward = execute_program(reward_code, wire)
    assert reward.success and reward.result == 1.375
    assert scope_risks(wire, reward, reward_question, reward_code) == []
    from unittest.mock import patch

    from src.llm_agent import LLMAgent

    client = LLMAgent(api_key="test-placeholder", pricing=(0, 0))
    with patch.object(
        client, "_request", return_value=json.dumps({"code": reward_code})
    ) as call:
        solved = client.solve_workbook(load_workbook_data(path), reward_question)
    assert call.call_count == 1 and solved.error is None
    assert float(solved.answer) == 1.375
    lookup = execute_program("result=num('State movements','I310')", wire)
    assert scope_risks(
        wire, lookup, "What is the Rate for Up?", "result=num('State movements','I310')"
    )


def test_cells_rows_cannot_silently_compare_equal_to_scalar() -> None:
    from src.excel_parser import CellData, SheetData, WorkbookData

    wire = WorkbookData(
        {
            "Status": SheetData(
                "Status",
                cells={
                    "H311": CellData("Active", "string"),
                    "H312": CellData("Closed", "string"),
                },
            )
        }
    ).wire()
    for compare in [
        "row == 'Active'",
        "'Active' != row",
        "row in ['Active','Closed']",
        "row not in ['Closed']",
        "'Active' in cells('Status','H311:H312')",
    ]:
        result = execute_program(
            f"row=cells('Status','H311')[0]; result={compare}", wire
        )
        assert not result.success and "2D list" in result.error
    correct = execute_program(
        "result=sum(1 for r in [311,312] if cell('Status','H'+str(r))=='Active')", wire
    )
    assert correct.success and correct.result == 1
    structured = execute_program(
        "row=cells('Status','H311')[0]; row.append('extra'); result=[row==['Active','extra'], cells('Status','H312')==[['Closed']], values('Status','H311')=='Active']",
        wire,
    )
    assert structured.success and structured.result == [True, True, False]
    membership = execute_program(
        "result=[cells('Status','H311')[0] in cells('Status','H311:H312'), 'Active' in values('Status','H311:H312')]",
        wire,
    )
    assert membership.success and membership.result == [True, True]


@pytest.mark.parametrize("audit", [False, True])
def test_review_mode_audits_code_without_supplying_computed_answer(
    tmp_path: Path, audit: bool
) -> None:
    from unittest.mock import patch

    from src.llm_agent import LLMAgent

    book = openpyxl.Workbook()
    book.active.append(["Entity", "Units"])
    book.active.append(["Cedar", 47])
    path = tmp_path / "audit.xlsx"
    book.save(path)
    book.close()
    workbook = load_workbook_data(path)
    code = "result=num('Sheet','B2')*3"
    first = {"code": code, "needs_verification": True, "format": {"kind": "integer"}}
    second = {**first, "needs_verification": False}
    agent = LLMAgent(api_key="test-placeholder", pricing=(0, 0), review_candidate=audit)
    verification_states = []
    replies = iter([json.dumps(first), json.dumps(second)])

    def respond(messages: list[dict[str, str]], state: object) -> str:
        from src.llm_agent import QueryState

        assert isinstance(state, QueryState)
        verification_states.append(state.verification)
        return next(replies)

    with patch.object(agent, "_request", side_effect=respond) as request:
        solved = agent.solve_workbook(workbook, "What is three times Units for Cedar?")
    assert solved.error is None and solved.answer == "141"
    assert verification_states == [False, True]
    review = request.call_args_list[1].args[0][-1]["content"]
    assert ("UNVERIFIED CANDIDATE" in review) is audit
    assert (code in review) is audit
    assert "141" not in review
    assert solved.attempts[0]["verification_mode"] == (
        "candidate-audit" if audit else "fresh-selection"
    )


def test_missing_record_filter_difference_cannot_count_verified_heading(
    tmp_path: Path,
) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    for row, values in enumerate(
        [["Participant", "Reading"], ["Birch", 19], ["Cedar", None], ["Pine", None]],
        407,
    ):
        for column, value in enumerate(values, 6):
            sheet.cell(row, column, value)
    path = tmp_path / "missing.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    question = "How many participants have a blank Reading? Answer only the integer."
    bad = (
        "all_rows=data_rows('Sheet',{'F':'string'})\n"
        "valid=set(data_rows('Sheet',{'F':'string','G':'number'}))\n"
        "result=len([r for r in all_rows if r not in valid])"
    )
    execution = execute_program(bad, wire)
    assert execution.success and execution.result == 3
    assert header_count_errors(wire, execution, question, bad, "result")
    good = bad.replace("{'F':'string'})", "{'F':'string'},start_row=408,end_row=410)")
    checked = execute_program(good, wire)
    assert checked.success and checked.result == 2
    assert not header_count_errors(wire, checked, question, good, "result")
    assert not header_count_errors(
        wire,
        execution,
        "How many rows, including headers, are not numeric?",
        bad,
        "result",
    )


def test_permutations_and_set_operations_remain_bounded() -> None:
    from src.excel_parser import CellData, SheetData, WorkbookData
    from src.executor import bounded_permutations

    wire = WorkbookData(
        {
            "Moved": SheetData(
                "Moved",
                cells={
                    "G401": CellData(17, "number"),
                    "G402": CellData(29, "number"),
                    "G403": CellData(8, "number"),
                },
            )
        }
    ).wire()
    result = execute_program(
        "from itertools import permutations\nv=values('Moved','G401:G403')\nresult=max(10*a+b for a,b in permutations(v,2))",
        wire,
    )
    assert result.success and result.result == 307
    result = execute_program(
        "a=frozenset(values('Moved','G401:G402')); b=set(values('Moved','G402:G403')); result=[sorted(a|b),sorted(a&b),sorted(a^b)]",
        wire,
    )
    assert result.success and result.result == [[8, 17, 29], [29], [8, 17]]
    assert list(bounded_permutations([1, 2])) == [(1, 2), (2, 1)]
    assert list(bounded_permutations([1, 2], 3)) == []
    assert list(bounded_permutations([], 0)) == [()]
    for count in (-1, 1.5, True):
        with pytest.raises((ValueError, TypeError)):
            bounded_permutations([1, 2], count)
    with pytest.raises(MemoryError, match="item budget"):
        bounded_permutations(range(1000), 1000)
    assert not execute_program("result=2**4097 | 1").success
    assert not execute_program("def f():\n    yield 1\nresult=list(f())").success


def test_extremum_requires_measure_and_returns_identity_instead_of_measure(
    tmp_path: Path,
) -> None:
    from src.verification import label_errors

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Buckets"
    sheet.append(["Duration category", "Share"])
    sheet.append(["Short", 17])
    sheet.append(["Long", 63])
    numeric = book.create_sheet("Numeric labels")
    numeric.append(["Category", "Share"])
    numeric.append([1, 21])
    numeric.append([2, 48])
    path = tmp_path / "buckets.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    query = "Which duration category had the highest share percentage? Answer with only the duration category."
    guessed = "result=cell('Buckets','A2')"
    assert extremum_evidence_errors(
        wire, execute_program(guessed, wire), query, guessed
    )
    wrong = execute_program("result=label('Buckets',argmax('Buckets','B2:B3'))", wire)
    assert wrong.success and label_errors(wrong, query, wire)
    correct = "best=argmax('Buckets','B2:B3'); result=cell('Buckets','A'+best[1:])"
    checked = execute_program(correct, wire)
    assert checked.result == "Long" and not label_errors(checked, query, wire)
    assert not extremum_evidence_errors(wire, checked, query, correct)
    numeric_answer = execute_program("result=cell('Numeric labels','A3')", wire)
    assert not label_errors(
        numeric_answer,
        "Which category has highest Share? Answer with only the category.",
        wire,
    )
