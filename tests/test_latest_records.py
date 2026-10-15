"""Latest-record selection is a typed table operation, independent of dataset IDs."""

import datetime as dt
from pathlib import Path

import openpyxl

from src.excel_parser import load_workbook_data
from src.executor import execute_program
from src.verification import latest_reference_columns


def test_latest_selection_precedes_status_filter_and_supports_merged_keys(
    tmp_path: Path,
) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    rows = [
        ["Member", "Action time", "State", "Amount"],
        ["Alpha", dt.datetime(2029, 1, 1), "Active", 99],
        [None, dt.datetime(2029, 1, 3), "Inactive", 110],
        ["Beta", dt.datetime(2029, 1, 1), "Inactive", 18],
        ["Beta", dt.datetime(2029, 1, 4), "Active", 43],
        ["Gamma", dt.datetime(2029, 1, 2), "Active", 37],
    ]
    for row, values in enumerate(rows, 307):
        for col, value in enumerate(values, 4):
            sheet.cell(row, col, value)
    sheet.merge_cells("D308:D309")
    path = tmp_path / "shifted.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    query = "For each Member use its latest Action time."
    assert latest_reference_columns(wire["Sheet"], query) == ("D", "E")
    code = "r=latest_rows('Sheet','D','E'); active=[i for i in r if cell('Sheet','F'+str(i))=='Active']; active.sort(key=lambda i:(-num('Sheet','G'+str(i)),label('Sheet','D'+str(i)))); result=[label('Sheet','D'+str(i)) for i in active]"
    execution = execute_program(code, wire)
    assert execution.success and execution.result == ["Beta", "Gamma"]
    assert "Sheet!E309" in execution.reads and "Sheet!D308" in execution.reads


def test_only_equal_maximum_timestamps_require_tie_policy(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Member", "Event", "Amount"])
    sheet.append(["A", "2029-01-01", 1])
    sheet.append(["A", "2029-01-01", 2])
    sheet.append(["A", "2029-01-03", 3])
    path = tmp_path / "events.xlsx"
    book.save(path)
    wire = load_workbook_data(path).wire()
    assert execute_program("result=latest_rows('Sheet','A','B')", wire).result == [4]
    sheet.append(["A", "2029-01-03", 4])
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    ambiguous = execute_program("result=latest_rows('Sheet','A','B')", wire)
    assert not ambiguous.success and "tie policy" in ambiguous.error
    assert execute_program(
        "result=latest_rows('Sheet','A','B',ties='last')", wire
    ).result == [5]
    assert execute_program(
        "result=latest_rows('Sheet','A','B',ties='first')", wire
    ).result == [4]


def test_missing_tail_timestamp_and_wrong_columns_are_not_silently_dropped(
    tmp_path: Path,
) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Member", "Event", "Amount"])
    sheet.append(["A", dt.datetime(2029, 1, 1), 1])
    sheet.append(["B", None, 2])
    path = tmp_path / "missing.xlsx"
    book.save(path)
    book.close()
    wire = load_workbook_data(path).wire()
    missing = execute_program("result=latest_rows('Sheet','A','B')", wire)
    assert not missing.success and "no timestamp" in missing.error
    wrong = execute_program("result=latest_rows('Sheet','X','Y')", wire)
    assert not wrong.success and "existing workbook columns" in wrong.error
    for bounds in ("start_row=False", "start_row=3,end_row=2", "end_row=1048577"):
        invalid = execute_program(f"result=latest_rows('Sheet','A','B',{bounds})", wire)
        assert not invalid.success and "bounds" in invalid.error
